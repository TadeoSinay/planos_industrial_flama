"""Genera salida/palletizado_PT.xlsx: pallets por camión, disposición por pallet y frecuencia de envíos.

Sólo cuentas (pedido del usuario): Datos -> Produccion -> Pallet -> Camion -> Escenarios. Todo lo calculado es
fórmula; negro y grises.
"""

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "salida")

# ------------------------------------------------------------------ datos de las fuentes
MP = "MP - Abastecimiento y Almacenamiento, hoja Pallets PT"
TIPOS = ["1 kg", "2,5 kg", "5 kg", "10 kg", "25 kg", "50 kg", "70 kg", "100 kg"]
ANIOS = list(range(2026, 2036))
# Proyecciones 5, Indexado Dim Técnico, filas 3 a 10 (matafuegos terminados) y 15 a 22 (cilindros vendidos vacíos)
TERMINADOS = [
    [4178.736543, 5997.303977, 6851.747758, 8467.135549, 12680.63271, 15620.25307, 20332.55326, 25132.07848, 28573.33614, 32072.63913],
    [295.0303572, 709.7274503, 986.8680576, 1500.638174, 2312.627726, 3252.869887, 4102.315524, 5104.306587, 6195.186243, 7424.858444],
    [2255.773025, 3846.978881, 5093.537591, 6823.115313, 10819.30694, 14605.96411, 18877.75984, 23368.18637, 26967.72455, 30750.77872],
    [91.83318153, 177.2472967, 252.9787798, 358.0204941, 602.7181519, 872.605032, 1127.160545, 1395.444812, 1612.350054, 1840.848079],
    [34.05179136, 58.46214745, 75.20853998, 100.5298529, 155.8760035, 211.9344827, 272.6682332, 337.8486707, 393.6274156, 453.2705332],
    [74.913941, 128.6167244, 165.458788, 221.1656765, 342.9272077, 466.255862, 599.8701129, 743.2670756, 865.9803143, 997.1951731],
    [16.34485985, 28.06183078, 36.10009919, 48.25432941, 74.82048167, 101.7285517, 130.8807519, 162.1673619, 188.9411595, 217.569856],
    [10.89657324, 18.70788719, 24.06673279, 32.16955294, 49.88032112, 67.81903447, 87.25383461, 108.1115746, 125.960773, 145.0465706],
]
CILINDROS = [
    [0, 0, 7662.4, 26694.8, 49434.8, 71680.2, 101340.6, 131001, 146819.8, 170548],
    [0, 0, 299.6, 1043.8, 1933, 2803, 3963, 5123, 5741.4, 6669],
    [0, 0, 1478.8, 5152, 9540.8, 13834.2, 19558.6, 25283, 28335.8, 32915],
    [0, 0, 90.4, 315, 583.4, 845.8, 1195.4, 1545, 1731.4, 2011],
    [0, 0, 14.4, 50.2, 93, 135, 191, 247, 277, 322],
    [0, 0, 31.8, 110.8, 205.2, 297.6, 420.8, 544, 609.6, 708],
    [0, 0, 7, 24.4, 45.2, 65.4, 92.2, 119, 133.4, 155],
    [0, 0, 4.6, 16, 29.6, 42.8, 60.4, 78, 87.6, 102],
]
# Dimensionamiento Técnico FLAMA, DATOS-BASE filas 102 a 104 (los mismos del MRP)
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
PERFILES = [("1. Automovilístico", [.12, .09, .08, .075, .07, .07, .075, .075, .075, .075, .075, .12]),
            ("2. Manual comercial", [.10, .08, .085, .085, .08, .075, .08, .08, .08, .075, .07, .11]),
            ("3. Industrial sobre ruedas", [.085, .075, .09, .09, .09, .08, .09, .09, .085, .075, .07, .08])]
PERFIL_DE = [1, 2, 2, 2, 3, 3, 3, 3]
CAJA = {"u": [6, 1, 1, 1], "L": [265, 135, 165, 195], "A": [190, 135, 165, 195], "H": [410, 435, 515, 685],
        "tara": [0.35, 0.25, 0.35, 0.5], "masa": [2.08, 4.63, 8.48, 16.23]}
CIL = {"D": [76.2, 124, 153, 181.5, 276.5, 320, 390, 390], "H": [297.5, 333, 414, 582, 672, 858, 908, 1128],
       "m": [0.76, 1.48, 2.78, 5.48, 15.92, 23.3, 46.86176284, 56.84631739]}
CARRO = {"L": [600, 700, 800, 800], "A": [550, 600, 700, 750], "m": [55, 100, 140, 185]}

# ------------------------------------------------------------------ estilo (sólo negro y grises)
F_N = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
F_T = Font(name="Arial", size=12, bold=True)
F_W = Font(name="Arial", size=10, bold=True, color="FFFFFF")
NEGRO = PatternFill("solid", fgColor="000000")
GRIS = PatternFill("solid", fgColor="D9D9D9")
GRIS_C = PatternFill("solid", fgColor="F2F2F2")
fino = Side(style="thin", color="808080")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)
DER = Alignment(horizontal="right", vertical="center")


def hoja(nombre, anchos):
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False
    for k, a in enumerate(anchos):
        ws.column_dimensions[L(k + 1)].width = a
    return ws


def tit(ws, r, t):
    ws.cell(r, 1, t).font = F_T


def cab(ws, r, textos):
    for k, t in enumerate(textos):
        c = ws.cell(r, 1 + k, t)
        c.font, c.fill, c.alignment, c.border = F_W, NEGRO, CEN, BORDE


def celda(ws, r, c, v, fmt=None, dato=False, neg=False, res=False):
    x = ws.cell(r, c, v)
    x.font = F_B if (neg or res) else F_N
    x.border = BORDE
    if dato:
        x.fill = GRIS
    elif res:
        x.fill = GRIS_C
    if fmt:
        x.number_format = fmt
    x.alignment = IZQ if isinstance(v, str) and not v.startswith("=") else (CEN if fmt is None else DER)
    return x


def fila(ws, r, nombre, unidad, valores, fmt, dato=False, res=False, fuente=None, c_fuente=None):
    celda(ws, r, 1, nombre, neg=res)
    celda(ws, r, 2, unidad).alignment = CEN
    for k, v in enumerate(valores):
        celda(ws, r, 3 + k, v, fmt, dato=dato, res=res)
    if fuente:
        celda(ws, r, c_fuente or 3 + len(valores), fuente)


wb = Workbook()
wb.remove(wb.active)

# ================================================================== Datos
ws = hoja("Datos", [44, 8, 11, 11, 11, 11, 11, 11, 11, 11, 70])
D = {}


def bloque(r, titulo_, columnas, filas_, fmt_de):
    """filas_: [(clave, nombre, unidad, valores, fuente)]. Devuelve la fila siguiente."""
    tit(ws, r, titulo_)
    cab(ws, r + 1, ["Parámetro", "Unidad"] + columnas + [""] * (8 - len(columnas)) + ["Fuente"])
    r += 2
    for clave, nom, un, vals, fu in filas_:
        fila(ws, r, nom, un, vals, fmt_de(clave), dato=True, fuente=fu, c_fuente=11)
        D[clave] = r
        r += 1
    return r + 1


r = bloque(1, "General", ["Valor"], [
    ("anio", "Año de análisis", "año", [2035], "Año de diseño de la planta"),
    ("sem_mes", "Semanas por mes", "sem", [4], "20 días hábiles ÷ 5 (MP, hoja Volumenes)"),
], lambda c: "0")
dv = DataValidation(type="list", formula1='"' + ",".join(map(str, ANIOS)) + '"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"C{D['anio']}")
r = bloque(r, "Pallet", ["Valor"], [
    ("PL", "Largo", "mm", [1200], "ARLOG IRAM 10016 (" + MP + ")"),
    ("PA", "Ancho", "mm", [1000], "ARLOG IRAM 10016 (" + MP + ")"),
    ("PH", "Alto", "mm", [145], "ARLOG IRAM 10016 (" + MP + ")"),
    ("PT", "Tara", "kg", [35], "ARLOG IRAM 10016 (" + MP + ")"),
    ("HMAX", "Altura máxima con la carga", "mm", [1600], "Envolvedora EDOS, 1.000 a 1.700 mm (" + MP + ")"),
    ("WMAX", "Peso máximo de la carga", "kg", [1100], "EDOS 600 a 1.200 kg; ARLOG 1.500 kg (" + MP + ")"),
    ("SE", "Separador de cartón: espesor", "mm", [3], MP + " (estimación)"),
    ("SM", "Separador de cartón: masa", "kg", [0.6], MP + " (estimación)"),
], lambda c: "0.0" if c == "SM" else "#,##0")
r = bloque(r, "Matafuegos manuales en caja", TIPOS[:4], [
    ("cu", "Matafuegos por caja", "u", CAJA["u"], "ET's Matafuegos (" + MP + ", sección B)"),
    ("cL", "Caja: largo", "mm", CAJA["L"], MP + ", sección B"),
    ("cA", "Caja: ancho", "mm", CAJA["A"], MP + ", sección B"),
    ("cH", "Caja: alto", "mm", CAJA["H"], MP + ", sección B"),
    ("ct", "Caja: tara", "kg", CAJA["tara"], MP + ", sección B (estimación)"),
    ("cm", "Matafuego lleno con accesorios", "kg", CAJA["masa"], MP + ", sección B"),
], lambda c: "0.00" if c in ("ct", "cm") else "#,##0")
r = bloque(r, "Cilindros vendidos vacíos", TIPOS, [
    ("iD", "Ø", "mm", CIL["D"], "MP, hoja Geometria Acero"),
    ("iH", "Alto con tapón", "mm", CIL["H"], "MP, hoja Geometria Acero"),
    ("im", "Masa", "kg", CIL["m"], "MP, hoja Geometria Acero"),
], lambda c: "0.00" if c == "im" else "#,##0.0")
r = bloque(r, "Rodantes terminados", TIPOS[4:], [
    ("rL", "Carro: largo", "mm", CARRO["L"], MP + ", sección C"),
    ("rA", "Carro: ancho", "mm", CARRO["A"], MP + ", sección C"),
    ("rm", "Peso cargado", "kg", CARRO["m"], "Ficha técnica Georgia ABC 90 (FLAMA_BOM, hoja Mercado)"),
], lambda c: "#,##0")
r = bloque(r, "Camión: Sider GyL Full para balancín (carga por los laterales)", ["Valor"], [
    ("SL", "Largo útil", "m", [7.40], "7,50 m (GyL, agroads.com.ar/detalle.asp?clasi=957986) − 0,10 m (criterio)"),
    ("SA", "Ancho útil", "m", [2.46], "2,60 m (GyL) − 0,14 m de parantes y lonas (criterio)"),
    ("SCU", "Carga útil", "kg", [12000], "Balancín sider 7,50 m: 12/14 t (logisticashc.com.ar/nosotros/nuestra-flota)"),
], lambda c: "#,##0" if c == "SCU" else "0.00")
r = bloque(r, "Furgón: Renault Master L3H2 (rodantes)", ["Valor"], [
    ("ML", "Largo de carga", "m", [3.857], "vansdirect.co.uk, Renault Master dimensions"),
    ("MA", "Ancho de carga", "m", [1.765], "vansdirect.co.uk"),
    ("MAp", "Ancho entre pasos de rueda", "m", [1.38], "vansdirect.co.uk (renault.com.mt da 1,73 para la Master nueva: se toma el menor)"),
    ("MLp", "Largo de los pasos de rueda", "m", [1.0], "Criterio"),
    ("MLt", "Largo libre detrás de los pasos de rueda", "m", [0.6], "Criterio"),
    ("MCU", "Carga útil", "kg", [1593], "autocosmos.com.ar, Master L3H2 Furgón 2.3"),
], lambda c: "#,##0" if c == "MCU" else "0.000")
ws.freeze_panes = "A3"


def d(clave, k=0):
    return f"Datos!${L(3 + k)}${D[clave]}"


# ================================================================== Produccion
ws = hoja("Produccion", [12] + [10] * 20)
tit(ws, 1, "Producción anual (Proyecciones 5, hoja Indexado Dim Técnico)")
COLS = ["Term. " + t for t in TIPOS] + ["Cil. " + t for t in TIPOS]
cab(ws, 2, ["Año"] + COLS)
a0 = 3
for j, a in enumerate(ANIOS):
    celda(ws, a0 + j, 1, a, "0", neg=True)
    for k in range(8):
        celda(ws, a0 + j, 2 + k, TERMINADOS[k][j], "#,##0", dato=True)
        celda(ws, a0 + j, 10 + k, CILINDROS[k][j], "#,##0", dato=True)
a1 = a0 + len(ANIOS) - 1
SEL = a1 + 1
celda(ws, SEL, 1, "Año elegido", neg=True)
for c in range(2, 18):
    celda(ws, SEL, c, f"=INDEX({L(c)}{a0}:{L(c)}{a1},MATCH({d('anio')},$A${a0}:$A${a1},0))", "#,##0", res=True)
celda(ws, SEL + 1, 1, "Perfil", neg=True)
for c in range(2, 18):
    celda(ws, SEL + 1, c, PERFIL_DE[(c - 2) % 8], "0", dato=True)

r = SEL + 3
tit(ws, r, "Estacionalidad (Dimensionamiento Técnico FLAMA, DATOS-BASE filas 102 a 104)")
cab(ws, r + 1, ["Perfil"] + MESES + ["Total"])
P0 = r + 2
for k, (nom, v) in enumerate(PERFILES):
    celda(ws, P0 + k, 1, nom)
    for j, x in enumerate(v):
        celda(ws, P0 + k, 2 + j, x, "0.0%", dato=True)
    celda(ws, P0 + k, 14, f"=SUM(B{P0 + k}:M{P0 + k})", "0%")
ws.column_dimensions["A"].width = 22

r = P0 + 4
tit(ws, r, "Producción por semana del año elegido (u/semana = anual × % del mes ÷ semanas por mes)")
GRUPOS = [("Mix term. 2,5/5/10", [1, 2, 3], 0), ("Mix term. 25 a 100", [4, 5, 6, 7], 0),
          ("Mix cil. 2,5/5/10", [1, 2, 3], 8), ("Mix cil. 25 a 100", [4, 5, 6, 7], 8)]
cab(ws, r + 1, ["Mes"] + COLS + [g[0] for g in GRUPOS])
S0 = r + 2
for j in range(12):
    rr = S0 + j
    celda(ws, rr, 1, MESES[j], neg=True)
    for c in range(2, 18):
        celda(ws, rr, c, f"={L(c)}${SEL}*INDEX($B${P0}:$M${P0 + 2},{L(c)}${SEL + 1},{j + 1})/{d('sem_mes')}",
              "#,##0")
    for g, (_, ks, off) in enumerate(GRUPOS):
        celda(ws, rr, 18 + g, "=" + "+".join(f"{L(2 + off + k)}{rr}" for k in ks), "#,##0")
S1 = S0 + 11
celda(ws, S1 + 1, 1, "Semana pico", neg=True)
for c in range(2, 22):
    celda(ws, S1 + 1, c, f"=MAX({L(c)}{S0}:{L(c)}{S1})", "#,##0", res=True)
ws.freeze_panes = "B3"


def pcol(k, cil=False):
    return L(2 + k + (8 if cil else 0))


# ================================================================== Pallet
ws = hoja("Pallet", [36, 7] + [10] * 8)
PL, PA, PH, PT, HM, WM = d("PL"), d("PA"), d("PH"), d("PT"), d("HMAX"), d("WMAX")
tit(ws, 1, "Matafuegos manuales en caja")
cab(ws, 2, ["", "Unidad"] + TIPOS[:4])
PC = {c: 3 + i for i, c in enumerate(["pc", "nl", "na", "cc", "ch", "cp", "ca", "cj", "U", "W", "Hp"])}


def colrow(clave, k):
    return f"{L(3 + k)}{PC[clave]}"


def caja_f(clave, k):
    cL, cA, cH = d("cL", k), d("cA", k), d("cH", k)
    o1 = f"INT({PL}/{cL})*INT({PA}/{cA})"
    o2 = f"INT({PL}/{cA})*INT({PA}/{cL})"
    c = lambda x: colrow(x, k)  # noqa: E731
    return {
        "pc": f"={d('cu', k)}*{d('cm', k)}+{d('ct', k)}",
        "nl": f"=IF({o1}>={o2},INT({PL}/{cL}),INT({PL}/{cA}))",
        "na": f"=IF({o1}>={o2},INT({PA}/{cA}),INT({PA}/{cL}))",
        "cc": f"={c('nl')}*{c('na')}",
        "ch": f"=INT(({HM}-{PH})/{cH})",
        "cp": f"=INT({WM}/({c('cc')}*{c('pc')}))",
        "ca": f"=MIN({c('ch')},{c('cp')})",
        "cj": f"={c('cc')}*{c('ca')}",
        "U": f"={c('cj')}*{d('cu', k)}",
        "W": f"={c('cj')}*{c('pc')}+{PT}",
        "Hp": f"={PH}+{c('ca')}*{cH}",
    }[clave]


for clave, nom, un, fmt in [("pc", "Peso por caja", "kg", "0.00"), ("nl", "Cajas a lo largo (1.200)", "u", "#,##0"),
                            ("na", "Cajas a lo ancho (1.000)", "u", "#,##0"), ("cc", "Cajas por capa", "u", "#,##0"),
                            ("ch", "Capas por altura", "u", "#,##0"), ("cp", "Capas por peso", "u", "#,##0"),
                            ("ca", "Capas", "u", "#,##0"), ("cj", "Cajas por pallet", "u", "#,##0"),
                            ("U", "Matafuegos por pallet", "u", "#,##0"), ("W", "Peso del pallet", "kg", "#,##0"),
                            ("Hp", "Altura del pallet", "mm", "#,##0")]:
    fila(ws, PC[clave], nom, un, [caja_f(clave, k) for k in range(4)], fmt, res=clave in ("cc", "ca", "U"))

r0 = PC["Hp"] + 2
tit(ws, r0, "Cilindros vendidos vacíos (parados, separador entre capas)")
cab(ws, r0 + 1, ["", "Unidad"] + TIPOS)
PI = {c: r0 + 2 + i for i, c in enumerate(["nl", "na", "cc", "ch", "cp", "ca", "U", "W", "Hp"])}


def cil_f(clave, k):
    iD, iH, im = d("iD", k), d("iH", k), d("im", k)
    c = lambda x: f"{L(3 + k)}{PI[x]}"  # noqa: E731
    return {
        "nl": f"=INT({PL}/{iD})",
        "na": f"=INT({PA}/{iD})",
        "cc": f"={c('nl')}*{c('na')}",
        "ch": f"=INT(({HM}-{PH})/({iH}+{d('SE')}))",
        "cp": f"=INT({WM}/({c('cc')}*{im}+{d('SM')}))",
        "ca": f"=MAX(1,MIN({c('ch')},{c('cp')}))",
        "U": f"={c('cc')}*{c('ca')}",
        "W": f"={c('U')}*{im}+({c('ca')}+1)*{d('SM')}+{PT}",
        "Hp": f"={PH}+{c('ca')}*{iH}+({c('ca')}+1)*{d('SE')}",
    }[clave]


for clave, nom, un, fmt in [("nl", "Cilindros a lo largo (1.200)", "u", "#,##0"),
                            ("na", "Cilindros a lo ancho (1.000)", "u", "#,##0"),
                            ("cc", "Cilindros por capa", "u", "#,##0"), ("ch", "Capas por altura", "u", "#,##0"),
                            ("cp", "Capas por peso", "u", "#,##0"), ("ca", "Capas", "u", "#,##0"),
                            ("U", "Cilindros por pallet", "u", "#,##0"), ("W", "Peso del pallet", "kg", "#,##0"),
                            ("Hp", "Altura del pallet", "mm", "#,##0")]:
    fila(ws, PI[clave], nom, un, [cil_f(clave, k) for k in range(8)], fmt, res=clave in ("cc", "ca", "U"))

# ================================================================== Camion
ws = hoja("Camion", [36, 7] + [10] * 4)
tit(ws, 1, "Sider GyL Full para balancín")
cab(ws, 2, ["", "Unidad", "Valor"])
SL, SA = d("SL"), d("SA")
o1 = f"INT({SL}*1000/{PA})*INT({SA}*1000/{PL})"
o2 = f"INT({SL}*1000/{PL})*INT({SA}*1000/{PA})"
fila(ws, 3, "Pallets a lo largo", "u", [f"=IF({o1}>={o2},INT({SL}*1000/{PA}),INT({SL}*1000/{PL}))"], "#,##0")
fila(ws, 4, "Pallets a lo ancho", "u", [f"=IF({o1}>={o2},INT({SA}*1000/{PL}),INT({SA}*1000/{PA}))"], "#,##0")
fila(ws, 5, "Pallets por camión", "u", ["=C3*C4"], "#,##0", res=True)
SIDER = "Camion!$C$5"

tit(ws, 7, "Renault Master L3H2")
cab(ws, 8, ["", "Unidad"] + TIPOS[4:])
CM = {"zf": 9, "zp": 10, "zt": 11, "lug": 12, "pes": 13, "N": 14}


def zona(largo, ancho, k):
    l, a = f"{d('rL', k)}/1000", f"{d('rA', k)}/1000"
    return (f"=MAX(INT(ROUND({largo}/({l}),6))*INT(ROUND({ancho}/({a}),6)),"
            f"INT(ROUND({largo}/({a}),6))*INT(ROUND({ancho}/({l}),6)))")


fila(ws, 9, "Carros en la zona delantera", "u",
     [zona(f"({d('ML')}-{d('MLp')}-{d('MLt')})", d("MA"), k) for k in range(4)], "#,##0")
fila(ws, 10, "Carros entre los pasos de rueda", "u", [zona(d("MLp"), d("MAp"), k) for k in range(4)], "#,##0")
fila(ws, 11, "Carros en la zona trasera", "u", [zona(d("MLt"), d("MA"), k) for k in range(4)], "#,##0")
fila(ws, 12, "Carros por lugar", "u", [f"=SUM({L(3 + k)}9:{L(3 + k)}11)" for k in range(4)], "#,##0")
fila(ws, 13, "Carros por peso", "u", [f"=INT({d('MCU')}/{d('rm', k)})" for k in range(4)], "#,##0")
fila(ws, 14, "Carros por furgón", "u", [f"=MIN({L(3 + k)}12,{L(3 + k)}13)" for k in range(4)], "#,##0", res=True)

# ================================================================== Escenarios
ws = hoja("Escenarios", [26, 8] + [17] * 5)


def frecuencia(v):
    """Texto de la frecuencia y carga del vehículo; con menos de 1 viaje por semana sale cada INT(1/v) semanas."""
    txt = (f'=IF({v}=0,"-",IF({v}>=1,ROUNDUP({v},0)&" por semana",'
           f'IF(INT(1/{v})=1,"1 por semana","1 cada "&INT(1/{v})&" semanas")))')
    carga = f"=IF({v}=0,0,IF({v}>=1,{v}/ROUNDUP({v},0),{v}*INT(1/{v})))"
    return txt, carga


def bloque_esc(r, titulo_, ks, cil, furgon):
    tit(ws, r, titulo_)
    nombres = [f"{i + 1}. Todo {TIPOS[k]}" for i, k in enumerate(ks)] + [
        "5. Mix " + ("2,5 + 5 + 10" if ks[0] == 0 else "25 a 100")]
    cab(ws, r + 1, ["", "Unidad"] + nombres)
    mixk = ks[1:] if ks[0] == 0 else ks
    g = [i for i, gg in enumerate(GRUPOS) if gg[1] == mixk and (gg[2] == 8) == cil][0]
    r += 2
    if furgon:
        R = {"N": r, "UV": r + 1}
        N = [f"=Camion!{L(3 + k - 4)}{CM['N']}" for k in ks]
        prods = [f"Produccion!{pcol(k)}${SEL}" for k in mixk]
        sp = "+".join(prods)
        den = "+".join(f"{p}/Camion!{L(3 + k - 4)}{CM['lug']}" for k, p in zip(mixk, prods))
        peso = "+".join(f"{p}*{d('rm', k - 4)}" for k, p in zip(mixk, prods))
        N.append(f"=IF(({sp})=0,0,MIN(INT(({sp})/({den})),INT({d('MCU')}/(({peso})/({sp})))))")
        fila(ws, R["N"], "Carros por furgón", "u", N, "#,##0", res=True)
        fila(ws, R["UV"], "Unidades por viaje", "u", [f"={L(3 + i)}{R['N']}" for i in range(5)], "#,##0", res=True)
        r += 2
    else:
        R = {"U": r, "W": r + 1, "N": r + 2, "UV": r + 3}
        hojap = "Pallet"
        RR = PI if cil else PC
        U = [f"={hojap}!{L(3 + k)}{RR['U']}" for k in ks]
        W = [f"={hojap}!{L(3 + k)}{RR['W']}" for k in ks]
        prods = [f"Produccion!{pcol(k, cil)}${SEL}" for k in mixk]
        sp = "+".join(prods)
        den = "+".join(f"IF({hojap}!{L(3 + k)}{RR['U']}=0,0,{p}/{hojap}!{L(3 + k)}{RR['U']})"
                       for k, p in zip(mixk, prods))
        U.append(f"=IF(({den})=0,0,({sp})/({den}))")
        W.append(f"=IF(({den})=0,0,(" + "+".join(f"{p}/{hojap}!{L(3 + k)}{RR['U']}*{hojap}!{L(3 + k)}{RR['W']}"
                                                  for k, p in zip(mixk, prods)) + f")/({den}))")
        fila(ws, R["U"], "Unidades por pallet", "u", U, "#,##0")
        fila(ws, R["W"], "Peso del pallet", "kg", W, "#,##0")
        fila(ws, R["N"], "Pallets por camión", "u",
             [f"=IF({L(3 + i)}{R['W']}=0,0,MIN({SIDER},INT({d('SCU')}/{L(3 + i)}{R['W']})))" for i in range(5)],
             "#,##0", res=True)
        fila(ws, R["UV"], "Unidades por viaje", "u", [f"={L(3 + i)}{R['N']}*{L(3 + i)}{R['U']}" for i in range(5)],
             "#,##0", res=True)
        r += 4
    celda(ws, r, 1, "Viajes por semana", neg=True)
    r += 1
    m0 = r
    cols_p = [pcol(k, cil) for k in ks] + [L(18 + g)]
    for j in range(12):
        fila(ws, m0 + j, MESES[j], "viajes",
             [f"=IF({L(3 + i)}${R['UV']}=0,0,Produccion!{cols_p[i]}{S0 + j}/{L(3 + i)}${R['UV']})" for i in range(5)],
             "0.00")
    rp = m0 + 12
    fila(ws, rp, "Mes pico", "viajes", [f"=MAX({L(3 + i)}{m0}:{L(3 + i)}{m0 + 11})" for i in range(5)], "0.00",
         res=True)
    tx, ca = zip(*[frecuencia(f"{L(3 + i)}{rp}") for i in range(5)])
    fila(ws, rp + 1, "Frecuencia (mes pico)", "", list(tx), None, res=True)
    fila(ws, rp + 2, "Carga del vehículo", "%", list(ca), "0%")
    return rp + 4


r = bloque_esc(1, "Matafuegos manuales en caja (sider)", [0, 1, 2, 3], False, False)
r = bloque_esc(r, "Rodantes (Master L3H2)", [4, 5, 6, 7], False, True)
r = bloque_esc(r, "Cilindros de manuales (sider)", [0, 1, 2, 3], True, False)
r = bloque_esc(r, "Cilindros de rodantes (sider)", [4, 5, 6, 7], True, False)

salida = sys.argv[1] if len(sys.argv) > 1 else os.path.join(SALIDA, "palletizado_PT.xlsx")
os.makedirs(os.path.dirname(salida), exist_ok=True)
wb.save(salida)
print(salida)
