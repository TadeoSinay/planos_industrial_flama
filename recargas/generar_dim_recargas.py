"""Genera salida/Dimensionamiento_Recargas_Planta_Comercial.xlsx y ..._Industrial.xlsx.

Mismo modelo para las dos plantas; cada archivo aplica su propio % de recargas (Proyecciones, Indexado Dim
Técnico V61 / V63; PBA Interior queda afuera porque la atienden distribuidores ajenos). Hojas: Links ->
PROY_Indexado -> datos_base -> parametros_recargas -> sobrestock_sustitutos -> Plan mensual -> Demanda por puesto
-> Resumen Capacidad (aprovechamiento consolidado). Todo lo calculado es fórmula; los IMPORTRANGE leen el link de la
hoja Links (los de Manuales y Carros quedan para pegar). Máquinas y turnos se cargan con la regla del
aprovechamiento (ningún puesto pasa el 90 %) y se pueden cambiar a mano. Negro y grises.
"""

import math
import os
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "salida")

# ------------------------------------------------------------------ datos de las fuentes
ANIOS = list(range(2026, 2036))
# IMPORTRANGE a Proyecciones, con el formato con que Google lo exporta a .xlsx (al subirlo a Sheets vuelve a ser
# IMPORTRANGE; en Excel muestra el último valor)
URL_PROY = "https://docs.google.com/spreadsheets/d/1xc4xDAwSYT3l-RulPYtcePxiCninrbQGzaBUsgO0YEs/edit?gid=1764491891#gid=1764491891"
IMP_REC = "Indexado Dim Técnico!B27:K43"
IMP_ZON = "Indexado Dim Técnico!U61:U74"
# hoja Links: (archivo, link, qué se importa). Los que dicen PEGAR se completan cuando estén los Sheets nuevos.
PEGAR = "PEGAR ACÁ EL LINK DEL GOOGLE SHEET"
LINKS = [("Proyecciones", URL_PROY,
          "Indexado Dim Técnico: recargas por tipo (B27:K43), % por zona (U61:U74), % por planta (V61 y V63)"),
         ("Dimensionamiento Manuales", PEGAR,
          "parametros_manuales K32, K33, K38, K39, K40 y K42: tiempos de PH, secado, carga, ensamblaje, presurización "
          "y etiquetado (equivalencia del rodante)"),
         ("Dimensionamiento Carros", PEGAR,
          "parametros_carros I29, I30, I33 e I34: tiempos ponderados de PH, secado, carga y presurización, y montaje "
          "y etiquetado (equivalencia del rodante)"),
         ("MRP (FLAMA_MRP)", PEGAR,
          "ESTACIONALIDAD C4:N6: perfiles de producción E01 (1 kg), E02 (2,5 a 10 kg) y E03 (rodantes), con los que "
          "llegan de fábrica los sustitutos ABC nuevos")]
LK_PROY, LK_MAN, LK_CAR, LK_MRP = "Links!$B$4", "Links!$B$5", "Links!$B$6", "Links!$B$7"
# Proyecciones, Indexado Dim Técnico V61 (comercial) y V63 (industrial): suma de las zonas de cada planta
PCT_PLANTA = {"Comercial": 0.5545197389, "Industrial": 0.2882365055}
ANIO_2T = 2032  # año desde el que se agrega turno antes que máquina (Dimensionamiento Técnico V2 y Manuales)
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
# Proyecciones, Indexado Dim Técnico, filas 27 a 43 (recargas por tipo, 2026 a 2035)
REC = [
    ("ABC 1 kg", [3477.482124, 16970.58564, 30546.08951, 56048.31106, 84270.22613, 117407.8231, 147610.6356, 171404.3607, 183690.566, 196282.6907]),
    ("ABC 2,5 kg", [329.996668, 1287.414024, 2947.169216, 5426.043269, 7850.875806, 10423.94365, 12809.89571, 14263.06831, 15432.86983, 16620.76117]),
    ("ABC 5 kg", [1950.394097, 7852.61852, 17775.85701, 33605.56835, 48407.57189, 64323.04276, 78721.10039, 87237.2124, 94103.66761, 101095.1762]),
    ("ABC 10 kg", [118.343632, 437.716578, 1024.974821, 1960.409862, 2844.976846, 3810.390622, 4663.67807, 5164.514584, 5571.544328, 5985.950855]),
    ("Agua 10 l", [159.308736, 643.423542, 1474.362741, 2766.451205, 3991.465221, 5335.75578, 6538.958887, 7269.819681, 7855.44139, 8450.814081]),
    ("AFFF 10 l", [22.758391, 161.315751, 316.607276, 527.39462, 743.788946, 962.874335, 1194.942207, 1351.617183, 1477.54886, 1604.441422]),
    ("AFFF 50 l s/ruedas", [11.379195, 83.22352, 164.463834, 281.341145, 398.902509, 514.527284, 629.065505, 711.178024, 773.41098, 836.350805]),
    ("BC 5 kg", [6.827517, 47.073112, 88.889676, 159.502491, 219.796202, 291.869853, 368.844169, 408.408953, 445.781944, 483.479203]),
    ("Sales K 6 l", [81.930207, 335.398855, 746.598352, 1358.899837, 1961.125522, 2598.541412, 3197.263477, 3578.539452, 3870.747997, 4167.559622]),
    ("CO₂ 2 kg", [450.616139, 1697.18149, 3955.196858, 7561.26033, 10920.14481, 14514.88522, 17746.49927, 19617.20546, 21161.80625, 22734.50171]),
    ("CO₂ 5 kg", [125.17115, 643.630371, 1375.626719, 2460.284265, 3545.257683, 4723.461784, 5853.905178, 6608.459909, 7185.708397, 7769.577127]),
    ("HCFC/HFC 5 kg", [18.206713, 112.075122, 230.854041, 400.715334, 574.507789, 754.345623, 944.84884, 1072.234671, 1171.695753, 1271.941099]),
    ("Clase D 9 l", [6.827517, 47.073112, 88.889676, 148.279182, 205.226053, 255.765145, 328.417902, 366.908071, 401.681707, 436.686782]),
    ("ABC 25 kg s/ruedas", [17.637753, 85.987819, 186.869567, 336.583621, 484.132598, 643.220385, 792.431742, 889.677707, 965.31111, 1041.939702]),
    ("ABC 50 kg s/ruedas", [38.803056, 189.173203, 411.113048, 740.483966, 1065.091715, 1415.084848, 1743.349833, 1957.290955, 2123.684441, 2292.267343]),
    ("ABC 70 kg s/ruedas", [8.466121, 41.274153, 89.697392, 161.560138, 232.383647, 308.745785, 380.367236, 427.045299, 463.349333, 500.131057]),
    ("ABC 100 kg s/ruedas", [5.644081, 27.516102, 59.798262, 107.706759, 154.922431, 205.830523, 253.578158, 284.696866, 308.899555, 333.420704]),
]
# Proyecciones, Indexado Dim Técnico, filas 61 a 74, columna U (% zonal 2035) y planta según V61 / V63. PBA Interior
# queda fuera: la atienden distribuidores ajenos
ZONAS = [
    ("CABA Norte", 0.0521797399, "Comercial"),
    ("CABA Este", 0.0798836162, "Comercial"),
    ("CABA Sur", 0.0598172860, "Industrial"),
    ("CABA Oeste", 0.0637650712, "Comercial"),
    ("PBA Sur 1er (Avellaneda)", 0.0676801452, "Industrial"),
    ("PBA Sur 2do cordón", 0.0932172985, "Industrial"),
    ("PBA Sur 3er cordón", 0.0675217758, "Industrial"),
    ("PBA Norte 1er cordón", 0.0533558834, "Comercial"),
    ("PBA Norte 2do cordón", 0.0885682049, "Comercial"),
    ("PBA Norte 3er cordón", 0.0471361107, "Comercial"),
    ("PBA Oeste 1er cordón", 0.0848900664, "Comercial"),
    ("PBA Oeste 2do cordón", 0.0672643024, "Comercial"),
    ("PBA Oeste 3er cordón", 0.0174767438, "Comercial"),
    ("PBA Interior", 0.1572437555, "Fuera"),
]
# Proyecciones, hojas PROYECCION <rubro>: cociente sobre RECARGAS An (es el mismo en todos los tipos y rubros)
TASAS = [
    ("Mat Usados ÷ RECARGAS An", 0.05, "Se usa: equipos que llegan descargados por uso"),
    ("Mat PH ÷ RECARGAS An", 0.02, "Se usa: equipos que no pasan la PH y se inutilizan"),
    ("Mat Rot ÷ RECARGAS An", 0.01, "Se usa: equipos rotos que se inutilizan"),
    ("Mat Fin VU ÷ RECARGAS An", 0.03302041313770093,
     "No se usa: Proyecciones aplica la misma vida útil a todos los tipos; acá va la del Anexo E (20 años, CO₂ 30)"),
]
AGENTES = ["ABC", "BC", "Clase D", "Agua", "AFFF", "Sales K", "CO₂", "Halogenado"]
# agente de ruta, rodante (u. equivalentes), manguera con PH anual (IRAM 3517-2 9.7.1.4)
SIN_M = "No: la manguera no tiene cierre en el extremo"
TIPOS = [
    ("ABC", 0, 0, "No: tobera sin manguera"), ("ABC", 0, 0, SIN_M), ("ABC", 0, 0, SIN_M), ("ABC", 0, 0, SIN_M),
    ("Agua", 0, 0, SIN_M), ("AFFF", 0, 0, SIN_M),
    ("AFFF", 1, 1, "Sí, 9.7.1.4 a): lanza con cierre (rodante, IRAM 3541)"),
    ("BC", 0, 0, SIN_M), ("Sales K", 0, 0, SIN_M),
    ("CO₂", 0, 0, "No: difusor de brazo, sin manguera (catálogo FLAMA)"),
    ("CO₂", 0, 1, "Sí, 9.7.1.4 b): manguera de CO₂ (difusor con manga)"),
    ("Halogenado", 0, 0, SIN_M), ("Clase D", 0, 0, SIN_M),
] + [("ABC", 1, 1, "Sí, 9.7.1.4 a): válvula de cierre en la manguera (rodante, IRAM 3550)")] * 4

# Filas fijas de las hojas
P_MIX = 57     # PROY_Indexado: primera fila de la matriz de esta planta (tipos en filas, rubros en columnas)
FILA_RUB = 35      # primera fila de la estacionalidad por rubro (7 rubros)
# datos de la estacionalidad: escuelas estatales, patentamientos 2025 por mes, 2024, flota y % exento de VTV
F_ESTATAL, F_PAT, F_P24, F_FLOTA, F_EXENTO = 44, 45, 46, 47, 48
FILA_TIPO = 55     # primera fila de la tabla de tipos (con su perfil mensual en E a P)
FILA_PUESTO = 76   # primera fila de la tabla de puestos
N_PUESTOS = 29
FILA_PCIL = 109    # primera fila del perfil de llegada de los sustitutos ABC nuevos (E01, E02, E03 del MRP)
# Equivalencia de un rodante: operaciones que comparten la fábrica y la recarga. (operación, celdas de
# parametros_manuales con su valor (s/u), celda de parametros_carros con su valor (min/u, ponderado con el mix 2035))
FILA_EQ = 116
EQUIV = [("Prueba hidráulica", [("K32", 60)], ("I29", 8.979935)),
         ("Secado", [("K33", 120)], ("I30", 10.058961)),
         ("Carga de polvo y presurización", [("K38", 45), ("K40", 30)], ("I33", 11.33)),
         ("Ensamblaje / montaje, precinto y etiquetado", [("K39", 100), ("K42", 60)], ("I34", 6.32))]
# Parámetros: (fila, texto, valor o fórmula, formato, dato a mano, fuente)
P_GEN = [
    (2, "Días hábiles por mes", 20, "0", True, "240 días hábiles por año ÷ 12"),
    (3, "Horas por día (1 turno)", 8, "0", True, ""),
    (4, "Refrigerio (min/turno)", 30, "0", True, ""),
    (5, "Horas productivas por turno", "=B3-B4/60", "0.00", False, ""),
    (6, "Horas productivas por mes (1 turno)", "=B2*B5", "0", False, ""),
    (7, "Umbral de aprovechamiento", 0.9, "0%", True,
     "10 % de holgura para variabilidad, averías y mantenimiento"),
    (8, "Turnos máximos por día", 2, "0", True, "hasta 2 turnos de 8 h; después se agrega máquina"),
    (9, "Equivalencia de un rodante (u. equivalentes)", f"=D{FILA_EQ + len(EQUIV)}", "0.0", False,
     "Promedio de la relación de tiempos carros ÷ manuales de FLAMA en las operaciones comunes (bloque EQUIVALENCIA "
     "DE UN RODANTE, al final de esta hoja)"),
    (10, "Año de inicio del 2° turno (antes: más máquinas)", ANIO_2T, "0", True,
     "Prioridad del aprovechamiento: antes de ese año se agrega máquina u operario; desde ese año, turno"),
]
# Criterios de recarga por agente (parametros_recargas filas 13 a 28, agentes en E a L como en la tabla de puestos):
# (fila, qué le pasa al equipo, unidad, fuente, estado, valor por agente, explicación). "{c}" = columna del agente.
F_PH, F_INU, F_USO, F_CAMBIO, F_VALV, F_PINT, F_HUM, F_LABD, F_LAB = 16, 21, 22, 23, 24, 25, 26, 27, 28
CRITERIOS = [
    (13, "Intervalo de prueba hidráulica", "años", "IRAM 3517-2:2020, Anexo E", "Norma", [5, 5, 5, 2, 2, 2, 5, 5],
     "Polvos, CO₂ y halogenados cada 5 años; agua, AFFF y sales K cada 2"),
    (14, "Vence la prueba hidráulica", "% por año", "1 ÷ intervalo", "Norma", ["=1/{c}13"] * 8,
     "9.7.2: en régimen vence la PH de 1 de cada 5 (o 1 de cada 2) equipos por año"),
    (15, "Golpe o abolladura leve: va a PH", "% por año", "Criterio del grupo", "A confirmar", [0.05] * 8,
     "9.7.1.2 y Word (inspección visual). El Dimensionamiento Técnico V2 usaba 25 % de PH en polvos: 20 % por "
     "vencimiento + 5 % por daño"),
    (16, "Pasa por prueba hidráulica", "% por año", "Vence + golpe", "Cálculo", ["={c}14+{c}15"] * 8, ""),
    (17, "Vida útil", "años", "IRAM 3517-2:2020, Anexo E", "Norma", [20, 20, 20, 20, 20, 20, 30, 20],
     "CO₂: 30 años (IRAM 3509)"),
    (18, "Llega al fin de la vida útil", "% por año", "1 ÷ vida útil", "Norma", ["=1/{c}17"] * 8,
     "Se inutiliza en la recepción (Word, operación 1)"),
    (19, "No pasa la prueba hidráulica", "% por año", "Proyecciones (Mat PH)", "Dato",
     ["=PROY_Indexado!$B$47"] * 8,
     "Equivale a 1 de cada 10 de los que se prueban (2 % ÷ 20 %), en línea con lo que informan los servicios de PH"),
    (20, "Se rompe en el uso", "% por año", "Proyecciones (Mat Rot)", "Dato", ["=PROY_Indexado!$B$48"] * 8, ""),
    (21, "Se inutiliza", "% por año", "Fin de vida + no pasa la PH + roto", "Cálculo", ["={c}18+{c}19+{c}20"] * 8,
     "9.12 y Anexo G"),
    (22, "Llega descargado por uso", "% por año", "Proyecciones (Mat Usados)", "Dato", ["=PROY_Indexado!$B$46"] * 8,
     "9.2.2: se recarga después de usarlo. El polvo sin presión no se descarga; el CO₂ descargado va directo a la "
     "carga"),
    (23, "Cambio obligatorio del agente", "% por año", "Norma y Word", "Norma", ["=1/10", 1, 1, 1, 1, 1, 0, 0],
     "ABC: dígito del año de fabricación (9.9.4 a). BC: siempre polvo nuevo (Word). Clase D: se descarga y carga a "
     "mano. Agua, AFFF y sales K: el agente se cambia cada año (Anexo E). CO₂ y halogenados se reutilizan"),
    (24, "Hay que abrir la válvula (daño, pérdida de peso o presión)", "% por año", "Criterio", "A confirmar",
     [0, 0, 0, 0, 0, 0, 0.05, 0.05],
     "Solo CO₂ y halogenados: el resto se abre siempre. Word: se recupera el agente por daño, cambio de piezas o "
     "pérdida de peso o presión"),
    (25, "Se pinta", "% por año", "Criterio del grupo", "A confirmar", [0.15] * 8,
     "9.10: óxido, metal a la vista, pérdida de brillo o repintado parcial; no fija el %. El Dimensionamiento "
     "Técnico V2 usaba 15 %"),
    (26, "Humedad interior al desarmar (secado extra)", "% por año", "Criterio", "A confirmar",
     [0.01, 0.01, 0.01, 0, 0, 0, 0, 0], "Word, operación 7: se seca también si en el desarme aparece humedad"),
    (27, "Ensayos de laboratorio por día", "ensayos/día", "Criterio (Word)", "A confirmar", [1, 0, 0, 0, 0, 0, 0, 0],
     "Word, operación 5: control periódico de la tolva de polvo recuperado (uno por día), dudas en el visor y lotes "
     "nuevos sin Sello IRAM. El BC se carga con polvo nuevo con sello"),
    (28, "Ensayos de laboratorio por equipo", "ensayos/u.", "Ensayos por día × días del año ÷ recargas del agente "
     "en 2035", "Cálculo",
     [f"=IF({{c}}27=0,0,{{c}}27*$B$2*12/SUMIF($B${FILA_TIPO}:$B${FILA_TIPO + 16},{{c}}$12,"
      f"datos_base!$K$2:$K$18))"] * 8,
     "1 control diario de la tolva repartido en las recargas del agente en 2035"),
]
# Sustituto del mismo tipo que el equipo retirado; el ABC 1 kg automotor no lleva (sin canje, criterio del usuario)
LLEVA = [0] + [1] * 16
DIA_RETIRO, DIA_DEVOL = 1, 1  # el sustituto sale con el camión que retira y vuelve con el que devuelve
SS = "sobrestock_sustitutos"
S_RUB = 17     # paso 1: lote, tiempo y ciclo por rubro (rubros en columnas B a H)
S_TIPO = 30    # paso 2: tipos
S_STOCK = 51   # paso 3: stock
S_ALTA = 73    # paso 4: alta anual (fábrica / compra); Manuales y Carros importan sus filas ABC
# Proyecciones, hojas PROYECCION <rubro>, 2035: recargas de mercado por rubro y tipo en las zonas de cada planta
# (industrial = CABA Sur y PBA Sur; comercial = el resto de CABA y PBA sin Interior). Los carros se reparten
# 25/55/12/8 % entre 25, 50, 70 y 100 kg como en PxQ (Corregido), filas 66 a 69.
RUBROS = ["Automotor", "Edificios", "Educación", "Hoteles", "Industrias / almacenes", "Salud", "Supermercados"]
MIX_RUBRO = {
    "Comercial": [
        [2251135, 0, 69628, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0.0, 0.0, 0.0, 0.0],
        [0, 100435.9, 732591.4, 44900.8, 56716.8, 0, 0, 0, 20087.2, 183147.8, 25995.2, 0, 0, 4431.0, 9748.2, 2126.9,
         1417.9],
        [0, 14173.8, 52786.7, 5673.3, 4876.4, 0, 0, 0, 4459.0, 19050.2, 891.8, 0, 0, 199.2, 438.3, 95.6, 63.8],
        [3578.5, 45655.1, 8847.8, 2123.5, 1415.6, 0, 0, 0, 3539.1, 3539.1, 1769.6, 0, 0, 177.0, 389.3, 84.9, 56.6],
        [0, 0, 113431.7, 0, 12568.5, 15671.5, 9151.4, 4634.7, 8169.7, 6402.2, 32678.6, 9269.4, 3535.0, 4055.3, 8921.8,
         1946.6, 1297.7],
        [0, 7336.9, 46903.8, 6270.5, 8958.8, 0, 0, 0, 2804.4, 10549.9, 11830.5, 2425.5, 0, 1364.4, 3001.6, 654.9,
         436.6],
        [530, 6153, 5342, 0, 562, 0, 0, 0, 0, 0, 2530, 0, 0, 0.0, 0.0, 0.0, 0.0],
    ],
    "Industrial": [
        [1188664, 0, 36760, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0.0, 0.0, 0.0, 0.0],
        [0, 51231.6, 373689.5, 22903.5, 28930.8, 0, 0, 0, 10246.3, 93422.4, 13260.0, 0, 0, 2260.2, 4972.5, 1084.9,
         723.3],
        [0, 8719.2, 32472.3, 3490.0, 2999.8, 0, 0, 0, 2743.0, 11718.9, 548.6, 0, 0, 122.6, 269.6, 58.8, 39.2],
        [1533.8, 23764.4, 6552.8, 1572.7, 1048.4, 0, 0, 0, 2621.1, 2621.1, 1310.5, 0, 0, 131.1, 288.3, 62.9, 41.9],
        [0, 0, 44086.3, 0, 4377.3, 6232.7, 3220.9, 1816.2, 3384.2, 2600.2, 13536.8, 3632.5, 1568.0, 1589.2, 3496.3,
         762.8, 508.5],
        [0, 4173.4, 26680.3, 3566.9, 5096.0, 0, 0, 0, 1595.2, 6001.1, 6729.6, 1379.7, 0, 776.1, 1707.4, 372.5,
         248.3],
        [331, 3794, 3289, 0, 346, 0, 0, 0, 0, 0, 1558, 0, 0, 0.0, 0.0, 0.0, 0.0],
    ],
}
# Equipos que se retiran juntos en cada cliente (u. por cliente de Recargas.xlsx, hoja Supuestos, dato del grupo).
# Control de sentido: Dec. 351/79, art. 176 (un matafuego cada 200 m² como mínimo)
LOTE = [25, 14, 18, 20, 45, 35, 16]
M2_MATAFUEGO = 200
# Perfiles mensuales del MRP (FLAMA_MRP, hoja ESTACIONALIDAD, filas 4 a 6; IMPORTRANGE con estos valores):
# E01 automotor 1 kg, E02 manual comercial (2,5 a 10 kg), E03 rodantes (25 a 100 kg)
E_MRP = [("E01 automotor (1 kg)", "ESTACIONALIDAD!C4:N4",
          [0.12, 0.09, 0.08, 0.075, 0.07, 0.07, 0.075, 0.075, 0.075, 0.075, 0.075, 0.12]),
         ("E02 manual comercial (2,5 a 10 kg)", "ESTACIONALIDAD!C5:N5",
          [0.1, 0.08, 0.085, 0.085, 0.08, 0.075, 0.08, 0.08, 0.08, 0.075, 0.07, 0.11]),
         ("E03 rodantes (25 a 100 kg)", "ESTACIONALIDAD!C6:N6",
          [0.085, 0.075, 0.09, 0.09, 0.09, 0.08, 0.09, 0.09, 0.085, 0.075, 0.07, 0.08])]
FAM_ABC = {0: 0, 1: 1, 2: 1, 3: 1, 13: 2, 14: 2, 15: 2, 16: 2}   # tipo ABC -> perfil del MRP de su línea
# Estacionalidad de las recargas por rubro (IRAM 3517-2:2020 9.2.1 y 9.4.14: cada equipo vuelve en el mes de su último
# mantenimiento, así que solo hay estacionalidad donde una regla o un calendario fija el mes para todo el rubro)
ESTATAL = 0.682   # unidades educativas de gestión estatal, PBA (HCD PBA, proyecto 13-14 D-1382, con el Relevamiento Anual)
# Automotor: patentamientos 2025 por mes (ACARA, total del mercado), 2024 (ACARA) y flota circulante 2025 (AFAC)
PAT_2025 = [68988, 42918, 48389, 55017, 55363, 52212, 62123, 54664, 55827, 51982, 35424, 23997]
PAT_2024 = 414041
FLOTA_2025 = 15784385
EST_RUBRO = [
    ("Automotor", "vtv",
     "VTV de PBA y CABA: el mes sale de la terminación de la patente (2 en febrero ... 0 en octubre y 1 en noviembre; "
     "diciembre y enero quedan para rezagados) y la VTV revisa que el matafuego tenga la carga vigente: se recarga en "
     "ese mes. Los exentos (PBA: menos de 2 años de patentado) recargan en el aniversario de la compra"),
    ("Edificios", "plano",
     "Sin mes fijo: la recarga vale 1 año desde que se hace y queda en la tarjeta (CABA, Ley 6116; PBA, Dec. 4992/90). "
     "La declaración jurada anual del administrador (CABA, Ley 941; PBA, Ley 14.701, hasta el 31/3) no fija el mes"),
    ("Educación", "escuelas",
     "Las escuelas estatales recargan en febrero, con los fondos de inicio del ciclo lectivo que la DGCyE gira a los "
     "Consejos Escolares antes del comienzo de clases (2 de marzo); las privadas, parejo"),
    ("Hoteles", "plano",
     "Sin mes fijo: el registro de alojamientos (CABA, Ley 4631) no fija un mes para incendio; rige la vigencia de 1 "
     "año de la recarga (Ley 6116) y, en PBA, el control de Bomberos cada 6 meses (Dec. 12/2005)"),
    ("Industrias / almacenes", "plano",
     "Sin mes fijo: el relevamiento de riesgos (RGRL, Res. SRT 463/09 y 529/09) se renueva con el contrato de la ART "
     "y Bomberos PBA controla cada 6 meses (Dec. 12/2005). Las paradas de enero (INDEC) no fijan el mes de la recarga"),
    ("Salud", "plano",
     "Sin mes fijo: funcionan todo el año; rigen la vigencia de 1 año de la recarga (Ley 6116, Dec. 4992/90) y el "
     "control de Bomberos cada 6 meses (Dec. 12/2005)"),
    ("Supermercados", "plano",
     "Sin mes fijo: constancia antisiniestral vigente con control de Bomberos cada 6 meses (Dec. 12/2005, PBA) y "
     "recarga anual desde su fecha (Ley 6116, Dec. 4992/90)"),
]
K_ABC = [0, 1, 2, 3, 13, 14, 15, 16]   # tipos ABC que fabrica la planta industrial (manuales y carros)
CIL_VACIO, TERMINADO, COMPRA = "Fábrica: cilindro vacío", "Fábrica: terminado", "Compra"
# columnas de la tabla de puestos: agentes E a L, cilindro nuevo M, tiempo N, u/h O, u/mes P, desglose Q, nota R
C_CIL, C_T, C_UH, C_UM = "M", "N", "O", "P"

def todos(v):
    return [v] * 8


def agentes(**kw):
    claves = ["abc", "bc", "d", "agua", "afff", "k", "co2", "hal"]
    return [kw.get(k, 0) for k in claves]


# Rutas por agente: fórmulas sobre los criterios de la misma columna ({c})
R_PH = f"={{c}}${F_PH}"
R_SECA = f"={{c}}${F_PH}+{{c}}${F_HUM}"
R_ABRE_CO2 = f"={{c}}${F_PH}+{{c}}${F_VALV}"
R_ABRE_HAL = f"={{c}}${F_PH}+{{c}}${F_USO}+{{c}}${F_VALV}"
R_CARGA_CO2 = f"={{c}}${F_PH}+{{c}}${F_VALV}+{{c}}${F_USO}"
R_RECUP = f"=(1-{{c}}${F_USO})*(1-{{c}}${F_CAMBIO})"
R_CAMBIA = f"=(1-{{c}}${F_USO})*{{c}}${F_CAMBIO}"
LIQS = dict(agua=R_CAMBIA, afff=R_CAMBIA, k=R_CAMBIA)

# (puesto, máquina / equipo, recurso, ruta por agente, tiempo tipo (fórmula de pasos), desglose, nota)
PUESTOS = [
    ("Recepción y registro", "Autoelevador + tablet e impresora de etiquetas; jaulas A y B", "Manual", todos(1),
     "=12.5+12.5+40+10", "Cargar jaulas 12,5 + descargar y controlar contra remito 12,5 + datos en el sistema 40 + "
     "separar por agente 10. Lo que se hace en el cliente no cuenta.", ""),
    ("Inspección visual", "Mesa con goma, lámpara con lupa, PC", "Manual", todos(1),
     "=5+30+25+10+15+5", "Escanear 5 + exterior 30 + válvula, manómetro y manguera 25 + fecha PH 10 + registro 15 + "
     "bandeja 5", "Antes 70 s; los pasos suman 90 s"),
    ("Descarga ABC — ruta recuperable", "Máquina neumática Yukon (tolva 240 kg)", "Máquina",
     agentes(abc=R_RECUP), "=15+60+10",
     "Escanear y agitado 15 + descargar mirando el visor 60 + tapar 10 (el retiro de válvula se cuenta en Desarme)", ""),
    ("Descarga ABC — ruta descarte", "Tambor Yukon M-000030", "Máquina", agentes(abc=R_CAMBIA), "=15+60+10",
     "Igual que la ruta recuperable, al tambor", ""),
    ("Descarga BC", "Tambor Yukon M-000030 dedicado BC", "Máquina", agentes(bc=R_CAMBIA), "=15+60+10",
     "Igual que ABC; todo a descarte", ""),
    ("Descarga de agua / AFFF / sales K", "IBC 1.000 l (AFFF) y tanque 600 l (agua)", "Máquina", agentes(**LIQS),
     "=10+10+15+60+10", "Escanear 10 + quitar lanza 10 + acoplar 15 + vaciar 60 + llevar a desarme 10",
     "Antes 180 s; los pasos suman 105 s. Sales K: misma ruta que agua y AFFF"),
    ("Recuperación de CO₂", "Trasvasadora CO₂ 103/D1 (misma máquina que la carga)", "Máquina",
     agentes(co2=R_ABRE_CO2), "=20+20+120+20+15+10",
     "Pesar 20 + conectar 20 + bombear 120 + verificar y pesar 20 + desconectar 15 + registrar 10", ""),
    ("Recuperación de halogenado", "Unidad THB/BP (misma máquina que la carga)", "Máquina",
     agentes(hal=R_ABRE_HAL), "=60+20+30+300+40+15",
     "Preparar y verificar cilindro de paso 60 + pesar 20 + conectar 30 + recuperar 300 + pesar ambos 40 + anotar 15",
     ""),
    ("Desarme y limpieza — recinto", "Lanza de soplado + aspiradora industrial + llaves", "Manual",
     agentes(abc=1, bc=1, d=1), f"=5+30+5+10+20+120*$E${F_PH}+5",
     "Verificar 5 + retirar válvula 30 + marbete 5 + bandeja 10 + soplar 20 + limpieza profunda 120 × % ABC con PH "
     "+ tapar 5", "Antes 90 s con 25 % de PH; los pasos suman 105 s con el mismo 25 %"),
    ("Desarme e inspección interior — zona general", "Boroscopio Klein ET19 + llaves", "Manual",
     agentes(co2=R_ABRE_CO2, hal=R_ABRE_HAL, agua=1, afff=1, k=1), "=10+40+60+90*0.4+15+15",
     "Confirmar vacío 10 + retirar válvula 40 + boroscopio 60 + enjuague 90 × 40 % + foto 15 + tapar 15", ""),
    ("Laboratorio — Puffer y mufla", "Puffer Yukon M-100150 + mufla ORL-0", "Máquina",
     agentes(abc=f"={{c}}${F_LAB}", bc=f"={{c}}${F_LAB}"), "=300+600+1800+60+120+300",
     "Tiempo por ENSAYO: registro 300 + mufla 600 + Puffer 1.800 + dictamen 60 + carga 120 + liberar tolva 300", ""),
    ("Prueba hidráulica", "Máquina PH Yukon 300 kg/cm²", "Máquina",
     todos(R_PH), "=12+20+90+90+75+15",
     "Verificar con patrón 12 + escanear 20 + llenar y conectar 90 + ensayo 90 + vaciar y retirar 75 + registrar 15. "
     "CO₂ lleva además ultrasonido y acuñado", ""),
    ("Secado", "Secadora Yukon de 6 recipientes", "Máquina",
     agentes(abc=R_SECA, bc=R_SECA, d=R_SECA, co2=R_PH, hal=R_PH), "=15+15*60/6+15",
     "Colocar 15 + 15 min ÷ 6 recipientes = 150 + retirar 15", "Agua, AFFF y sales K no se secan: se vuelven a llenar"),
    ("Inspección interior con videoscopio", "Videoscopio + mesa con soporte en V", "Manual",
     agentes(abc=R_SECA, bc=R_SECA, d=R_SECA, co2=R_PH, hal=R_PH), "=60+15+10", "Recorrer con la sonda 60 + guardar foto 15 + tapar 10", ""),
    ("Pintura", "Cabina con extracción, HVLP, lámparas IR, amoladora", "Máquina", todos(f"={{c}}${F_PINT}"), 750,
     "", "Sin desglose de pasos en el original"),
    ("Carga de polvo ABC", "Máquina neumática Yukon con corte por peso", "Máquina", agentes(abc=1),
     "=6+15+45+15+15+10", "Contraste por turno 6 + tara 15 + carga 45 + pesar 15 + rosca 15 + registrar 10",
     "Antes 90 s; los pasos suman 106 s"),
    ("Carga de polvo BC", "Cargadora portátil Yukon dedicada BC", "Máquina", agentes(bc=1), "=10+15+75+15+15+10",
     "Contraste 10 + tara 15 + carga portátil 75 + pesar 15 + rosca 15 + registrar 10", ""),
    ("Descarga y carga Clase D", "Recipiente exclusivo + embudo + balanza", "Manual", agentes(d=1),
     "=180+15+240+30+15+10", "Descarga manual 180 + tara 15 + carga con embudo 240 + ajuste 30 + rosca 15 + registro 10",
     "Supuesto del original (no está en el documento de procesos)"),
    ("Carga de agua / AFFF / sales K", "Red + tanque de premezcla con bomba 40 l/min + balanza", "Máquina",
     agentes(**LIQS), "=20+10+30+15+15+10",
     "Contraste y premezcla 20 + tara 10 + cargar 30 + corregir 15 + rosca 15 + registrar 10", ""),
    ("Carga de CO₂", "Trasvasadora CO₂ 103/D1 + manifold", "Máquina",
     agentes(co2=R_CARGA_CO2), "=30+10+15+150+25+10",
     "Tarar y programar 30 + elegir origen 10 + abrir manifold 15 + llenar 150 + pesar 25 + registrar 10", ""),
    ("Ensamblaje y marbete — recinto", "Morsa + llave dinamométrica", "Manual", agentes(abc=1, bc=1, d=1),
     "=10+30+20+10+30+20", "Escanear 10 + revisar piezas 30 + o-ring 20 + marbete 10 + enroscar 30 + manguera y "
     "repuestos 20", ""),
    ("Presurización — recinto", "Pinza Yukon M-000150 + regulador", "Máquina", agentes(abc=1, bc=1, d=1),
     "=3+15+10+20+15+10", "Contraste por turno 3 + fijar presión 15 + colocar 10 + presurizar 20 + pasador 15 + "
     "registrar 10", ""),
    ("Ensamblaje y presurización — zona general", "Morsa + dinamométrica + pinza Yukon", "Manual",
     agentes(co2=R_ABRE_CO2, hal=R_ABRE_HAL, agua=1, afff=1, k=1), "=120+50",
     "Ensamblaje 120 + presurización 50. CO₂ no se presuriza y halogenados se presurizan en la carga", ""),
    ("Carga y presurización de halogenado", "Unidad THB/BP + pinza", "Máquina", agentes(hal=R_ABRE_HAL),
     "=300+40+60+30", "Cargar 300 + presurizar 40 + agitar y repetir 60 + pesar y registrar 30", ""),
    ("Control de peso (CO₂ y halogenados sin recarga)", "Balanza", "Manual",
     agentes(co2="=1-(" + R_CARGA_CO2[1:] + ")", hal="=1-(" + R_ABRE_HAL[1:] + ")"), "=30+10",
     "Pesar y comparar 30 + registrar 10", "Anexo E: control anual de la masa"),
    ("Ensayo de pérdidas", "Tanque de inmersión inox (6 posiciones); spray en rodantes", "Máquina", todos(1),
     "=5+45+10+10+10", "Preparar por turno 5 + bajar y observar en tandas de 6: 45 + sacar 10 + secar 10 + registrar 10",
     ""),
    ("Precinto, etiqueta y despacho", "Dispensador + impresora + jaulas", "Manual", todos(1),
     "=15+25+15+10+15+6", "Precinto 15 + etiqueta 25 + revisión final 15 + cerrar OT 10 + jaula 15 + remito 6",
     "Antes 80 s; los pasos suman 86 s"),
    ("Inutilización y disposición final", "Taladro de columna con prensa en V", "Máquina",
     todos(f"={{c}}${F_INU}"), "=30+120+90+60+90",
     "Confirmar 30 + perforar 120 + desarmar y cortar manguera 90 + residuos 60 + foto y acta 90",
     "9.12 y Anexo G: acta de inutilización"),
    ("PH de mangueras (CO₂ y rodantes)", "Máquina PH Yukon 300 kg/cm² (misma que la PH de recipientes)", "Máquina",
     None, "=12+20+90+90+75+15",
     "Mismos pasos que la PH del recipiente: patrón 12 + escanear 20 + llenar y conectar 90 + ensayo 90 + vaciar y "
     "retirar 75 + registrar 15",
     "9.7.1.4: PH anual de mangueras de CO₂ y con cierre en el extremo (tabla 5). Demanda = fila Mangueras del Plan "
     "mensual. Criterio: tiempo igual al de la PH del recipiente; falta el rótulo plástico"),
]
# Cuándo pasa cada puesto (Word de recargas, tabla de operaciones; IRAM 3517-2:2020)
CUANDO = {
    "Recepción y registro": "Siempre (operación 1)",
    "Inspección visual": "Siempre (operación 2)",
    "Descarga ABC — ruta recuperable": "ABC que llega con presión y no cambia el polvo (operación 3). Sin presión pasa "
                                       "directo a desarme",
    "Descarga ABC — ruta descarte": "ABC con presión al que le toca el cambio de polvo por el dígito del año (9.9.4 a)",
    "Descarga BC": "BC con presión: el polvo BC se cambia siempre",
    "Descarga de agua / AFFF / sales K": "Con presión: el agente se cambia todos los años (Anexo E)",
    "Recuperación de CO₂": "CO₂ que va a PH o que hay que abrir por daño o pérdida de peso",
    "Recuperación de halogenado": "Halogenado que va a PH, que llega usado o que hay que abrir",
    "Desarme y limpieza — recinto": "Siempre en polvos (operación 4); la limpieza profunda, solo si va a PH",
    "Desarme e inspección interior — zona general": "CO₂ y halogenados cuando se abren; agua, AFFF y sales K siempre",
    "Laboratorio — Puffer y mufla": "Control periódico de la tolva, dudas en el visor y lotes de polvo nuevo sin Sello "
                                    "IRAM (operación 5; 9.9.3)",
    "Prueba hidráulica": "PH vencida o que vence antes del próximo mantenimiento, o daño menor (operación 6; 9.7.2, "
                         "Anexo E y 9.7.1.2)",
    "Secado": "Después de la PH, el mismo día; en polvos también si hay humedad al desarmar (operación 7)",
    "Inspección interior con videoscopio": "Con el secado (operación 7)",
    "Pintura": "Óxido, metal a la vista, pérdida de brillo o color, o repintado parcial (operación 8; 9.10)",
    "Carga de polvo ABC": "Siempre (operación 9)",
    "Carga de polvo BC": "Siempre",
    "Descarga y carga Clase D": "Siempre: carga manual con embudo para no contaminar el agente",
    "Carga de agua / AFFF / sales K": "Siempre: el agente se cambia cada año (Anexo E)",
    "Carga de CO₂": "CO₂ que se abrió o que llega usado",
    "Ensamblaje y marbete — recinto": "Siempre (operación 10)",
    "Presurización — recinto": "Siempre (operación 11)",
    "Ensamblaje y presurización — zona general": "CO₂ y halogenados que se abrieron; agua, AFFF y sales K siempre",
    "Carga y presurización de halogenado": "Halogenado que se abrió o que llega usado",
    "Control de peso (CO₂ y halogenados sin recarga)": "CO₂ y halogenados que no se abren: control anual de la masa "
                                                       "(Anexo E)",
    "Ensayo de pérdidas": "Siempre (operación 12)",
    "Precinto, etiqueta y despacho": "Siempre (operación 13)",
    "Inutilización y disposición final": "Vida útil vencida, no pasa la PH o roto (operación 14; 9.12)",
    "PH de mangueras (CO₂ y rodantes)": "Mangueras de CO₂ y con cierre en el extremo, todos los años (9.7.1.4)",
}
# Ruta del cilindro nuevo (sustituto ABC que llega vacío de la planta industrial con PH, secado y pintura de fábrica):
# se registra, se carga, se arma, se presuriza, se prueba y se etiqueta (Word de recargas, puestos 1, 9 a 13)
CILINDRO_NUEVO = {"Recepción y registro", "Carga de polvo ABC", "Ensamblaje y marbete — recinto",
                  "Presurización — recinto", "Ensayo de pérdidas", "Precinto, etiqueta y despacho"}

# ------------------------------------------------------------------ estilo (sólo negro y grises)
F_N = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
F_T = Font(name="Arial", size=12, bold=True)
F_W = Font(name="Arial", size=10, bold=True, color="FFFFFF")
NEGRO = PatternFill("solid", fgColor="000000")
GRIS = PatternFill("solid", fgColor="D9D9D9")
GRIS_C = PatternFill("solid", fgColor="F2F2F2")
GRIS_O = PatternFill("solid", fgColor="A6A6A6")
fino = Side(style="thin", color="808080")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)
DER = Alignment(horizontal="right", vertical="center")


def hoja(wb, nombre, anchos):
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False
    for k, a in enumerate(anchos):
        ws.column_dimensions[L(k + 1)].width = a
    return ws


def tit(ws, r, t, c=1):
    ws.cell(r, c, t).font = F_T


def nota(ws, r, t, c=1):
    x = ws.cell(r, c, t)
    x.font = Font(name="Arial", size=9, italic=True)
    x.alignment = Alignment(horizontal="left", vertical="center")


def cab(ws, r, textos, c0=1):
    for k, t in enumerate(textos):
        c = ws.cell(r, c0 + k, t)
        c.font, c.fill, c.alignment, c.border = F_W, NEGRO, CEN, BORDE


def celda(ws, r, c, v, fmt=None, dato=False, neg=False, res=False):
    x = ws.cell(r, c, v)
    x.font = F_B if (neg or res) else F_N
    x.border = BORDE
    if fmt:
        x.number_format = fmt
    x.alignment = IZQ if isinstance(v, str) and not v.startswith("=") else (CEN if fmt is None else DER)
    return x


# ------------------------------------------------------------------ hojas
def importrange(ws, r0, c0, link, rango, valores, fmt):
    """Bloque IMPORTRANGE como lo exporta Google: la primera celda lleva la fórmula (el link sale de la hoja Links);
    el resto, el valor calculado."""
    for i, fila in enumerate(valores):
        for j, x in enumerate(fila):
            if i == 0 and j == 0:
                f = f'=IFERROR(__xludf.DUMMYFUNCTION("IMPORTRANGE({link},""{rango}"")"),{x!r})'
            else:
                f = f'=IFERROR(__xludf.DUMMYFUNCTION("""COMPUTED_VALUE"""),{x!r})'
            celda(ws, r0 + i, c0 + j, f, fmt)


def h_links(wb, filas):
    ws = hoja(wb, "Links", [30, 70, 96])
    tit(ws, 1, "Links de los IMPORTRANGE")
    nota(ws, 2, "Pegar en la columna B el link del Google Sheet de cada archivo; todas las fórmulas IMPORTRANGE leen el "
                "link de esta hoja. Hasta que se pegue, esas celdas muestran el último valor conocido o un error de referencia.")
    for k, t in enumerate(["Archivo", "Link", "Qué se importa"]):
        c = ws.cell(3, 1 + k, t)
        c.font, c.fill, c.alignment, c.border = F_B, GRIS_C, CEN, BORDE
    for i, (a, u, q) in enumerate(filas):
        celda(ws, 4 + i, 1, a, neg=True)
        celda(ws, 4 + i, 2, u, dato=True)
        celda(ws, 4 + i, 3, q)
    return ws


def h_proy(wb, planta):
    ws = hoja(wb, "PROY_Indexado", [34, 13, 13, 13, 13, 13, 13, 13, 13, 13, 13, 50])
    tit(ws, 1, "Proyecciones — Indexado Dim Técnico")
    cab(ws, 3, ["Recargas por tipo"] + ANIOS)
    for i, (n, v) in enumerate(REC):
        celda(ws, 4 + i, 1, n)
    importrange(ws, 4, 2, LK_PROY, IMP_REC, [v for n, v in REC], "#,##0")
    celda(ws, 21, 1, "Recargas total", neg=True)
    for k in range(10):
        celda(ws, 21, 2 + k, f"=SUM({L(2 + k)}4:{L(2 + k)}20)", "#,##0", res=True)

    tit(ws, 23, "Recargas por zona y planta (PBA Interior queda fuera: la atienden distribuidores ajenos)")
    cab(ws, 24, ["Zona", "% zonal 2035", "Planta"])
    for i, (z, p, pl) in enumerate(ZONAS):
        celda(ws, 25 + i, 1, z)
        celda(ws, 25 + i, 3, pl).alignment = CEN
    importrange(ws, 25, 2, LK_PROY, IMP_ZON, [[p] for z, p, pl in ZONAS], "0.0%")
    dv = DataValidation(type="list", formula1='"Comercial,Industrial,Fuera"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C25:C38")
    cab(ws, 39, ["Planta", "% de recargas", "Control: suma de sus zonas"])
    celda(ws, 40, 1, "Comercial (Proyecciones V61)", neg=True)
    importrange(ws, 40, 2, LK_PROY, "Indexado Dim Técnico!V61", [[PCT_PLANTA["Comercial"]]], "0.0%")
    celda(ws, 40, 3, '=SUMIF($C$25:$C$38,"Comercial",$B$25:$B$38)', "0.0%")
    celda(ws, 41, 1, "Industrial (Proyecciones V63)", neg=True)
    importrange(ws, 41, 2, LK_PROY, "Indexado Dim Técnico!V63", [[PCT_PLANTA["Industrial"]]], "0.0%")
    celda(ws, 41, 3, '=SUMIF($C$25:$C$38,"Industrial",$B$25:$B$38)', "0.0%")
    celda(ws, 42, 1, "Fuera: PBA Interior (distribuidores ajenos)")
    celda(ws, 42, 2, '=SUMIF($C$25:$C$38,"Fuera",$B$25:$B$38)', "0.0%")
    celda(ws, 43, 1, "Total (tiene que dar 100 %)")
    celda(ws, 43, 2, "=SUM(B40:B42)", "0.0%")

    tit(ws, 44, "Tasas de Proyecciones (hojas PROYECCION por rubro; iguales para todos los tipos)")
    cab(ws, 45, ["Concepto", "Valor", "Uso en este archivo"])
    ws.merge_cells("C45:L45")
    for i, (t, v, u) in enumerate(TASAS):
        celda(ws, 46 + i, 1, t)
        celda(ws, 46 + i, 2, v, "0.0%")
        celda(ws, 46 + i, 3, u)
        ws.merge_cells(start_row=46 + i, start_column=3, end_row=46 + i, end_column=12)

    celda(ws, 51, 1, "Planta de este archivo", neg=True)
    celda(ws, 51, 2, planta).alignment = CEN
    celda(ws, 52, 1, "% de recargas de esta planta", neg=True)
    celda(ws, 52, 2, '=IF(B51="Comercial",B40,B41)', "0.0%", res=True)

    tit(ws, P_MIX - 3, f"Recargas de mercado 2035 por tipo y rubro en las zonas de la planta {planta.lower()} "
                       f"(hojas PROYECCION <rubro>)")
    cab(ws, P_MIX - 1, ["Tipo"] + RUBROS + ["Total"])
    for k in range(17):
        r = P_MIX + k
        celda(ws, r, 1, f"=A{4 + k}")
        for c in range(7):
            celda(ws, r, 2 + c, MIX_RUBRO[planta][c][k], "#,##0")
        celda(ws, r, 9, f"=SUM(B{r}:H{r})", "#,##0", res=True)
    nota(ws, P_MIX + 17, "Carros repartidos 25/55/12/8 % entre 25, 50, 70 y 100 kg, como PxQ (Corregido) filas 66 a 69. "
                         "Comercial: CABA y PBA sin zona Sur ni Interior; industrial: CABA Sur y PBA Sur.")
    ws.freeze_panes = "B4"


def h_datos(wb):
    ws = hoja(wb, "datos_base", [24, 11, 11, 11, 11, 11, 11, 11, 11, 11, 11])
    cab(ws, 1, ["Tipo de matafuego"] + [f"=PROY_Indexado!{L(2 + k)}3" for k in range(10)])
    for i in range(17):
        celda(ws, 2 + i, 1, f"=PROY_Indexado!A{4 + i}")
        for k in range(10):
            celda(ws, 2 + i, 2 + k, f"=PROY_Indexado!{L(2 + k)}{4 + i}*PROY_Indexado!$B$52", "#,##0")
    celda(ws, 19, 1, "Recargas Total", neg=True)
    for k in range(10):
        celda(ws, 19, 2 + k, f"=SUM({L(2 + k)}2:{L(2 + k)}18)", "#,##0", res=True)
    ws.freeze_panes = "B2"


def h_param(wb, planta):
    ws = hoja(wb, "parametros_recargas",
              [52, 30, 34, 11, 10, 10, 10, 10, 10, 10, 10, 10, 11, 11, 10, 11, 70, 60, 70])
    cab(ws, 1, ["Parámetro", "Valor", "Fuente / criterio"])
    ws.merge_cells("C1:Q1")
    for r, t, v, f, d, fu in P_GEN:
        celda(ws, r, 1, t)
        celda(ws, r, 2, v, f, dato=d, res=not d)
        celda(ws, r, 3, fu)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=17)
    # criterios de recarga: una fila por cosa que le pasa al equipo, una columna por agente
    tit(ws, 11, f"CRITERIOS DE RECARGA POR AGENTE — PLANTA {planta.upper()} (qué le pasa al equipo en el año)")
    cab(ws, 12, ["Qué le pasa al equipo", "Unidad", "Fuente", "Estado"] + AGENTES + ["Explicación"])
    ws.merge_cells("M12:S12")
    fmts = {"años": "0", "% por año": "0.0%", "ensayos/día": "0", "ensayos/u.": "0.0000"}
    for r, t, u, fu, est, vals, ex in CRITERIOS:
        calc = est == "Cálculo"
        celda(ws, r, 1, t, neg=calc)
        celda(ws, r, 2, u).alignment = CEN
        celda(ws, r, 3, fu)
        celda(ws, r, 4, est).alignment = CEN
        for j, v in enumerate(vals):
            x = celda(ws, r, 5 + j, v.replace("{c}", L(5 + j)) if isinstance(v, str) else v, fmts[u], res=calc)
            if est == "A confirmar":
                x.fill = GRIS_C
        celda(ws, r, 13, ex)
        ws.merge_cells(start_row=r, start_column=13, end_row=r, end_column=19)
        ws.row_dimensions[r].height = 30
    nota(ws, 29, "Estado: Norma = lo fija IRAM 3517-2; Dato = Proyecciones; A confirmar (fondo gris) = criterio del "
                 "grupo que falta confirmar en la visita al recargador; Cálculo = sale de las filas de arriba.")
    nota(ws, 30, "Cada puesto de CAPACIDAD POR PUESTO toma estas filas en la columna de su agente (ver la columna "
                 "Cuándo pasa). Cambiar un valor acá cambia la ruta de todos los puestos que lo usan.")

    # estacionalidad de las recargas por rubro
    tit(ws, FILA_RUB - 3, "ESTACIONALIDAD DE LAS RECARGAS POR RUBRO (% de las recargas del año en cada mes)")
    nota(ws, FILA_RUB - 2, "IRAM 3517-2:2020 9.2.1 (mantenimiento al menos anual) y 9.4.14 (la tarjeta fija el mes y "
                           "año del próximo): cada equipo vuelve en el mes de su último mantenimiento. Solo hay "
                           "estacionalidad donde una regla o un calendario fija el mes para todo el rubro.")
    cab(ws, FILA_RUB - 1, ["Rubro", "Fuente / criterio", "", ""] + MESES + ["Total"])
    ws.merge_cells(start_row=FILA_RUB - 1, start_column=2, end_row=FILA_RUB - 1, end_column=4)
    for c, (rb, regla, fu) in enumerate(EST_RUBRO):
        r = FILA_RUB + c
        celda(ws, r, 1, rb, neg=True)
        celda(ws, r, 2, fu)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 54
        ex, pat = f"$B${F_EXENTO}", f"$Q${F_PAT}"
        for m in range(12):
            x = L(5 + m)
            if regla == "plano":
                v = "=1/12"
            elif regla == "escuelas":
                v = f"=$B${F_ESTATAL}+(1-$B${F_ESTATAL})/12" if m == 1 else f"=(1-$B${F_ESTATAL})/12"
            else:
                v = (f"=(1-{ex})/10+{ex}*{x}${F_PAT}/{pat}" if 1 <= m <= 10 else f"={ex}*{x}${F_PAT}/{pat}")
            celda(ws, r, 5 + m, v, "0.0%")
        celda(ws, r, 17, f"=SUM(E{r}:P{r})", "0%", res=True)
    cab(ws, F_ESTATAL - 1, ["Dato de la estacionalidad", "Valor", "Fuente", ""] + MESES + ["Total"])
    ws.merge_cells(start_row=F_ESTATAL - 1, start_column=3, end_row=F_ESTATAL - 1, end_column=4)
    datos = [(F_ESTATAL, "Escuelas de gestión estatal (% de las unidades educativas)", ESTATAL, "0.0%", True,
              "PBA: 68,2 % estatal y 31,8 % privada (HCD PBA, proyecto 13-14 D-1382, con el Relevamiento Anual)"),
             (F_PAT, "Patentamientos 2025 por mes (u.)", None, "#,##0", True, "ACARA, total del mercado"),
             (F_P24, "Patentamientos 2024 (u.)", PAT_2024, "#,##0", True, "ACARA"),
             (F_FLOTA, "Flota circulante 2025 (u.)", FLOTA_2025, "#,##0", True, "AFAC, Flota Vehicular Circulante 2025"),
             (F_EXENTO, "Exentos de VTV (% de la flota)", f"=(B{F_P24}+Q{F_PAT})/B{F_FLOTA}", "0.0%", False,
              "PBA: VTV anual desde los 2 años de patentado → patentados 2024 + 2025 ÷ flota")]
    for r, t, v, fmt, dato, fu in datos:
        celda(ws, r, 1, t)
        if v is not None:
            celda(ws, r, 2, v, fmt, dato=dato, res=not dato)
        celda(ws, r, 3, fu)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    for m in range(12):
        celda(ws, F_PAT, 5 + m, PAT_2025[m], "#,##0", dato=True)
    celda(ws, F_PAT, 17, f"=SUM(E{F_PAT}:P{F_PAT})", "#,##0", res=True)
    nota(ws, F_EXENTO + 1, "Fuentes: VTV PBA (vtv.gba.gob.ar, plazos; anexo del cronograma en normas.gba.gob.ar) y CABA "
                           "(buenosaires.gob.ar); revisión del matafuego en la VTV (iProfesional); 2 años de exención en "
                           "PBA (Ámbito, oct. 2025); ACARA (El Economista, MDZ); AFAC 2025. DGCyE: fondos de inicio del "
                           "ciclo lectivo a los Consejos Escolares (Ámbito 8/2/2022; 2026); calendario escolar PBA 2026.")
    nota(ws, F_EXENTO + 2, "Edificios y comercios: CABA Leyes 941, 4631 y 6116; PBA Ley 14.701 y Decretos 4992/90 y "
                           "12/2005; Res. SRT 463/09 y 529/09. Ninguna fija un mes del calendario: la recarga vence a "
                           "los 12 meses de hecha (IRAM 3517-2 9.2.1 y 9.4.14), así que esos rubros quedan parejos.")
    nota(ws, F_EXENTO + 3, "Pico de verano en autos: lo informan recargadores de Córdoba, Mendoza, Paraná y Rafaela "
                           "(prensa local; Córdoba: 12 % de las ITV del año en enero). En el AMBA no hay dato: a "
                           "confirmar en la visita al recargador.")

    # tipos con su perfil mensual
    tit(ws, FILA_TIPO - 2, "TIPOS DE MATAFUEGO Y PERFIL MENSUAL DE RECARGAS (mix de rubros de la planta × estacionalidad "
                           "de cada rubro)")
    cab(ws, FILA_TIPO - 1, ["Tipo", "Agente (ruta)", "u. equivalentes por recipiente", "Manguera con PH anual (1 = sí)"]
        + MESES + ["Mes pico", "Fuente de la manguera"])
    for i, (ag, rod, mang, fu) in enumerate(TIPOS):
        r = FILA_TIPO + i
        mix = f"PROY_Indexado!$B${P_MIX + i}:$H${P_MIX + i}"
        celda(ws, r, 1, f"=datos_base!A{2 + i}")
        celda(ws, r, 2, ag).alignment = CEN
        celda(ws, r, 3, "=$B$9" if rod else 1, "0.0" if rod else "0", res=bool(rod), dato=not rod)
        celda(ws, r, 4, mang, "0", dato=True)
        for m in range(12):
            x = L(5 + m)
            suma = "+".join(f"PROY_Indexado!${L(2 + c)}${P_MIX + i}*{x}${FILA_RUB + c}" for c in range(7))
            celda(ws, r, 5 + m, f"=IF(SUM({mix})=0,1/12,({suma})/SUM({mix}))", "0.0%")
        celda(ws, r, 17, f"=MAX(E{r}:P{r})", "0.0%", res=True)
        celda(ws, r, 18, fu)
    r = FILA_TIPO + 17
    celda(ws, r, 1, "Toda la planta (recargas 2035; lo importan el MRP y Recargas.xlsx)", neg=True)
    for m in range(12):
        x = L(5 + m)
        celda(ws, r, 5 + m, f"=SUMPRODUCT(datos_base!$K$2:$K$18,{x}{FILA_TIPO}:{x}{FILA_TIPO + 16})"
                            f"/SUM(datos_base!$K$2:$K$18)", "0.0%", res=True)
    celda(ws, r, 17, f"=MAX(E{r}:P{r})", "0.0%", res=True)
    dv = DataValidation(type="list", formula1=f"$E${FILA_PUESTO - 1}:$L${FILA_PUESTO - 1}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"B{FILA_TIPO}:B{FILA_TIPO + 16}")

    # puestos
    tit(ws, FILA_PUESTO - 2, "CAPACIDAD POR PUESTO (fracción de cada agente que pasa por el puesto; cilindro nuevo: 1 = pasa)")
    cab(ws, FILA_PUESTO - 1, ["N°", "Puesto / operación", "Máquina / equipo", "Recurso"] + AGENTES
        + ["Cilindro nuevo (alta de sustituto)", "Tiempo tipo (s/u)", "Cap. real (u/h)", "Cap. real (u/mes, 1 turno)",
           "Desglose del tiempo (s)", "Nota", "Cuándo pasa (Word de recargas, tabla de operaciones, y norma)"])
    for k, (pu, maq, rec, ruta, t, des, no) in enumerate(PUESTOS):
        r = FILA_PUESTO + k
        celda(ws, r, 1, k + 1, "0").alignment = CEN
        celda(ws, r, 2, pu)
        celda(ws, r, 3, maq)
        celda(ws, r, 4, rec, dato=True).alignment = CEN
        for j in range(8):
            if ruta is None:
                celda(ws, r, 5 + j, "—").alignment = CEN
            else:
                v = ruta[j].replace("{c}", L(5 + j)) if isinstance(ruta[j], str) else ruta[j]
                celda(ws, r, 5 + j, v, "0.0%")
        celda(ws, r, 13, 1 if pu in CILINDRO_NUEVO else 0, "0", dato=True)
        celda(ws, r, 14, t, "0", dato=True)
        celda(ws, r, 15, f"=IF({C_T}{r}>0,3600/{C_T}{r},0)", "0.0", res=True)
        celda(ws, r, 16, f"={C_UH}{r}*$B$6", "#,##0", res=True)
        celda(ws, r, 17, des)
        celda(ws, r, 18, no)
        celda(ws, r, 19, CUANDO[pu])
    dv = DataValidation(type="list", formula1='"Máquina,Manual,Tercerizado"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"D{FILA_PUESTO}:D{FILA_PUESTO + N_PUESTOS - 1}")

    # perfil de llegada de los sustitutos ABC nuevos (perfiles de producción del MRP)
    tit(ws, FILA_PCIL - 3, "PERFIL DE LLEGADA DE LOS SUSTITUTOS ABC NUEVOS (perfiles de producción del MRP)")
    nota(ws, FILA_PCIL - 2, "Solo pasan por la ruta del cilindro nuevo los que llegan como cilindro vacío (planta "
                            "comercial). IMPORTRANGE de la hoja ESTACIONALIDAD del MRP (link en la hoja Links).")
    cab(ws, FILA_PCIL - 1, ["Perfil del MRP", "Tipos ABC que lo siguen", "", ""] + MESES + ["Total"])
    ws.merge_cells(start_row=FILA_PCIL - 1, start_column=2, end_row=FILA_PCIL - 1, end_column=4)
    for f, (nom, rng, vals) in enumerate(E_MRP):
        r = FILA_PCIL + f
        celda(ws, r, 1, nom)
        celda(ws, r, 2, ", ".join(REC[k][0] for k, ff in FAM_ABC.items() if ff == f))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        importrange(ws, r, 5, LK_MRP, rng, [vals], "0.0%")
        celda(ws, r, 17, f"=SUM(E{r}:P{r})", "0%", res=True)

    # equivalencia de un rodante con los tiempos de fábrica de FLAMA
    tit(ws, FILA_EQ - 3, "EQUIVALENCIA DE UN RODANTE (tiempos de fábrica de FLAMA: Manuales y Carros)")
    nota(ws, FILA_EQ - 2, "Cuántos manuales equivale un carro en las operaciones que tienen la fábrica y la recarga: "
                          "tiempo del carro (ponderado con el mix 2035) ÷ tiempo del manual. IMPORTRANGE: los links van "
                          "en la hoja Links.")
    cab(ws, FILA_EQ - 1, ["Operación común", "Manual (s/u)", "Carro (min/u)", "Carro ÷ manual", "Origen"])
    ws.merge_cells(start_row=FILA_EQ - 1, start_column=5, end_row=FILA_EQ - 1, end_column=18)
    for k, (op, man, (cc, cv)) in enumerate(EQUIV):
        r = FILA_EQ + k
        celda(ws, r, 1, op)
        fm = "+".join(f'IMPORTRANGE({LK_MAN},""parametros_manuales!{c}"")' for c, _ in man)
        celda(ws, r, 2, f'=IFERROR(__xludf.DUMMYFUNCTION("{fm}"),{sum(v for _, v in man)!r})', "0")
        celda(ws, r, 3, f'=IFERROR(__xludf.DUMMYFUNCTION("IMPORTRANGE({LK_CAR},""parametros_carros!{cc}"")"),'
                        f'{cv!r})', "0.00")
        celda(ws, r, 4, f"=IF(B{r}=0,0,C{r}*60/B{r})", "0.0", res=True)
        celda(ws, r, 5, "Manuales, parametros_manuales " + " + ".join(c for c, _ in man) + " (tiempo tipo); Carros, "
                        f"parametros_carros {cc} (tiempo ponderado)")
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=18)
    r = FILA_EQ + len(EQUIV)
    celda(ws, r, 1, "Equivalencia de un rodante (promedio)", neg=True)
    celda(ws, r, 4, f"=AVERAGE(D{FILA_EQ}:D{r - 1})", "0.0", res=True)
    celda(ws, r, 5, "Va a B9 y de ahí a la columna u. equivalentes de los rodantes (AFFF 50 l y ABC 25 a 100 kg). "
                    "El Dimensionamiento Técnico V2 usaba 3 (supuesto)")
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=18)
    ws.freeze_panes = "B2"


def bloque_plan(r0):
    """Filas del bloque de un año en Plan mensual: tipos, total, agentes, mangueras y cilindros nuevos."""
    return {"t1": r0 + 1, "t17": r0 + 17, "tot": r0 + 18, "ag1": r0 + 20, "mang": r0 + 28, "cil": r0 + 29}


def h_plan(wb):
    ws = hoja(wb, "Plan mensual", [40, 12, 9] + [9] * 12 + [10])
    pr = "parametros_recargas"
    for i, a in enumerate(ANIOS):
        r0 = 1 + 32 * i
        b = bloque_plan(r0)
        cab(ws, r0, [a, "Agente", "u. equiv."] + MESES + ["Total"])
        ws.cell(r0, 1).number_format = "0"
        y = L(2 + i)
        for k in range(17):
            r = b["t1"] + k
            # mantenimiento en planta de los sustitutos: el stock del año anterior (9.4.5 y 9.2.1), 1/12 por mes
            ms = f"+{SS}!${L(1 + i)}${S_STOCK + k}/12" if i else ""
            celda(ws, r, 1, f"={pr}!A{FILA_TIPO + k}")
            celda(ws, r, 2, f"={pr}!B{FILA_TIPO + k}").alignment = CEN
            celda(ws, r, 3, f"={pr}!C{FILA_TIPO + k}", "0.0")
            for m in range(12):
                celda(ws, r, 4 + m, f"=datos_base!${y}${2 + k}*{pr}!{L(5 + m)}${FILA_TIPO + k}{ms}", "#,##0.0")
            celda(ws, r, 16, f"=SUM(D{r}:O{r})", "#,##0", res=True)
        celda(ws, b["tot"], 1, "Total recipientes (con mantenimiento de sustitutos)", neg=True)
        for m in range(13):
            c = L(4 + m)
            celda(ws, b["tot"], 4 + m, f"=SUM({c}{b['t1']}:{c}{b['t17']})", "#,##0", res=True)
        celda(ws, b["tot"] + 1, 1, "Por agente (u. equivalentes)", neg=True)
        for j, ag in enumerate(AGENTES):
            r = b["ag1"] + j
            celda(ws, r, 1, ag)
            for m in range(13):
                c = L(4 + m)
                celda(ws, r, 4 + m, f"=SUMPRODUCT(($B${b['t1']}:$B${b['t17']}=$A{r})*{c}${b['t1']}:{c}${b['t17']}"
                                    f"*$C${b['t1']}:$C${b['t17']})", "#,##0.0")
        r = b["mang"]
        celda(ws, r, 1, "Mangueras con PH anual (u.)", neg=True)
        for m in range(13):
            c = L(4 + m)
            celda(ws, r, 4 + m, f"=SUMPRODUCT({c}${b['t1']}:{c}${b['t17']}*{pr}!$D${FILA_TIPO}:"
                                f"$D${FILA_TIPO + 16})", "#,##0.0", res=True)
        r = b["cil"]
        celda(ws, r, 1, "Sustitutos ABC nuevos como cilindro vacío (u. equiv.)", neg=True)
        for m in range(12):
            partes = []
            for f in range(len(E_MRP)):
                ks = [k for k, ff in FAM_ABC.items() if ff == f]
                a1, a2 = S_ALTA + ks[0], S_ALTA + ks[-1]
                partes.append(f"SUMPRODUCT({SS}!${y}${a1}:${y}${a2},--({SS}!$F${S_TIPO + ks[0]}:$F${S_TIPO + ks[-1]}"
                              f'="{CIL_VACIO}"),{pr}!$C${FILA_TIPO + ks[0]}:$C${FILA_TIPO + ks[-1]})'
                              f"*{pr}!{L(5 + m)}${FILA_PCIL + f}")
            celda(ws, r, 4 + m, "=" + "+".join(partes), "#,##0.0", res=True)
        celda(ws, r, 16, f"=SUM(D{r}:O{r})", "#,##0", res=True)
    ws.freeze_panes = "D2"


def fila_dem(i, k):
    return 2 + 32 * i + k


def h_demanda(wb):
    ws = hoja(wb, "Demanda por puesto", [44] + [9] * 12 + [11, 13, 12])
    for i, a in enumerate(ANIOS):
        r0 = 1 + 32 * i
        b = bloque_plan(r0)
        cab(ws, r0, [f"Año {a}"] + MESES + ["Total anual", "Demanda mes pico (u/mes)", "Demanda (u/h, 1 turno)"])
        for k in range(N_PUESTOS):
            r = fila_dem(i, k)
            rp = FILA_PUESTO + k
            celda(ws, r, 1, f"=parametros_recargas!$B${rp}")
            for m in range(12):
                c = L(4 + m)
                if PUESTOS[k][3] is None:
                    f = f"='Plan mensual'!{c}{b['mang']}"
                else:
                    f = ("=" + "+".join(f"parametros_recargas!${L(5 + j)}${rp}*'Plan mensual'!{c}{b['ag1'] + j}"
                                        for j in range(8))
                         + f"+parametros_recargas!${C_CIL}${rp}*'Plan mensual'!{c}{b['cil']}")
                celda(ws, r, 2 + m, f, "#,##0.0")
            celda(ws, r, 14, f"=SUM(B{r}:M{r})", "#,##0", res=True)
            celda(ws, r, 15, f"=MAX(B{r}:M{r})", "#,##0.0", res=True)
            celda(ws, r, 16, f"=O{r}/parametros_recargas!$B$6", "0.00", res=True)
    ws.freeze_panes = "B2"


ACCIONES = ["OK", "Agrega turno", "Agrega máquina", "Agrega operario (puesto manual)", "Tercerizado", "Revisar"]
ACC_TXT = ["Aprovechamiento dentro del umbral",
           "Solo en el año del 2° turno, a todos los puestos que lo necesitan hasta el final del horizonte",
           "Antes del año del 2° turno, o en ese año si con 2 turnos no alcanza hasta el final; después, nunca",
           "Lo mismo que máquina, en un puesto manual",
           "Recurso cargado como Tercerizado en parametros",
           "Se cargó a mano un plan que pasa el umbral"]


# colores de la acción aplicada en el aprovechamiento (los mismos de Dimensionamiento Manuales)
COLOR_ACC = {"Agrega turno": "4A86E8", "Agrega máquina": "FF00FF", "Agrega operario (puesto manual)": "E06666",
             "Tercerizado": "FF9900", "Revisar": "FFFF00"}


def plan_capacidad(dem, cap, recurso, umbral=0.9, t_max=2, anio_2t=ANIO_2T, piso=None):
    """Máquinas y turnos por puesto y año. Nunca se baja y se parte del piso (1 × 1 o lo cargado a mano; 0 máquinas
    es un año tercerizado). Antes de anio_2t, si el puesto pasa el umbral se agrega máquina u operario. En anio_2t
    se dimensiona hasta el final del horizonte: los puestos que lo necesitan (o que ya tenían turno cargado) pasan a
    t_max turnos ese mismo año y se agregan las máquinas que hagan falta para que ningún año pase el umbral. Después
    no se agrega nada: todos los turnos entran el mismo año y no se compra máquina al final."""
    iy = ANIOS.index(anio_2t) if anio_2t in ANIOS else len(ANIOS)
    plan = {}
    for k in range(len(cap)):
        m, t = 1, 1
        activo = cap[k] > 0 and recurso[k] != "Tercerizado"
        for i in range(len(ANIOS)):
            pm, pt = piso[(k, i)] if piso else (1, 1)
            if pm == 0:
                plan[(k, i)] = (0, max(1, pt))
                continue
            if i < iy:
                m = max(m, pm)
                while activo and dem[k][i] / (cap[k] * m * t) > umbral:
                    m += 1
            elif i == iy:
                resto = range(iy, len(ANIOS))
                p_m = max(piso[(k, j)][0] for j in resto) if piso else 1
                p_t = max(piso[(k, j)][1] for j in resto) if piso else 1
                m = max(m, p_m)
                n = max(dem[k][j] for j in resto) / (cap[k] * umbral) if activo else 0
                if p_t > 1 or n > m * t + 1e-9:
                    t = t_max
                m = max(m, math.ceil(n / t - 1e-9))
            plan[(k, i)] = (m, t)
    return plan


def colores_aprovechamiento(ws, r_util, r_maq, r_tur, n, col_rec, umbral, sin_maquina_es_terc=False):
    """Formato condicional del aprovechamiento: el color sale de comparar las tablas de máquinas y turnos con el año
    anterior (2026 contra 1 × 1)."""
    for rango, c, prev in ((f"B{r_util}:B{r_util + n - 1}", "B", None),
                           (f"C{r_util}:K{r_util + n - 1}", "C", "B")):
        pm = f"{prev}{r_maq}" if prev else "1"
        pt = f"{prev}{r_tur}" if prev else "1"
        terc = f"{c}{r_maq}=0" if sin_maquina_es_terc else f'${col_rec}{r_util}="Tercerizado"'
        reglas = [("Revisar", f"{c}{r_util}>{umbral}"),
                  ("Tercerizado", terc),
                  ("Agrega turno", f"{c}{r_tur}>{pt}"),
                  ("Agrega máquina", f'AND({c}{r_maq}>{pm},${col_rec}{r_util}<>"Manual")'),
                  ("Agrega operario (puesto manual)", f'AND({c}{r_maq}>{pm},${col_rec}{r_util}="Manual")')]
        for acc, f in reglas:
            ws.conditional_formatting.add(rango, FormulaRule(
                formula=[f], stopIfTrue=True, fill=PatternFill("solid", fgColor=COLOR_ACC[acc], bgColor=COLOR_ACC[acc])))


def h_resumen(wb, plan):
    ws = hoja(wb, "Resumen Capacidad", [48] + [10] * 10 + [13, 3, 30, 60])
    bloques = {"util": 1, "maq": 33, "tur": 65}
    # colores de la acción aplicada (formato condicional del aprovechamiento)
    cab(ws, 1, ["Color", "Cuándo"], 14)
    for k, (a, t) in enumerate(zip(ACCIONES, ACC_TXT)):
        x = celda(ws, 2 + k, 14, a)
        if a in COLOR_ACC:
            x.fill = PatternFill("solid", fgColor=COLOR_ACC[a])
        celda(ws, 2 + k, 15, t)
    celda(ws, 9, 14, "Umbral de aprovechamiento", neg=True)
    celda(ws, 9, 15, "=parametros_recargas!B7", "0%", res=True)
    celda(ws, 10, 14, "Turnos máximos por día", neg=True)
    celda(ws, 10, 15, "=parametros_recargas!B8", "0", res=True)
    celda(ws, 11, 14, "Año de inicio del 2° turno", neg=True)
    celda(ws, 11, 15, "=parametros_recargas!B10", "0", res=True)
    nota(ws, 13, "Antes del año del 2° turno se agrega máquina u operario; en ese año entran todos los turnos y las "
                 "máquinas que hagan falta hasta 2035; después no se agrega nada. Ningún puesto pasa el umbral. Se "
                 "pueden cambiar a mano; los colores se recalculan solos.", 14)

    titulos = {"util": "Aprovechamiento consolidado (recargas + mantenimiento y alta de sustitutos)",
               "maq": "Máquinas en el puesto", "tur": "Turnos del puesto"}
    for clave, r0 in bloques.items():
        cab(ws, r0, [titulos[clave]] + ANIOS)
        for k in range(N_PUESTOS):
            r = r0 + 1 + k
            rp = FILA_PUESTO + k
            celda(ws, r, 1, f"=parametros_recargas!$B${rp}")
            for i in range(10):
                c = L(2 + i)
                rd = fila_dem(i, k)
                dem = f"'Demanda por puesto'!$P${rd}"
                cap = f"parametros_recargas!${C_UH}${rp}"
                mq = f"{c}{bloques['maq'] + 1 + k}"
                tu = f"{c}{bloques['tur'] + 1 + k}"
                if clave == "util":
                    celda(ws, r, 2 + i, f"=IF({cap}*{mq}*{tu}=0,0,{dem}/({cap}*{mq}*{tu}))", "0%")
                else:
                    celda(ws, r, 2 + i, plan[(k, i)][0 if clave == "maq" else 1], "0")
        rt = r0 + 1 + N_PUESTOS
        if clave == "util":
            celda(ws, rt, 1, "Aprovechamiento máximo", neg=True)
            for i in range(10):
                c = L(2 + i)
                celda(ws, rt, 2 + i, f"=MAX({c}{r0 + 1}:{c}{rt - 1})", "0%", res=True)
            celda(ws, r0, 12, "Recurso").font = F_B
            for k in range(N_PUESTOS):
                celda(ws, r0 + 1 + k, 12, f"=parametros_recargas!$D${FILA_PUESTO + k}").alignment = CEN
            colores_aprovechamiento(ws, r0 + 1, bloques["maq"] + 1, bloques["tur"] + 1, N_PUESTOS, "L", "$O$9")
        elif clave == "maq":
            celda(ws, rt, 1, "TOTAL máquinas / puestos", neg=True)
            for i in range(10):
                c = L(2 + i)
                celda(ws, rt, 2 + i, f"=SUM({c}{r0 + 1}:{c}{rt - 1})", "0", res=True)
        elif clave == "tur":
            celda(ws, rt, 1, "Turnos de la planta (máximo)", neg=True)
            for i in range(10):
                c = L(2 + i)
                celda(ws, rt, 2 + i, f"=MAX({c}{r0 + 1}:{c}{rt - 1})", "0", res=True)
    ws.freeze_panes = "B2"


def h_sust(wb):
    """Sustitutos (9.4.5): lote y ciclo por rubro, stock, alta (fábrica / compra), mantenimiento y pedido a fábrica."""
    ws = hoja(wb, SS, [44, 12, 12, 12, 12, 12, 12, 12, 11, 11, 11, 24])
    tit(ws, 1, "SUSTITUTOS (IRAM 3517-2:2020, 9.4.5)")
    nota(ws, 2, "Mismo tipo que el equipo retirado. Sale con el camión que retira el lote y vuelve con el que devuelve "
                "los recargados: nunca se deja en otro cliente en ese viaje.")
    cab(ws, 4, ["Parámetro", "Valor", "Fuente / criterio"])
    ws.merge_cells("C4:L4")
    filas = [("Días hábiles por mes", "=parametros_recargas!B2", "0", False, ""),
             ("Día de retiro", DIA_RETIRO, "0", True, "El sustituto sale con el camión que retira el lote"),
             ("Día de devolución", DIA_DEVOL, "0", True, "Vuelve con el camión que devuelve el lote recargado"),
             ("Minutos productivos por turno", "=parametros_recargas!B5*60", "0", False, ""),
             ("Planta", "=PROY_Indexado!B51", None, False,
              "Comercial: los ABC llegan de la planta industrial como cilindro vacío y se cargan acá. Industrial: los "
              "ABC salen terminados de la línea. El resto se compra"),
             ("Turnos por día desde el año del 2° turno", "=parametros_recargas!B8", "0", False,
              "Antes de ese año la planta trabaja 1 turno"),
             ("Año del 2° turno", "=parametros_recargas!B10", "0", False, "El mismo del aprovechamiento"),
             ("Superficie por matafuego (m²)", M2_MATAFUEGO, "0", True,
              "Dec. 351/79, art. 176: un matafuego cada 200 m² como mínimo (control de sentido del lote)")]
    for j, (t, v, f, d, fu) in enumerate(filas):
        r = 5 + j
        celda(ws, r, 1, t)
        celda(ws, r, 2, v, f, dato=d, res=not d).alignment = CEN
        celda(ws, r, 3, fu)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=12)

    # paso 1: lote, tiempo y ciclo por rubro
    tipos_rng = f"$C${S_TIPO}:$C${S_TIPO + 16}"
    tit(ws, S_RUB - 3, "PASO 1 — Lote, días en planta y ciclo por rubro")
    nota(ws, S_RUB - 2, "El lote del cliente viaja junto en su jaula: días en planta = minutos del lote ÷ minutos de "
                        "planta por día, redondeado hacia arriba (mínimo 1). Ciclo = retiro + días en planta + devolución.")
    cab(ws, S_RUB - 1, ["Concepto"] + RUBROS)
    conceptos = [("% de las recargas (mercado, zonas de la planta)", "0.0%"),
                 ("Lote por cliente (equipos que se retiran juntos)", "0"),
                 ("Superficie mínima que protege el lote (m²)", "#,##0"),
                 ("Minutos de operación por equipo (mix del rubro)", "0.0"),
                 ("Minutos del lote en planta", "#,##0"),
                 ("Días en planta con 1 turno", "0"),
                 ("Días en planta con 2 turnos", "0"),
                 ("Ciclo del sustituto con 1 turno (días)", "0"),
                 ("Ciclo del sustituto con 2 turnos (días)", "0")]
    for j, (t, fmt) in enumerate(conceptos):
        r = S_RUB + j
        celda(ws, r, 1, t, neg=(j >= 7))
        for c in range(7):
            col = L(2 + c)
            mix = f"PROY_Indexado!{col}${P_MIX}:{col}${P_MIX + 16}"
            f = [f"=SUM({mix})/SUM(PROY_Indexado!$B${P_MIX}:$H${P_MIX + 16})",
                 LOTE[c],
                 f"={col}{S_RUB + 1}*$B$12",
                 f"=IF(SUM({mix})=0,0,SUMPRODUCT({mix},{tipos_rng})/SUM({mix}))",
                 f"={col}{S_RUB + 1}*{col}{S_RUB + 3}",
                 f"=MAX(1,ROUNDUP({col}{S_RUB + 4}/$B$8,0))",
                 f"=MAX(1,ROUNDUP({col}{S_RUB + 4}/($B$8*$B$10),0))",
                 f"=$B$6+{col}{S_RUB + 5}+$B$7",
                 f"=$B$6+{col}{S_RUB + 6}+$B$7"][j]
            celda(ws, r, 2 + c, f, fmt, dato=(j == 1), res=(j >= 7))
    nota(ws, S_RUB + 9, "Lote: u. por cliente de Recargas.xlsx (hoja Supuestos), dato del grupo. La superficie es un "
                        "control de sentido con la dotación mínima del Dec. 351/79, art. 176.")

    # paso 2: tipos
    tit(ws, S_TIPO - 2, "PASO 2 — Tipos (el ciclo de cada tipo pesa los rubros según sus recargas)")
    cab(ws, S_TIPO - 1, ["Tipo", "Lleva sustituto (1 = sí)", "Minutos de operación por equipo",
                         "Ciclo con 1 turno (días)", "Ciclo con 2 turnos (días)", "Origen del alta", "Nota"])
    ws.merge_cells(f"G{S_TIPO - 1}:L{S_TIPO - 1}")
    pr = "parametros_recargas"
    for k in range(17):
        r = S_TIPO + k
        rt = FILA_TIPO + k
        fila_mix = f"PROY_Indexado!$B${P_MIX + k}:$H${P_MIX + k}"
        celda(ws, r, 1, f"=datos_base!A{2 + k}")
        celda(ws, r, 2, LLEVA[k], "0", dato=True)
        fm = (f"=(SUMPRODUCT(INDEX({pr}!$E${FILA_PUESTO}:$L${FILA_PUESTO + 27},0,MATCH({pr}!$B${rt},"
              f"{pr}!$E${FILA_PUESTO - 1}:$L${FILA_PUESTO - 1},0)),{pr}!${C_T}${FILA_PUESTO}:${C_T}$"
              f"{FILA_PUESTO + 27})*{pr}!$C${rt}+{pr}!$D${rt}*{pr}!${C_T}${FILA_PUESTO + 28})/60")
        celda(ws, r, 3, None, "0.0").value = ArrayFormula(f"C{r}", fm)
        for j, fr in enumerate((S_RUB + 7, S_RUB + 8)):
            celda(ws, r, 4 + j, f"=IF(SUM({fila_mix})=0,$B$6+1+$B$7,SUMPRODUCT({fila_mix},$B${fr}:$H${fr})"
                                f"/SUM({fila_mix}))", "0.0", res=True)
        celda(ws, r, 6, f'=IF(LEFT(A{r},3)="ABC",IF($B$9="Comercial","{CIL_VACIO}","{TERMINADO}"),"{COMPRA}")'
              ).alignment = CEN
        celda(ws, r, 7, "" if LLEVA[k] else "Automotor: sin canje")
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=12)

    origen = f"$F${S_TIPO}:$F${S_TIPO + 16}"

    def bloque(r0, titulo, explica, f_celda, pie, fmt):
        tit(ws, r0 - 3, titulo)
        nota(ws, r0 - 2, explica)
        cab(ws, r0 - 1, ["Tipo"] + ANIOS)
        for k in range(17):
            r = r0 + k
            celda(ws, r, 1, f"=datos_base!A{2 + k}")
            for i in range(10):
                celda(ws, r, 2 + i, f_celda(k, i), fmt)
        for j, (txt, f) in enumerate(pie):
            r = r0 + 17 + j
            celda(ws, r, 1, txt, neg=True)
            for i in range(10):
                c = L(2 + i)
                celda(ws, r, 2 + i, f(c, r0, r0 + 16), fmt, res=True)

    total = [("Total", lambda c, a, b: f"=SUM({c}{a}:{c}{b})")]
    bloque(S_STOCK, "PASO 3 — Stock (u.)",
           "Stock = recargas del año × % del mes pico del tipo ÷ días del mes × ciclo (con 2 turnos desde el año del 2° "
           "turno), redondeado hacia arriba. El stock del año anterior se mantiene en planta (9.4.5 y 9.2.1): entra "
           "al Plan mensual por la ruta de su tipo, 1/12 por mes",
           lambda k, i: (f"=ROUNDUP(datos_base!{L(2 + i)}{2 + k}*$B${S_TIPO + k}*IF({L(2 + i)}${S_STOCK - 1}>=$B$11,"
                         f"$E${S_TIPO + k},$D${S_TIPO + k})*{pr}!$Q${FILA_TIPO + k}/$B$5,0)"),
           total, "#,##0")
    bloque(S_ALTA, "PASO 4 — Alta anual (u.)",
           f"Alta = lo que crece el stock + stock del año anterior × % que se inutiliza (parametros_recargas fila "
           f"{F_INU}, columna de su agente). Cada alta lleva faja amarilla y tarjeta",
           lambda k, i: (f"={L(2 + i)}{S_STOCK + k}" if i == 0 else
                         f"=MAX(0,{L(2 + i)}{S_STOCK + k}-{L(1 + i)}{S_STOCK + k})+{L(1 + i)}{S_STOCK + k}"
                         f"*INDEX({pr}!$E${F_INU}:$L${F_INU},MATCH({pr}!$B${FILA_TIPO + k},{pr}!$E$12:$L$12,0))"),
           [(f"{CIL_VACIO} (se carga en esta planta)",
             lambda c, a, b: f'=SUMIF({origen},"{CIL_VACIO}",{c}{a}:{c}{b})'),
            (f"{TERMINADO} (sale de la línea)", lambda c, a, b: f'=SUMIF({origen},"{TERMINADO}",{c}{a}:{c}{b})'),
            ("Compra", lambda c, a, b: f'=SUMIF({origen},"{COMPRA}",{c}{a}:{c}{b})'),
            ("Kits de sustituto", lambda c, a, b: f"=SUM({c}{a}:{c}{b})")], "#,##0.0")
    nota(ws, S_ALTA + 22, f"Manuales importa B{S_ALTA}:K{S_ALTA + 3} (ABC 1 a 10 kg) y Carros B{S_ALTA + 13}:"
                          f"K{S_ALTA + 16} (25 a 100 kg). Comercial: los recibe como cilindro vacío y los carga acá "
                          f"(ruta del cilindro nuevo); industrial: salen terminados de la línea (proceso completo).")
    ws.freeze_panes = "B5"


def generar(planta, plan=None):
    if plan is None:
        plan = {(k, i): (1, 1) for k in range(N_PUESTOS) for i in range(10)}
    wb = Workbook()
    wb.remove(wb.active)
    h_links(wb, LINKS)
    h_proy(wb, planta)
    h_datos(wb)
    h_param(wb, planta)
    h_sust(wb)
    h_plan(wb)
    h_demanda(wb)
    h_resumen(wb, plan)
    wb.active = wb.sheetnames.index("Resumen Capacidad")
    os.makedirs(SALIDA, exist_ok=True)
    ruta = os.path.join(SALIDA, f"Dimensionamiento_Recargas_Planta_{planta}.xlsx")
    wb.save(ruta)
    return ruta


def modelo_demanda(planta):
    """Réplica del Excel en Python (demanda por puesto en u/h de un turno, mes pico; stock y alta de sustitutos)
    para cargar el plan de capacidad."""
    proy = {46: TASAS[0][1], 47: TASAS[1][1], 48: TASAS[2][1]}
    share = PCT_PLANTA[planta]
    filas = {r: v for r, t, v, f, d, fu in P_GEN}
    crit = {r: vals for r, t, u, fu, est, vals, ex in CRITERIOS}
    # recargas 2035 de cada agente (datos_base K), para los ensayos de laboratorio por equipo
    rec35 = [sum(REC[k][1][9] * share for k in range(17) if TIPOS[k][0] == a) for a in AGENTES]
    B, C = {}, {}

    def val(r):
        if r not in B:
            if r == 9:
                B[r] = sum(cv * 60 / sum(v for _, v in man) for _, man, (_, cv) in EQUIV) / len(EQUIV)
            else:
                v = filas[r]
                B[r] = ev(v) if isinstance(v, str) else v
        return B[r]

    def cr(r, j):
        if (r, j) not in C:
            v = crit[r][j]
            if r == F_LAB:
                C[(r, j)] = 0 if cr(F_LABD, j) == 0 else cr(F_LABD, j) * val(2) * 12 / rec35[j]
            else:
                C[(r, j)] = ev(v, j) if isinstance(v, str) else v
        return C[(r, j)]

    def ev(f, j=None):
        expr = f[1:].replace("{c}", L(5 + j)) if j is not None else f[1:]
        expr = re.sub(r"PROY_Indexado!\$?B\$?(\d+)", lambda m: repr(proy[int(m.group(1))]), expr)
        expr = re.sub(r"\$?([E-L])\$?(\d+)", lambda m: repr(cr(int(m.group(2)), ord(m.group(1)) - 69)), expr)
        expr = re.sub(r"\$?B\$?(\d+)", lambda m: repr(val(int(m.group(1)))), expr)
        return eval(expr)

    def techo(x):
        return math.ceil(x - 1e-9)

    tiempos = [ev(t) if isinstance(t, str) else t for (_, _, _, _, t, _, _) in PUESTOS]
    cap = [3600 / t if t else 0 for t in tiempos]
    fr = [[0] * 8 if ruta is None else [ev(x, j) if isinstance(x, str) else x for j, x in enumerate(ruta)]
          for (_, _, _, ruta, _, _, _) in PUESTOS]
    cil = [1 if p[0] in CILINDRO_NUEVO else 0 for p in PUESTOS]
    eq = [val(9) if rod else 1 for (_, rod, _, _) in TIPOS]
    mins = [(sum(fr[s][AGENTES.index(TIPOS[k][0])] * tiempos[s] for s in range(28)) * eq[k]
             + TIPOS[k][2] * tiempos[28]) / 60 for k in range(17)]
    mix = MIX_RUBRO[planta]
    turno, t2, a2 = val(5) * 60, val(8), val(10)
    min_r = [sum(mix[c][k] * mins[k] for k in range(17)) / sum(mix[c]) for c in range(7)]
    ciclo_r = {n: [DIA_RETIRO + max(1, techo(LOTE[c] * min_r[c] / (turno * n))) + DIA_DEVOL for c in range(7)]
               for n in (1, t2)}
    ciclo_t = {n: [sum(mix[c][k] * ciclo_r[n][c] for c in range(7)) / sum(mix[c][k] for c in range(7))
                   if sum(mix[c][k] for c in range(7)) else DIA_RETIRO + 1 + DIA_DEVOL for k in range(17)]
               for n in (1, t2)}
    inut = [cr(F_INU, AGENTES.index(TIPOS[k][0])) for k in range(17)]
    vacio = [planta == "Comercial" and REC[k][0].startswith("ABC") for k in range(17)]
    perfil = {k: E_MRP[FAM_ABC[k]][2] for k in K_ABC}
    # estacionalidad por rubro y perfil de cada tipo (mix de rubros de la planta)
    prof_r = []
    for rb, regla, fu in EST_RUBRO:
        if regla == "vtv":
            ex = (PAT_2024 + sum(PAT_2025)) / FLOTA_2025
            prof_r.append([(1 - ex) / 10 * (1 <= m <= 10) + ex * PAT_2025[m] / sum(PAT_2025) for m in range(12)])
        elif regla == "escuelas":
            prof_r.append([(1 - ESTATAL) / 12 + (ESTATAL if m == 1 else 0) for m in range(12)])
        else:
            prof_r.append([1 / 12] * 12)
    prof_t = []
    for k in range(17):
        tot = sum(mix[c][k] for c in range(7))
        prof_t.append([sum(mix[c][k] * prof_r[c][m] for c in range(7)) / tot if tot else 1 / 12 for m in range(12)])
    dem = [[0] * 10 for _ in range(N_PUESTOS)]
    stock, alta = [], []
    for i in range(10):
        rec = [v[i] * share for n, v in REC]
        n = t2 if ANIOS[i] >= a2 else 1
        st = [techo(rec[k] * LLEVA[k] * ciclo_t[n][k] * max(prof_t[k]) / val(2)) for k in range(17)]
        al = st if i == 0 else [max(0, st[k] - stock[-1][k]) + stock[-1][k] * inut[k] for k in range(17)]
        mant = stock[-1] if i else [0] * 17
        stock.append(st)
        alta.append(al)
        mes = [[rec[k] * prof_t[k][m] + mant[k] / 12 for k in range(17)] for m in range(12)]
        ag = [[sum(mes[m][k] * eq[k] for k in range(17) if TIPOS[k][0] == a) for a in AGENTES] for m in range(12)]
        mang = [sum(mes[m][k] * TIPOS[k][2] for k in range(17)) for m in range(12)]
        cil_m = [sum(al[k] * eq[k] * perfil[k][m] for k in K_ABC if vacio[k]) for m in range(12)]
        for s in range(N_PUESTOS):
            if PUESTOS[s][3] is None:
                pico = max(mang)
            else:
                pico = max(sum(fr[s][j] * ag[m][j] for j in range(8)) + cil[s] * cil_m[m] for m in range(12))
            dem[s][i] = pico / val(6)
    extra = {"stock": stock, "alta": alta, "ciclo_r": ciclo_r, "min_r": min_r, "ciclo_t": ciclo_t, "fr": fr,
             "tiempos": tiempos, "crit": {(r, j): cr(r, j) for r in crit for j in range(8)}, "eq": val(9),
             "perfil": prof_t}
    return dem, cap, [r for (_, _, r, _, _, _, _) in PUESTOS], val(7), val(8), val(10), extra


if __name__ == "__main__":
    for p in ("Comercial", "Industrial"):
        dem, cap, recurso, umbral, t_max, anio_2t, _ = modelo_demanda(p)
        print(generar(p, plan_capacidad(dem, cap, recurso, umbral, t_max, anio_2t)))
