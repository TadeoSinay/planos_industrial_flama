"""Actualiza los Aprovechamientos de Manuales y Carros del grupo (referencias/) y los deja en salida/.

- Hoja Links primera, con la tabla de links y todos los IMPORTRANGE del libro (Proyecciones, los dos
  Aprovechamientos de recargas, el MRP y, en Manuales, el Excel «Análisis alternativas corte chapa»). Las demás
  hojas leen Links: alcanza con pegar los links.
- Manuales: la versión 4 del grupo (26 puestos: G1 Ermaksan y G2 QC12Y, cilindradora Firesafer y soldadura
  longitudinal TIG del 1 kg). datos_base con Proyecciones y el alta (paso 4) de sustitutos ABC de cada planta;
  estacionalidad del MRP; piezas, golpes y tandas por hoja del Excel de corte.
- Carros: datos_base con IMPORTRANGE de Proyecciones y el bloque de sustitutos; estacionalidad del MRP; guillotinas
  nuevas (los tiempos siguen siendo los del W-CAR).
- Aprovechamiento indexado a las tablas de máquinas y turnos, cargadas con la regla del 90 % (2.° turno en un solo
  año) y colores por formato condicional.
"""

import os
import subprocess
import sys
import zipfile
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generar_dim_recargas as g  # noqa: E402

RAIZ = g.RAIZ
REF = os.path.join(RAIZ, "referencias")
RECALC = os.environ.get("RECALC", "/root/.claude/skills/synced/25c356cd-d646-4fc6-a248-83c7f917a94a_"
                                  "ed033efe-807e-46d6-9687-95e22d26e759/xlsx/scripts/recalc.py")
ANIO_2T = g.ANIO_2T

# Proyecciones, Indexado Dim Técnico (valores que trae el IMPORTRANGE)
TERM_MAN = [[4178.736543, 5997.303977, 6851.747758, 8467.135549, 12680.63271, 15620.25307, 20332.55326, 25132.07848,
             28573.33614, 32072.63913],
            [295.0303572, 709.7274503, 986.8680576, 1500.638174, 2312.627726, 3252.869887, 4102.315524, 5104.306587,
             6195.186243, 7424.858444],
            [2255.773025, 3846.978881, 5093.537591, 6823.115313, 10819.30694, 14605.96411, 18877.75984, 23368.18637,
             26967.72455, 30750.77872],
            [91.83318153, 177.2472967, 252.9787798, 358.0204941, 602.7181519, 872.605032, 1127.160545, 1395.444812,
             1612.350054, 1840.848079]]
CIL_MAN = [[0.0, 0.0, 7662.4, 26694.8, 49434.8, 71680.2, 101340.6, 131001.0, 146819.8, 170548.0],
           [0.0, 0.0, 299.6, 1043.8, 1933.0, 2803.0, 3963.0, 5123.0, 5741.4, 6669.0],
           [0.0, 0.0, 1478.8, 5152.0, 9540.8, 13834.2, 19558.6, 25283.0, 28335.8, 32915.0],
           [0.0, 0.0, 90.4, 315.0, 583.4, 845.8, 1195.4, 1545.0, 1731.4, 2011.0]]
TERM_CAR = [[34.05179136, 58.46214745, 75.20853998, 100.5298529, 155.8760035, 211.9344827, 272.6682332, 337.8486707,
             393.6274156, 453.2705332],
            [74.913941, 128.6167244, 165.458788, 221.1656765, 342.9272077, 466.255862, 599.8701129, 743.2670756,
             865.9803143, 997.1951731],
            [16.34485985, 28.06183078, 36.10009919, 48.25432941, 74.82048167, 101.7285517, 130.8807519, 162.1673619,
             188.9411595, 217.569856],
            [10.89657324, 18.70788719, 24.06673279, 32.16955294, 49.88032112, 67.81903447, 87.25383461, 108.1115746,
             125.960773, 145.0465706]]
CIL_CAR = [[0.0, 0.0, 14.4, 50.2, 93.0, 135.0, 191.0, 247.0, 277.0, 322.0],
           [0.0, 0.0, 31.8, 110.8, 205.2, 297.6, 420.8, 544.0, 609.6, 708.0],
           [0.0, 0.0, 7.0, 24.4, 45.2, 65.4, 92.2, 119.0, 133.4, 155.0],
           [0.0, 0.0, 4.6, 16.0, 29.6, 42.8, 60.4, 78.0, 87.6, 102.0]]

CORTE = os.path.join(g.SALIDA, "Analisis_alternativas_corte_chapa.xlsx")
FUENTES_MAN = [
    ("proy", "Proyecciones", g.URL_PROY,
     "Indexado Dim Técnico B3:K10 y B15:K22 → datos_base y carros en las guillotinas compartidas"),
    ("com", "Aprovechamiento Recargas Planta Comercial", g.PEGAR,
     "sobrestock_sustitutos B73:K76 y B86:K89: alta de sustitutos ABC (cilindros vacíos)"),
    ("ind", "Aprovechamiento Recargas Planta Industrial", g.PEGAR,
     "sobrestock_sustitutos B73:K76 y B86:K89: alta de sustitutos ABC (terminados)"),
    ("mrp", "MRP (FLAMA_MRP)", g.PEGAR, "ESTACIONALIDAD C4:N6 (E01, E02 y E03) → parametros_manuales"),
    ("corte", "Análisis alternativas corte chapa", g.PEGAR,
     "Corte B19:G27 (piezas y golpes por hoja; golpes por minuto) → parametros_manuales filas 118 a 137")]
FUENTES_CAR = [
    ("proy", "Proyecciones", g.URL_PROY, "Indexado Dim Técnico B7:K10 y B19:K22 → datos_base"),
    ("com", "Aprovechamiento Recargas Planta Comercial", g.PEGAR,
     "sobrestock_sustitutos B86:K89: alta de sustitutos ABC (cilindros vacíos) → datos_base"),
    ("ind", "Aprovechamiento Recargas Planta Industrial", g.PEGAR,
     "sobrestock_sustitutos B86:K89: alta de sustitutos ABC (terminados) → datos_base"),
    ("mrp", "MRP (FLAMA_MRP)", g.PEGAR, "ESTACIONALIDAD C6:N6 (E03) → parametros_carros filas 12 a 15"),
    ("corte", "Análisis alternativas corte chapa", g.PEGAR,
     "Corte B23:G27 (piezas y golpes por hoja de 25 a 100 kg; golpes por minuto) → parametros_carros filas 66 a 73"),
    ("man", "Aprovechamiento Manuales", g.PEGAR,
     "Resumen cap. sin sobrestock filas 3, 4, 32, 33, 61 y 62 → guillotinas compartidas (consolidado)")]
# modelo de tiempos de guillotina (W-MAN y W-CAR): minutos por corte con el golpe incluido y manipulación por hoja
T_MAN = [(0.3, 2), (0.3, 2), (0.3, 2), (0.4, 3)]           # 1 kg (como 2,5 kg), 2,5, 5 y 10 kg
T_CAR = [(0.4, "=10.6-12*0.4"), (0.4, "=8.8-9*0.4"), (0.6, "=14.5-8*0.6"), (0.6, "=14.5-8*0.6")]
FILA_CORTE_CAR = [0, 1, 4, 4]        # 70 y 100 kg salen de la hoja mixta (fila 5 del bloque de rodantes)
ANIOS_TXT = [str(a) for a in g.ANIOS]
PED_MAN = f"sobrestock_sustitutos!B{g.S_ALTA}:K{g.S_ALTA + 3}"
PED_CAR = f"sobrestock_sustitutos!B{g.S_ALTA + 13}:K{g.S_ALTA + 16}"
# Recurso de cada puesto (para el color de máquina u operario), según el equipo de parametros y el Word de procesos
REC_MAN = ["Manual"] + ["Máquina"] * 12 + ["Manual", "Máquina", "Máquina", "Máquina", "Manual", "Manual", "Máquina",
                                           "Máquina", "Manual", "Máquina", "Máquina", "Máquina", "Manual"]
N_MAN, R_MAQ_M, R_TUR_M = 26, 31, 60      # puestos de Manuales y filas de sus tablas de máquinas y turnos
REC_CAR = ["Manual", "Máquina", "Máquina", "Máquina", "Manual", "Máquina", "Máquina", "Máquina", "Manual", "Manual",
           "Máquina", "Máquina", "Manual", "Manual", "Máquina", "Manual", "Manual", "Manual"]


def pedidos():
    """Sustitutos ABC por modelo y año que pide cada planta de recargas (paso 4, alta, de su Aprovechamiento)."""
    out = {}
    for p in ("Comercial", "Industrial"):
        alta = g.modelo_demanda(p)[-1]["alta"]
        out[p] = [[alta[i][k] for i in range(10)] for k in g.K_ABC]
    return out


def estilo(src, dst):
    dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format


def refs(ws, r0, c0, LB, pos, nf, nc, clave=None, fmt=None):
    """Celdas que leen un bloque de la hoja Links (la celda conserva su formato)."""
    for i in range(nf):
        for j in range(nc):
            x = ws.cell(r0 + i, c0 + j)
            x.value = "=" + LB.ref(pos, i, j, clave)
            if fmt:
                x.number_format = fmt


def estacionalidad(ws, LB, perfiles, nota_txt):
    """Filas 12 a 15 de parametros: el perfil del MRP de cada modelo, desde la hoja Links."""
    for j, e in enumerate(perfiles):
        nom, rng, vals = g.E_MRP[e]
        refs(ws, 12 + j, 2, LB, LB.imp("mrp", rng, [vals], "0.0%", nom, [nom], g.MESES), 1, 12)
        ws.cell(12 + j, 15).value = f"MRP, ESTACIONALIDAD {nom} (hoja Links)"
    ws["A10"].value = nota_txt


def plan_con_sustitutos(ws, fila_db, pm, verif):
    """Plan mensual del grupo: cada fila suma los sustitutos de su modelo (terminados de la planta industrial y
    cilindros de la comercial), con el mismo perfil."""
    for i in range(10):
        r0 = 1 + 11 * i
        y = L(2 + i)
        for j in range(4):
            for fila, base, extra in ((r0 + 1 + j, 2 + j, fila_db["Industrial"] + j),
                                      (r0 + 5 + j, 9 + j, fila_db["Comercial"] + j)):
                for m in range(12):
                    viejo = f"=datos_base!${y}${base}*{pm}!{L(2 + m)}${12 + j}"
                    if ws.cell(fila, 4 + m).value != viejo:
                        verif.append((ws.title, ws.cell(fila, 4 + m).coordinate, ws.cell(fila, 4 + m).value, viejo))
                    ws.cell(fila, 4 + m).value = f"=(datos_base!${y}${base}+datos_base!${y}${extra})*{pm}!{L(2 + m)}${12 + j}"


def grises(wb):
    """Encabezados azules a gris claro con letra negra y celdas verdes de resultado a gris, como el resto del archivo."""
    gris = {"FF1F4E78": "FFF2F2F2", "FFE2EFDA": "FFD9D9D9"}
    for ws in wb:
        for fila in ws.iter_rows():
            for c in fila:
                if c.fill.fill_type == "solid" and c.fill.fgColor.type == "rgb" and c.fill.fgColor.rgb in gris:
                    if c.fill.fgColor.rgb == "FF1F4E78":
                        f = copy(c.font)
                        f.color = "FF000000"
                        c.font = f
                    c.fill = PatternFill("solid", fgColor=gris[c.fill.fgColor.rgb])


def bloque_sustitutos(ws, r0, modelos, rango, ped, fila_estilo, fila_tot, nota_txt, LB):
    """datos_base: terminados (industrial) y cilindros (comercial) para sustitutos, desde la hoja Links."""
    ws.cell(r0, 1).value = "Sustitutos ABC para las plantas de recargas"
    for c in range(1, 12):
        estilo(ws.cell(1, c), ws.cell(r0, c))
    for i in range(10):
        ws.cell(r0, 2 + i).value = g.ANIOS[i]
    filas = {}
    for b, (pl, txt, lk) in enumerate([("Industrial", "terminado, planta industrial", "ind"),
                                       ("Comercial", "cilindro vacío, planta comercial", "com")]):
        r1 = r0 + 1 + 5 * b
        for j, mo in enumerate(modelos):
            ws.cell(r1 + j, 1).value = f"{mo} — sustituto ({txt})"
            for c in range(1, 12):
                estilo(ws.cell(fila_estilo, c), ws.cell(r1 + j, c))
        pos = LB.imp(lk, rango, ped[pl], "#,##0", f"alta de sustitutos ABC ({txt})", modelos, ANIOS_TXT)
        refs(ws, r1, 2, LB, pos, 4, 10, lk)
        rt = r1 + 4
        ws.cell(rt, 1).value = f"Total sustitutos {'terminados' if b == 0 else 'cilindros vacíos'}"
        for c in range(1, 12):
            estilo(ws.cell(fila_tot, c), ws.cell(rt, c))
        for i in range(10):
            col = L(2 + i)
            ws.cell(rt, 2 + i).value = f"=SUM({col}{r1}:{col}{r1 + 3})"
        filas[pl] = r1
    ws.cell(r0 + 11, 1).value = nota_txt
    ws.cell(r0 + 11, 1).font = Font(name="Arial", size=9, italic=True)
    return filas


def resumen(ws, n, r_maq, r_tur, util, recursos, umbral_ref, celda_umbral, legend_row, piso_solo=False):
    """Fórmulas de aprovechamiento indexadas a las tablas, columna Recurso, umbral local y colores."""
    for k in range(n):
        for i in range(10):
            ws.cell(2 + k, 2 + i).value = util(k, i, f"{L(2 + i)}{r_maq + k}", f"{L(2 + i)}{r_tur + k}")
            ws.cell(2 + k, 2 + i).fill = PatternFill(fill_type=None)
    assert ws["L1"].value is None, ws.title
    estilo(ws["K1"], ws["L1"])
    ws["L1"].value = "Recurso"
    for k in range(n):
        estilo(ws.cell(2 + k, 1), ws.cell(2 + k, 12))
        ws.cell(2 + k, 12).value = recursos[k]
    rm, rn = celda_umbral
    ws[rm].value = "Umbral de aprovechamiento"
    ws[rm].font = Font(name="Arial", size=10, bold=True)
    ws[rn].value = f"={umbral_ref}"
    ws[rn].number_format = "0%"
    ws.cell(legend_row, 13).value = "Colores automáticos contra el año anterior · Word de recargas § 9"
    ws.cell(legend_row, 13).font = Font(name="Arial", size=9, italic=True)
    g.colores_aprovechamiento(ws, 2, r_maq, r_tur, n, "L", "$" + rn[0] + "$" + rn[1:], sin_maquina_es_terc=True)


def piso(ws, n, r_maq, r_tur, mitad, multiplica):
    """Máquinas y turnos efectivos de lo que ya estaba cargado: el "/2" escrito a mano cuenta como 2 turnos (sin
    sobrestock la fórmula no leía la tabla; con sobrestock dividía la tabla por 2)."""
    out = {}
    for k in range(n):
        for i in range(10):
            m = ws.cell(r_maq + k, 2 + i).value
            t = ws.cell(r_tur + k, 2 + i).value or 1
            h = 2 if mitad(k, i) else 1
            out[(k, i)] = (1 if m is None else int(m), int(t * h if multiplica else max(t, h)))
    return out


def mitades(ws, n):
    """Celdas de aprovechamiento con "/2" escrito a mano."""
    c = {(k, i) for k in range(n) for i in range(10) if "/2" in str(ws.cell(2 + k, 2 + i).value)}
    return lambda k, i: (k, i) in c


def escribir_tablas(ws, n, r_maq, r_tur, plan):
    for k in range(n):
        for i in range(10):
            ws.cell(r_maq + k, 2 + i).value = plan[(k, i)][0]
            ws.cell(r_tur + k, 2 + i).value = plan[(k, i)][1]


def recalc(ruta):
    out = subprocess.run([sys.executable, RECALC, ruta, "180"], capture_output=True, text=True)
    print(out.stdout.strip()[:400])


def arreglar_dummy(ruta):
    """LibreOffice guarda __xludf.DUMMYFUNCTION en minúscula: se vuelve a poner como lo exporta Google."""
    tmp = ruta + ".tmp"
    with zipfile.ZipFile(ruta) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            d = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/"):
                d = d.replace(b"__xludf.dummyfunction", b"__xludf.DUMMYFUNCTION").replace(
                    b"_xludf.dummyfunction", b"_xludf.DUMMYFUNCTION")
            zo.writestr(it, d)
    os.replace(tmp, ruta)


def valores_corte(r0, n):
    """Piezas, golpes y tandas por hoja y golpes por minuto del Excel de corte (valor que trae el IMPORTRANGE)."""
    ws = load_workbook(CORTE, data_only=True)["Corte"]
    return [[ws.cell(r0 + i, 2 + j).value for j in range(6)] for i in range(n)]


def corte_manuales(pm, LB, ped):
    """Bloque CORTE DE CHAPA de parametros_manuales: tiempos de W-MAN y carga de carros en las guillotinas."""
    pm["A117"].value = ("CORTE DE CHAPA: DOS GUILLOTINAS EN SERIE (Word «Esquema de corte en guillotina»: G1 Ermaksan CNC "
                        "HGS 2600-6, G2 QC12Y 6 × 2500)")
    pm["A120"].value = "Parte de la manipulación por hoja que va a la guillotina 1 (el resto, a la 2)"
    pm["B120"].value = 0.5
    pm["B120"].number_format = "0%"
    pm["C120"].value = "Criterio: traer la hoja (G1) y recargar las franjas (G2) llevan lo mismo"
    pm["A121"].value = "Tiempos por corte y manipulación por hoja"
    pm["B121"].value = None
    pm["C121"].value = ("W-MAN, «Corte de chapa»: 0,3 min por corte y 2 min por hoja (0,4 y 3 en 2,0 mm); el 1 kg como el "
                        "2,5 kg. Carros: W-CAR (0,4 y 0,6 min por corte; min por hoja − cortes)")
    cab = ["Modelo", "Hoja (mm)", "Piezas por hoja", "Golpes G1 por hoja", "Golpes G2 por hoja",
           "Tiempo por corte (min)", "Manipulación por hoja (min)", "G1 (s/cuerpo)", "G2 (s/cuerpo)",
           "Recipientes 2035 (peso)"]
    for j, t in enumerate(cab):
        estilo(pm.cell(123, min(1 + j, 9)), pm.cell(123, 1 + j))
        pm.cell(123, 1 + j).value = t
    for k in range(4):
        r = 124 + k
        for col in "FGHIJ":
            estilo(pm[f"D{r}"], pm[f"{col}{r}"])
        pm[f"F{r}"].value, pm[f"G{r}"].value = T_MAN[k]
        pm[f"H{r}"].value = f"=(D{r}*MAX(F{r},1/$B$118)+G{r}*$B$120)*60/C{r}"
        pm[f"I{r}"].value = f"=(E{r}*MAX(F{r},1/$B$119)+G{r}*(1-$B$120))*60/C{r}"
        pm[f"J{r}"].value = f"=datos_base!$K${2 + k}+datos_base!$K${9 + k}"
        for col, fmt in (("F", "0.0"), ("G", "0.0"), ("H", "0.0"), ("I", "0.0"), ("J", "#,##0")):
            pm[f"{col}{r}"].number_format = fmt
    pm["G128"].value = None
    for col in "HI":
        estilo(pm["G124"], pm[f"{col}128"])
        pm[f"{col}128"].value = f"=SUMPRODUCT({col}124:{col}127,$J$124:$J$127)/SUM($J$124:$J$127)"
        pm[f"{col}128"].number_format = "0.0"
    pm["K20"].value, pm["K21"].value = "=$H$128", "=$I$128"
    pm["A129"].value = ("Piezas y golpes por hoja: Excel de corte (hoja Links). El tiempo tipo es uno solo por puesto: se "
                        "pondera por los recipientes de 2035 (terminados + cilindros vacíos).")
    # carros en las guillotinas compartidas
    estilo(pm["A117"], pm["A131"])
    pm["A131"].value = "GUILLOTINAS COMPARTIDAS: CARGA DE CARROS EN G1 Y G2 (los puestos 2 y 3 quedan consolidados)"
    pm["A132"].value = ("Mismo modelo de tiempos con los datos de W-CAR. La carga de carros pasa a unidades equivalentes de "
                        "manuales (tiempo de carros ÷ tiempo tipo del puesto) y se suma en las hojas Demanda puesto.")
    car = ["25 kg", "50 kg", "70 kg", "100 kg"]
    pos = LB.imp("corte", "Corte!B23:G27", valores_corte(23, 5), "0.00",
                 "piezas y golpes por hoja (25 a 100 kg y hoja mixta)", car + ["Hoja mixta 70 + 100 kg"],
                 ["Piezas por hoja", "Golpes G1 por hoja", "Golpes G2 por hoja", "Tandas G2 por hoja",
                  "Golpes/min G1", "Golpes/min G2"])
    cab2 = cab[:7] + ["G1 (min/cuerpo)", "G2 (min/cuerpo)"]
    for j, t in enumerate(cab2):
        estilo(pm.cell(123, 1 + j), pm.cell(133, 1 + j))
        pm.cell(133, 1 + j).value = t
    for k in range(4):
        r = 134 + k
        for col in "ABCDEFGHI":
            estilo(pm[f"{col}124"], pm[f"{col}{r}"])
        pm[f"A{r}"].value = car[k]
        pm[f"B{r}"].value = "1.500 × 3.000"
        for j, col in enumerate("CDE"):
            pm[f"{col}{r}"].value = "=" + LB.ref(pos, FILA_CORTE_CAR[k], j)
        pm[f"F{r}"].value, pm[f"G{r}"].value = T_CAR[k]
        pm[f"H{r}"].value = f"=(D{r}*MAX(F{r},1/$B$118)+G{r}*$B$120)/C{r}"
        pm[f"I{r}"].value = f"=(E{r}*MAX(F{r},1/$B$119)+G{r}*(1-$B$120))/C{r}"
        for col in "FGHI":
            pm[f"{col}{r}"].number_format = "0.00"
    # carros a fabricar por año
    tc = LB.imp("proy", "Indexado Dim Técnico!B7:K10", TERM_CAR, "#,##0", "carros terminados", car, ANIOS_TXT)
    cc = LB.imp("proy", "Indexado Dim Técnico!B19:K22", CIL_CAR, "#,##0", "cilindros de carros vacíos", car, ANIOS_TXT)
    ic = LB.imp("ind", PED_CAR, ped["Industrial"][4:], "#,##0", "alta de sustitutos ABC terminados (carros)", car,
                ANIOS_TXT)
    co = LB.imp("com", PED_CAR, ped["Comercial"][4:], "#,##0", "alta de sustitutos ABC cilindros (carros)", car,
                ANIOS_TXT)
    estilo(pm["A123"], pm["A139"])
    pm["A139"].value = "Carros a fabricar (terminados + cilindros + sustitutos)"
    for i in range(10):
        estilo(pm["B123"], pm.cell(139, 2 + i))
        pm.cell(139, 2 + i).value = g.ANIOS[i]
    for k in range(4):
        r = 140 + k
        estilo(pm["A124"], pm[f"A{r}"])
        pm[f"A{r}"].value = f"Carro {car[k]}"
        for i in range(10):
            x = pm.cell(r, 2 + i)
            estilo(pm["J124"], x)
            x.value = ("=" + LB.ref(tc, k, i) + "+" + LB.ref(cc, k, i) + "+" + LB.ref(ic, k, i, "ind") + "+"
                       + LB.ref(co, k, i, "com"))
    e03 = LB.imp("mrp", g.E_MRP[2][1], [g.E_MRP[2][2]], "0.0%", g.E_MRP[2][0], [g.E_MRP[2][0]], g.MESES)
    estilo(pm["A124"], pm["A145"])
    pm["A145"].value = "Perfil de producción de carros (MRP, E03)"
    for m in range(12):
        x = pm.cell(145, 2 + m)
        estilo(pm["F124"], x)
        x.value = "=" + LB.ref(e03, 0, m)
        x.number_format = "0.0%"
    for r0, col, kt, txt in ((147, "H", "K20", "G1"), (159, "I", "K21", "G2")):
        estilo(pm["A123"], pm[f"A{r0}"])
        pm[f"A{r0}"].value = f"Carga de carros en {txt} (unidades equivalentes de manuales por mes)"
        for m in range(12):
            estilo(pm["B123"], pm.cell(r0, 2 + m))
            pm.cell(r0, 2 + m).value = g.MESES[m]
        for i in range(10):
            r = r0 + 1 + i
            estilo(pm["A124"], pm[f"A{r}"])
            pm[f"A{r}"].value = g.ANIOS[i]
            y = L(2 + i)
            for m in range(12):
                x = pm.cell(r, 2 + m)
                estilo(pm["J124"], x)
                x.value = (f"=IF(${kt}=0,0,SUMPRODUCT(${y}$140:${y}$143,${col}$134:${col}$137)*60*"
                           f"{L(2 + m)}$145/${kt})")
                x.number_format = "#,##0.0"
    pm["A171"].value = ("Tiras de accesorios de carros en la G1: 8,6 min/día en el pico (W-CAR), en la disponibilidad "
                        "del puesto 2 (fila 53).")


def corte_carros(pc, LB):
    """parametros_carros: corte de cuerpos en G1 y G2 con el modelo de Manuales y los tiempos de W-CAR."""
    car = ["25 kg", "50 kg", "70 kg", "100 kg"]
    pos = LB.imp("corte", "Corte!B23:G27", valores_corte(23, 5), "0.00",
                 "piezas y golpes por hoja (25 a 100 kg y hoja mixta)", car + ["Hoja mixta 70 + 100 kg"],
                 ["Piezas por hoja", "Golpes G1 por hoja", "Golpes G2 por hoja", "Tandas G2 por hoja",
                  "Golpes/min G1", "Golpes/min G2"])
    estilo(pc["A43"], pc["A66"])
    pc["A66"].value = "CORTE DE CUERPOS EN LAS GUILLOTINAS G1 Y G2 (compartidas con manuales; mismo modelo que Manuales)"
    pc["A67"].value = "Parte de la manipulación por hoja que va a la guillotina 1"
    pc["B67"].value = 0.5
    pc["B67"].number_format = "0%"
    pc["C67"].value = "Criterio: igual que en Manuales (parametros_manuales B120)"
    cab = ["Modelo", "Piezas por hoja", "Golpes G1 por hoja", "Golpes G2 por hoja", "Golpes/min G1", "Golpes/min G2",
           "Tiempo por corte (min)", "Manipulación por hoja (min)", "G1 (min/cuerpo)", "G1 + G2 (min/cuerpo)",
           "G2 (min/cuerpo)"]
    for j, t in enumerate(cab):
        estilo(pc["B44"], pc.cell(68, 1 + j))
        pc.cell(68, 1 + j).value = t
    for k in range(4):
        r = 69 + k
        for j in range(11):
            estilo(pc["D45"], pc.cell(r, 1 + j))
        pc[f"A{r}"].value = f"Carro {car[k]}"
        for j, col in enumerate("BCD"):
            pc[f"{col}{r}"].value = "=" + LB.ref(pos, FILA_CORTE_CAR[k], j)
        pc[f"E{r}"].value = "=" + LB.ref(pos, 0, 4)
        pc[f"F{r}"].value = "=" + LB.ref(pos, 0, 5)
        pc[f"G{r}"].value, pc[f"H{r}"].value = T_CAR[k]
        pc[f"I{r}"].value = f"=(C{r}*MAX(G{r},1/E{r})+H{r}*$B$67)/B{r}"
        pc[f"K{r}"].value = f"=(D{r}*MAX(G{r},1/F{r})+H{r}*(1-$B$67))/B{r}"
        pc[f"J{r}"].value = f"=I{r}+K{r}"
        for col in "BCDGHIJK":
            pc[f"{col}{r}"].number_format = "0.00"
    pc["A73"].value = ("W-CAR, «Corte de chapas»: 0,4 min por corte en 3,2 mm y 0,6 en 4,75 mm; manipulación = min por "
                       "hoja − cortes × min por corte (25 kg: 10,6 y 12 cortes; 50 kg: 8,8 y 9; mixta: 14,5 y 8). En el "
                       "puesto 2 se suman las tiras y patas del W-CAR.")


def consolidado_carros(ws, LB):
    """Resumen Capacidad de Carros: aprovechamiento de G1 y G2 consolidado (lo calcula Manuales)."""
    nom = ["Guillotina 1 (manuales + carros)", "Guillotina 2 (manuales + carros)"]
    rv = load_workbook(os.path.join(g.SALIDA, "Aprovechamiento_Manuales.xlsx"), data_only=True)
    rv = rv["Resumen cap. sin sobrestock"]
    bloques = [("Aprovechamiento consolidado", "Resumen cap. sin sobrestock!B3:K4", 3, "0%"),
               ("Máquinas", "Resumen cap. sin sobrestock!B32:K33", 32, "0"),
               ("Turnos", "Resumen cap. sin sobrestock!B61:K62", 61, "0")]
    estilo(ws["A1"], ws["A65"])
    ws["A65"].value = "GUILLOTINAS COMPARTIDAS CON MANUALES (consolidado en el Aprovechamiento de Manuales)"
    r = 66
    for t, rng, r0, fmt in bloques:
        vals = [[rv.cell(r0 + k, 2 + i).value for i in range(10)] for k in range(2)]
        pos = LB.imp("man", rng, vals, fmt, t.lower(), nom, ANIOS_TXT)
        for k in range(2):
            estilo(ws["A2"], ws[f"A{r + k}"])
            ws[f"A{r + k}"].value = f"{t}: {nom[k]}"
            for i in range(10):
                x = ws.cell(r + k, 2 + i)
                estilo(ws["B23"], x)
                x.value = "=" + LB.ref(pos, k, i)
                x.number_format = fmt
        r += 2
    ws[f"A{r}"].value = "El puesto 2 de arriba es el operario de corte de carros; las máquinas se dimensionan en Manuales."


def manuales(ped):
    """Versión 4 del grupo (26 puestos) con la hoja Links y todo indexado."""
    wb = load_workbook(os.path.join(REF, "Dimensionamiento_Manuales.xlsx"))
    sal = os.path.join(g.SALIDA, "Aprovechamiento_Manuales.xlsx")
    LB = g.Links(FUENTES_MAN, iferror=("com", "ind"))
    grises(wb)
    db = wb["datos_base"]
    # sustitutos: solo los ABC de 1 a 10 kg (los demás tipos se compran); se sacan las filas de los otros tipos
    db.delete_rows(58, 13)
    db.delete_rows(37, 13)
    modelos = ["ABC 1 kg", "ABC 2,5 kg", "ABC 5 kg", "ABC 10 kg"]
    for r0, clave, rng, vals, tit in (
            (19, "proy", "Indexado Dim Técnico!B3:K6", TERM_MAN, "matafuegos terminados"),
            (26, "proy", "Indexado Dim Técnico!B15:K18", CIL_MAN, "cilindros vendidos vacíos"),
            (33, "ind", PED_MAN, ped["Industrial"][:4], "alta de sustitutos ABC terminados"),
            (41, "com", PED_MAN, ped["Comercial"][:4], "alta de sustitutos ABC cilindros vacíos")):
        refs(db, r0, 2, LB, LB.imp(clave, rng, vals, "#,##0", tit, modelos, ANIOS_TXT), 4, 10, clave)
    db["A18"].value = "PROYECCIONES — Matafuegos terminados (Indexado Dim Técnico B3:K6, hoja Links)"
    db["A25"].value = "PROYECCIONES — Cilindros vendidos vacíos (Indexado Dim Técnico B15:K18, hoja Links)"
    db["A32"].value = "PLANTA INDUSTRIAL — Sustitutos ABC terminados: alta anual (sobrestock_sustitutos, paso 4)"
    db["A40"].value = "PLANTA COMERCIAL — Sustitutos ABC cilindros vacíos: alta anual (sobrestock_sustitutos, paso 4)"
    db["A37"].value, db["A45"].value = "Total sustitutos planta industrial", "Total sustitutos planta comercial"
    db["A47"].value = "A la línea de manuales entran solo los ABC de 1 a 10 kg; los demás tipos se compran."
    for i in range(10):
        c = L(2 + i)
        db[f"{c}37"].value = f"=SUM({c}33:{c}36)"
        db[f"{c}45"].value = f"=SUM({c}41:{c}44)"
        for j in range(4):
            db[f"{c}{9 + j}"].value = f"={c}{26 + j}+{c}{41 + j}"
    pm = wb["parametros_manuales"]
    estacionalidad(pm, LB, [0, 1, 1, 1], "ESTACIONALIDAD (% de la demanda anual por mes; MRP, hoja Links)")
    pm["A8"].value = "Umbral de aprovechamiento"
    pm["C8"].value = "Límite duro (análisis de riesgos) · Word de recargas § 8"
    pm["A9"].value = "Año de inicio del 2° turno (antes: más máquinas)"
    pm["B9"].value = ANIO_2T
    pm["C9"].value = "Todos los turnos entran ese año · Word de recargas § 8"
    for c in "ABC":
        estilo(pm[f"{c}8"], pm[f"{c}9"])
    pm["B9"].number_format = "0"
    # corte de chapa: piezas y golpes por hoja y golpes por minuto del Excel de corte; tiempos de W-MAN
    pos = LB.imp("corte", "Corte!B19:G22", valores_corte(19, 4), "0.00",
                 "piezas, golpes y tandas por hoja y golpes por minuto (1 a 10 kg)", modelos,
                 ["Piezas por hoja", "Golpes G1 por hoja", "Golpes G2 por hoja", "Tandas G2 por hoja",
                  "Golpes/min G1", "Golpes/min G2"])
    refs(pm, 124, 3, LB, pos, 4, 3)
    pm["B118"].value = "=" + LB.ref(pos, 0, 4)
    pm["B119"].value = "=" + LB.ref(pos, 0, 5)
    corte_manuales(pm, LB, ped)
    pm["C20"].value = "Guillotina hidráulica Ermaksan CNC HGS 2600-6 (tope 1.000 mm), compartida con carros"
    pm["C21"].value = "Guillotina hidráulica QC12Y 6 × 2500 (tope 800 mm), compartida con carros"
    for r, t in ((20, "Corte de chapa — 1.er corte en franjas (guillotina 1, manuales + carros)"),
                 (21, "Corte de chapa — 2.º corte en piezas (guillotina 2, manuales + carros)")):
        pm[f"B{r}"].value = t
        pm[f"B{r + 33}"].value = t
    pm["O20"].value = ("(Golpes G1 × min por corte + parte de la manipulación por hoja) ÷ piezas; promedio ponderado "
                       "(fila 128). La carga de carros se suma en Demanda puesto (filas 148 a 157)")
    pm["O21"].value = ("(Golpes G2 × min por corte + resto de la manipulación) ÷ piezas; promedio ponderado (fila 128). "
                       "La carga de carros se suma en Demanda puesto (filas 160 a 169)")
    pm["E53"].value = "=C53*D53/60+8.6*$B$2/60"
    pm["G53"].value = "Más 8,6 min/día de tiras de accesorios de carros (W-CAR: placas 1,7 + patas 1,9 + 5,0)"
    # ensayo de pérdidas: tanque de 6 posiciones, en tandas (W-MAN); el Excel del grupo tenía 55 s por unidad
    pm["K42"].value = "=(60+120+60)/6"
    pm["O42"].value = ("W-MAN, «Ensayo de pérdidas»: tandas de 6 en el tanque; 1 min para bajarlos + 2 min de observación + "
                       "1 min para sacarlos = 4 min por tanda, 40 s por matafuego")
    hojas = [("Resumen cap. sin sobrestock", "Demanda puesto sin sobrestock"),
             ("Resumen cap. con sobrestock", "Demanda puesto con sobrestock")]
    uno = {(k, i): (1, 1) for k in range(N_MAN) for i in range(10)}
    for rn, dn in hojas:
        ws = wb[rn]

        def util(k, i, mq, tu, dn=dn):
            pr = 28 * i + 2 + k
            return (f"=IF({mq}*{tu}*parametros_manuales!$M${19 + k}*parametros_manuales!$B$6=0,0,'{dn}'!$P${pr}/"
                    f"(parametros_manuales!$M${19 + k}*{mq}*{tu}*parametros_manuales!$B$6))")
        resumen(ws, N_MAN, R_MAQ_M, R_TUR_M, util, REC_MAN, "parametros_manuales!$B$8", ("M13", "N13"), 14)
        escribir_tablas(ws, N_MAN, R_MAQ_M, R_TUR_M, uno)
    for dn in ("Demanda puesto sin sobrestock", "Demanda puesto con sobrestock"):
        d = wb[dn]
        for i in range(10):
            for k, r0 in ((1, 148), (2, 160)):
                for m in range(12):
                    x = d.cell(28 * i + 2 + k, 2 + m)
                    x.value = f"{x.value}+parametros_manuales!${L(2 + m)}${r0 + i}"
    wb["Resumen cap. con sobrestock"]["O1"].value = ("Máquinas y turnos: regla del 90 % con el 2.° turno en un solo "
                                                     "año (parametros_manuales B8 y B9); se pueden cambiar a mano.")
    LB.escribir(wb)
    wb.active = wb.sheetnames.index("Resumen cap. sin sobrestock")
    wb.save(sal)
    recalc(sal)
    # segunda pasada: máquinas y turnos con la regla del 90 % sobre la demanda nueva
    wv = load_workbook(sal, data_only=True)
    wb = load_workbook(sal)
    cap = [wv["parametros_manuales"].cell(19 + k, 13).value for k in range(N_MAN)]
    b6 = wv["parametros_manuales"]["B6"].value
    umbral = wv["parametros_manuales"]["B8"].value
    for rn, dn in hojas:
        dem = [[wv[dn].cell(28 * i + 2 + k, 16).value / b6 for i in range(10)] for k in range(N_MAN)]
        plan = g.plan_capacidad(dem, cap, ["Máquina"] * N_MAN, umbral, 2, ANIO_2T, piso=uno)
        escribir_tablas(wb[rn], N_MAN, R_MAQ_M, R_TUR_M, plan)
    wb.save(sal)
    recalc(sal)
    arreglar_dummy(sal)
    return sal, []


def carros(ped):
    ent = os.path.join(REF, "Dimensionamiento_Carros.xlsx")
    sal = os.path.join(g.SALIDA, "Aprovechamiento_Carros.xlsx")
    wb = load_workbook(ent)
    verif = []
    LB = g.Links(FUENTES_CAR, iferror=("com", "ind"))
    grises(wb)
    estacionalidad(wb["parametros_carros"], LB, [2, 2, 2, 2],
                   "ESTACIONALIDAD (% de la demanda anual por mes; MRP, hoja Links)")
    db = wb["datos_base"]
    modelos = ["Carro ABC 25 kg", "Carro ABC 50 kg", "Carro ABC 70 kg", "Carro ABC 100 kg"]
    refs(db, 2, 2, LB, LB.imp("proy", "Indexado Dim Técnico!B7:K10", TERM_CAR, "#,##0", "carros terminados", modelos,
                              ANIOS_TXT), 4, 10)
    refs(db, 9, 2, LB, LB.imp("proy", "Indexado Dim Técnico!B19:K22", CIL_CAR, "#,##0", "cilindros vendidos vacíos",
                              modelos, ANIOS_TXT), 4, 10)
    fila_db = bloque_sustitutos(db, 34, modelos, PED_CAR, {p: v[4:] for p, v in ped.items()}, 2, 6,
                                "Sustitutos: alta (paso 4) de cada Aprovechamiento de recargas · Word de recargas § 9",
                                LB)
    db["A15"].value = "Total recipientes a fabricar (con sustitutos)"
    for i in range(10):
        c = L(2 + i)
        db[f"{c}15"].value = f"={c}6+{c}13+{c}39+{c}44"
    db["A24"].value = ("Fuente: Proyecciones, Indexado Dim Técnico (hoja Links): terminados filas 7 a 10 y cilindros "
                       "filas 19 a 22.")
    plan_con_sustitutos(wb["Plan mensual"], fila_db, "parametros_carros", verif)
    pc = wb["parametros_carros"]
    pc["A9"].value = "Año de inicio del 2° turno (antes: más máquinas)"
    pc["B9"].value = ANIO_2T
    pc["C9"].value = "Todos los turnos entran ese año · Word de recargas § 8"
    for c in "ABC":
        estilo(pc[f"{c}8"], pc[f"{c}9"])
    pc["B9"].number_format = "0"
    pc["A8"].value = "Umbral de aprovechamiento"
    pc["C8"].value = "Límite duro (análisis de riesgos) · Word de recargas § 8"
    # guillotinas compartidas con manuales: el corte de cuerpos sale del mismo modelo que Manuales (W-CAR)
    pc["C20"].value = ("Guillotinas Ermaksan CNC HGS 2600-6 (1.er corte) y QC12Y 6 × 2500 (2.º corte), compartidas con "
                       "manuales + mesa elevadora de tijera 3 t + ventosas")
    corte_carros(pc, LB)
    for col, viejo in (("D", "10.6/9"), ("E", "8.8/6"), ("F", "14.5/4"), ("G", "14.5/4")):
        f = pc[f"{col}20"].value
        assert f.startswith("=" + viejo), f
        pc[f"{col}20"].value = f"=$J${69 + 'DEFG'.index(col)}" + f[1 + len(viejo):]
    pc["A42"].value = ((pc["A42"].value or "") + " Corte: cuerpos con G1 y G2 (filas 66 a 73, mismo modelo que Manuales) "
                       "más las tiras y patas del W-CAR; la carga de las guillotinas se consolida en Manuales.")
    pw = wb["Proceso (Word)"]
    pw["F6"].value = (pw["F6"].value or "").replace(
        "Guillotina CENA 8 × 3200 NC E21 (USD 36.800 + IVA, compartida con manuales)",
        "Guillotinas Ermaksan CNC HGS 2600-6 y QC12Y 6 × 2500 (compartidas con manuales)")
    ws = wb["Resumen Capacidad"]
    estilo(ws["M5"], ws["M6"])
    ws["M6"].value = "Revisar"
    ws["M6"].fill = PatternFill("solid", fgColor=g.COLOR_ACC["Revisar"], bgColor=g.COLOR_ACC["Revisar"])
    piso_c = piso(ws, 18, 23, 44, mitades(ws, 18), False)

    def util(k, i, mq, tu):
        return f"=IF({mq}*{tu}=0,0,'Demanda por puesto'!$P${20 * i + 2 + k}/({mq}*{tu}))"
    resumen(ws, 18, 23, 44, util, REC_CAR, "parametros_carros!$B$8", ("M8", "N8"), 9)
    consolidado_carros(ws, LB)
    LB.escribir(wb)
    wb.active = wb.sheetnames.index("Resumen Capacidad")
    wb.save(sal)
    recalc(sal)
    wv = load_workbook(sal, data_only=True)
    wb = load_workbook(sal)
    umbral = wv["parametros_carros"]["B8"].value
    dem = [[wv["Demanda por puesto"].cell(20 * i + 2 + k, 16).value for i in range(10)] for k in range(18)]
    plan = g.plan_capacidad(dem, [1] * 18, ["Máquina"] * 18, umbral, 2, ANIO_2T, piso=piso_c)
    escribir_tablas(wb["Resumen Capacidad"], 18, 23, 44, plan)
    wb.save(sal)
    recalc(sal)
    arreglar_dummy(sal)
    return sal, verif


if __name__ == "__main__":
    ped = pedidos()
    for f in (manuales, carros):
        ruta, verif = f(ped)
        print(ruta, "| fórmulas de demanda que no seguían el patrón:", len(verif))
        for v in verif[:5]:
            print("   ", v)
