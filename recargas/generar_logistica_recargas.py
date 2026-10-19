"""Genera salida/Recargas.xlsx: logística de recargas de las dos plantas (camión, viajes, sustitutos y pulmón) en el
mes pico de 2035.

Hojas: Links -> Resumen -> Supuestos -> Importados -> Cálculo. Lo que viene de la demanda sale por IMPORTRANGE de los
Dimensionamientos de recargas (recargas 2035 por tipo, perfil mensual, mix y estacionalidad por rubro, lote y
sustitutos); el resto son los supuestos de logística con su fuente. Las dos plantas van lado a lado.
"""

import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generar_dim_recargas as g  # noqa: E402
from generar_dim_recargas import cab, celda, hoja, nota, tit  # noqa: E402

PLANTAS = ["Comercial", "Industrial"]
LINKS = [("Dimensionamiento Recargas Planta Comercial", g.PEGAR, "Lo que trae la hoja Importados"),
         ("Dimensionamiento Recargas Planta Industrial", g.PEGAR, "Lo que trae la hoja Importados")]
LK = {"Comercial": "Links!$B$4", "Industrial": "Links!$B$5"}
TIPOS = [n for n, v in g.REC]
# peso lleno: BOM FLAMA (Resumen, peso de catálogo) y, para los carros, Word de carros (Tabla 0, carro terminado)
PESO = [1.9, 4.6, 8.5, 16.5, 12.55, 12.6, 100, 8.5, 8.25, 9, 20, 8.5, 16.5, 53, 90, 136, 179]
F_PESO = ["BOM FLAMA"] * 13 + ["Word de carros, Tabla 0"] * 4
# u. por jaula de 1,2 × 1,0 m: Word de recargas (A: 72 de 1 kg; B: 30 de 2,5, 24 de 5 o 16 de 10 kg); el resto, con la
# huella del catálogo FLAMA (ancho × profundidad) en la clase del Word que la contiene; rodantes por posición en planta
U_JAULA = [72, 30, 24, 16, 16, 16, 2, 24, 16, 30, 16, 24, 16, 2, 2, 2, 1]
JAULA = ["A"] + ["B"] * 5 + ["Rodantes"] + ["B"] * 6 + ["Rodantes"] * 4
MIN_U = [0.5, 1.5, 1.5, 1.5, 1.5, 1.5, 1.0]   # min por u. en el cliente (criterio del grupo; Word § 10.4)

# Supuestos: secciones con (clave, texto, valor, formato, fuente). "{clave}" en una fórmula = fila de esa clave.
SUP = [
    ("1. Camión", [
        ("modelo", "Vehículo", "Iveco Daily 55C17 con caja de baranda rebatible", None, "Word de recargas · § 10.1"),
        ("carga", "Carga útil (kg)", 2000, "#,##0", "Word de recargas: aprox. 2 t"),
        ("pos", "Posiciones de jaula 1,2 × 1,0 m en la caja", 6, "0", "2 × 3 en caja de 3,7 × 2,1 m · § 10.1"),
        ("pers", "Personas a bordo (chofer + ayudante)", 2, "0", "Retiro y colocación de sustitutos · § 10.1"),
        ("kgp", "Peso por persona (kg)", 75, "0", "Masa normalizada por ocupante (Dir. 2007/46/CE)"),
        ("kgj", "Peso de una jaula (kg)", 75, "0", "Jaula plegable 1,2 × 1,0 m (Palletower Hypacage)"),
        ("tara", "Tara de las jaulas (kg)", "=B{pos}*B{kgj}", "#,##0", "Posiciones × peso de jaula"),
        ("neta", "Carga neta por viaje (kg)", "=B{carga}-B{pers}*B{kgp}-B{tara}", "#,##0",
         "Carga útil − personas − jaulas"),
        ("precio", "Precio del camión con caja (ARS)", None, "#,##0", "A cotizar")]),
    ("2. Recorrido", [
        ("turnos", "Turnos de logística por día", 2, "0", "Los mismos de la planta (Ley 11.544) · § 10.2"),
        ("iv", "Ida y vuelta planta ↔ zona de clientes (h por viaje)", 1.5, "0.0", "2 × 25 km (TomTom 2024) · § 10.2"),
        ("parada", "Tiempo fijo por parada (min)", 20, "0", "Criterio · § 10.2"),
        ("manejo", "Manejo entre paradas (min)", 15, "0", "5 km a 20 km/h (TomTom 2024) · § 10.2"),
        ("umbral", "Ocupación máxima de la flota", 0.9, "0%", "El mismo límite de los puestos · § 8")]),
    ("3. Pulmón", [
        ("d_rec", "Días de llegada que esperan en recepción", 0.5, "0.0", "Entra en el turno siguiente · § 10.3"),
        ("d_desp", "Días de recargados que esperan el recorrido", 0.5, "0.0", "Sale en el turno siguiente · § 10.3"),
        ("medio", "Factor de jaulas a medio llenar", 1.15, "0.00", "Jaulas por cliente y agente · § 10.3"),
        ("vacias", "Jaulas vacías de reserva (% de las ocupadas)", 0.15, "0%", "Criterio · § 10.3"),
        ("niveles", "Niveles de la estantería de sustitutos", 3, "0", "Criterio · § 10.3"),
        ("apil", "Jaulas vacías apiladas por posición", 4, "0", "Hypacage: apilable de a 4"),
        ("m2pos", "Superficie por posición con pasillo (m²)", "=1.2*1*2", "0.0", "Pasillo compartido por dos filas")]),
]
S = {}
_r = 3
for _sec, _items in SUP:
    _r += 1
    for _it in _items:
        S[_it[0]] = _r
        _r += 1
    _r += 1


def sp(clave):
    return f"Supuestos!$B${S[clave]}"


# Importados: bloque de cada planta (filas relativas a su inicio)
BLK = 34
IMP0 = {"Comercial": 3, "Industrial": 3 + BLK}


def fi(p, que, k=0):
    """Fila de Importados: 'tipo' (k = 0..16), 'tot', 'lote', 'dias', 'horas', 'rub' (k = 0..6)."""
    r0 = IMP0[p]
    return {"tipo": r0 + 2 + k, "tot": r0 + 19, "lote": r0 + 20, "dias": r0 + 21, "horas": r0 + 22,
            "rub": r0 + 25 + k}[que]


def datos_dim(planta):
    """Valores actuales del Dimensionamiento de la planta (lo que trae cada IMPORTRANGE)."""
    wb = load_workbook(os.path.join(g.SALIDA, f"Dimensionamiento_Recargas_Planta_{planta}.xlsx"), data_only=True)
    db, px, ss, pr = wb["datos_base"], wb["PROY_Indexado"], wb[g.SS], wb["parametros_recargas"]
    return {"rec": [[db.cell(2 + k, 11).value] for k in range(17)],
            "lle": [[ss.cell(g.S_TIPO + k, 2).value] for k in range(17)],
            "stk": [[ss.cell(g.S_STOCK + k, 11).value] for k in range(17)],
            "mix": [[px.cell(g.P_MIX + k, 2 + c).value for c in range(7)] for k in range(17)],
            "per": [[pr.cell(g.FILA_TIPO + k, 5 + m).value for m in range(12)] for k in range(18)],
            "lot": [[ss.cell(g.S_RUB + 1, 2 + c).value for c in range(7)]],
            "dias": [[pr["B2"].value]], "horas": [[pr["B5"].value]],
            "rub": [[pr.cell(g.FILA_RUB + c, 5 + m).value for m in range(12)] for c in range(7)]}


def h_importados(wb, datos):
    ws = hoja(wb, "Importados", [36] + [7.5] * 12 + [12, 11, 13] + [10] * 7 + [10])
    tit(ws, 1, "Importados de los Dimensionamientos de recargas (IMPORTRANGE; el link va en la hoja Links)")
    for p in PLANTAS:
        r0, d, lk = IMP0[p], datos[p], LK[p]
        t1 = fi(p, "tipo")
        tit(ws, r0, f"Planta {p.lower()}")
        cab(ws, r0 + 1, ["Tipo: perfil mensual de recargas (% del año)"] + g.MESES
            + ["Recargas 2035 (u)", "Lleva sustituto (1 = sí)", "Stock de sustitutos 2035 (u)"] + g.RUBROS
            + ["Total mix"])
        for k in range(17):
            r = t1 + k
            celda(ws, r, 1, TIPOS[k])
            celda(ws, r, 24, f"=SUM(Q{r}:W{r})", "#,##0", res=True)
        g.importrange(ws, t1, 2, lk, f"parametros_recargas!E{g.FILA_TIPO}:P{g.FILA_TIPO + 17}", d["per"], "0.0%")
        g.importrange(ws, t1, 14, lk, "datos_base!K2:K18", d["rec"], "#,##0")
        g.importrange(ws, t1, 15, lk, f"sobrestock_sustitutos!B{g.S_TIPO}:B{g.S_TIPO + 16}", d["lle"], "0")
        g.importrange(ws, t1, 16, lk, f"sobrestock_sustitutos!K{g.S_STOCK}:K{g.S_STOCK + 16}", d["stk"], "#,##0")
        g.importrange(ws, t1, 17, lk, f"PROY_Indexado!B{g.P_MIX}:H{g.P_MIX + 16}", d["mix"], "#,##0")
        r = fi(p, "tot")
        celda(ws, r, 1, "Toda la planta", neg=True)
        celda(ws, r, 14, f"=SUM(N{t1}:N{t1 + 16})", "#,##0", res=True)
        celda(ws, r, 16, f"=SUM(P{t1}:P{t1 + 16})", "#,##0", res=True)
        celda(ws, fi(p, "lote"), 1, "Lote por cliente (u.)")
        g.importrange(ws, fi(p, "lote"), 17, lk, f"sobrestock_sustitutos!B{g.S_RUB + 1}:H{g.S_RUB + 1}", d["lot"], "0")
        celda(ws, fi(p, "dias"), 1, "Días hábiles por mes")
        g.importrange(ws, fi(p, "dias"), 2, lk, "parametros_recargas!B2", d["dias"], "0")
        celda(ws, fi(p, "horas"), 1, "Horas productivas por turno")
        g.importrange(ws, fi(p, "horas"), 2, lk, "parametros_recargas!B5", d["horas"], "0.00")
        cab(ws, fi(p, "rub") - 1, ["Rubro: estacionalidad (% del año)"] + g.MESES)
        for c, rb in enumerate(g.RUBROS):
            celda(ws, fi(p, "rub", c), 1, rb)
        g.importrange(ws, fi(p, "rub"), 2, lk, f"parametros_recargas!E{g.FILA_RUB}:P{g.FILA_RUB + 6}", d["rub"],
                      "0.0%")
    ws.freeze_panes = "B3"


def h_supuestos(wb):
    ws = hoja(wb, "Supuestos", [52, 16, 56])
    tit(ws, 1, "Supuestos de logística")
    nota(ws, 2, "Demanda, estacionalidad, mix, lote y sustitutos vienen de Importados · § 10")
    for sec, items in SUP:
        cab(ws, S[items[0][0]] - 1, [sec, "Valor", "Fuente / criterio"])
        for clave, t, v, fmt, fu in items:
            r = S[clave]
            formula = isinstance(v, str) and v.startswith("=")
            celda(ws, r, 1, t)
            celda(ws, r, 2, v.format(**S) if formula else v, fmt, dato=not formula, res=formula).alignment = g.CEN
            celda(ws, r, 3, fu)


# Cálculo: primera fila de cada punto
C_MES, C_T1, C_R1, C_CAM = 5, 13, 37, 49


def h_calculo(wb):
    ws = hoja(wb, "Cálculo", [44, 12, 11, 11, 11, 12, 12, 12, 12, 12, 12, 12, 40])
    tit(ws, 1, "Logística en el mes pico de 2035 (las dos plantas lado a lado)")
    nota(ws, 2, "Cada planta en su mes pico · § 10")
    col = {"Comercial": "B", "Industrial": "C"}

    # 1. mes pico
    cab(ws, C_MES - 1, ["1. Mes pico", "Comercial", "Industrial"])
    filas = [("N° del mes pico (mayor % de Toda la planta)",
              lambda p: (f"=MATCH(MAX(Importados!$B${fi(p, 'tot')}:$M${fi(p, 'tot')}),"
                         f"Importados!$B${fi(p, 'tot')}:$M${fi(p, 'tot')},0)"), "0"),
             ("Mes pico", lambda p: f"=INDEX(Importados!$B${IMP0[p] + 1}:$M${IMP0[p] + 1},{col[p]}{C_MES})", None),
             ("% de las recargas del año en el mes pico",
              lambda p: f"=INDEX(Importados!$B${fi(p, 'tot')}:$M${fi(p, 'tot')},{col[p]}{C_MES})", "0.0%"),
             ("Días hábiles por mes", lambda p: f"=Importados!$B${fi(p, 'dias')}", "0"),
             ("Horas productivas por turno", lambda p: f"=Importados!$B${fi(p, 'horas')}", "0.00")]
    for j, (t, fx, fmt) in enumerate(filas):
        r = C_MES + j
        celda(ws, r, 1, t, neg=(j == 1))
        for p in PLANTAS:
            celda(ws, r, 2 + PLANTAS.index(p), fx(p), fmt, res=True).alignment = g.CEN
    nmes = {p: f"${col[p]}${C_MES}" for p in PLANTAS}
    dias = {p: f"${col[p]}${C_MES + 3}" for p in PLANTAS}
    horas = {p: f"${col[p]}${C_MES + 4}" for p in PLANTAS}

    # 2. por tipo
    t17 = C_T1 + 16
    cab(ws, C_T1 - 2, ["2. Por tipo", "Peso lleno (kg)", "u. por jaula", "Jaula", "Rodante (1 = sí)",
                       "Comercial", "", "", "Industrial", "", "", "", "Fuente del peso"])
    ws.merge_cells(start_row=C_T1 - 2, start_column=6, end_row=C_T1 - 2, end_column=8)
    ws.merge_cells(start_row=C_T1 - 2, start_column=9, end_row=C_T1 - 2, end_column=11)
    cab(ws, C_T1 - 1, ["Tipo", "", "", "", "", "u/día", "Sustitutos por día", "Jaulas por día",
                       "u/día", "Sustitutos por día", "Jaulas por día", "", "§ 10.3"])
    for k in range(17):
        r = C_T1 + k
        celda(ws, r, 1, TIPOS[k])
        celda(ws, r, 2, PESO[k], "0.0", dato=True)
        celda(ws, r, 3, U_JAULA[k], "0", dato=True)
        celda(ws, r, 4, JAULA[k]).alignment = g.CEN
        celda(ws, r, 5, 1 if JAULA[k] == "Rodantes" else 0, "0")
        for i, p in enumerate(PLANTAS):
            c0 = 6 + 3 * i
            ri = fi(p, "tipo", k)
            u = L(c0)
            celda(ws, r, c0, f"=Importados!$N${ri}*INDEX(Importados!$B${ri}:$M${ri},{nmes[p]})/{dias[p]}", "#,##0.0")
            celda(ws, r, c0 + 1, f"={u}{r}*Importados!$O${ri}", "#,##0.0")
            celda(ws, r, c0 + 2, f"={u}{r}/$C{r}", "0.00")
        celda(ws, r, 13, F_PESO[k])
    tot_t = t17 + 1
    celda(ws, tot_t, 1, "Total", neg=True)
    for c in range(6, 12):
        celda(ws, tot_t, c, f"=SUM({L(c)}{C_T1}:{L(c)}{t17})", "#,##0.0", res=True)
    celda(ws, tot_t + 1, 1, "Peso medio por u. (kg)", neg=True)
    for i in range(2):
        u = L(6 + 3 * i)
        celda(ws, tot_t + 1, 6 + 3 * i, f"=SUMPRODUCT({u}{C_T1}:{u}{t17},$B${C_T1}:$B${t17})/{u}{tot_t}", "0.0",
              res=True)

    # 3. por rubro
    r7 = C_R1 + 6
    cab(ws, C_R1 - 2, ["3. Por rubro", "min por u. en el cliente", "Comercial", "", "", "", "Industrial", "", "", "",
                       "", "", "Fuente"])
    ws.merge_cells(start_row=C_R1 - 2, start_column=3, end_row=C_R1 - 2, end_column=6)
    ws.merge_cells(start_row=C_R1 - 2, start_column=7, end_row=C_R1 - 2, end_column=10)
    cab(ws, C_R1 - 1, ["Rubro", "", "u/día", "Lote (u/cliente)", "Paradas por día", "Horas de parada",
                       "u/día", "Lote (u/cliente)", "Paradas por día", "Horas de parada", "", "", ""])
    for c, rb in enumerate(g.RUBROS):
        r = C_R1 + c
        celda(ws, r, 1, rb)
        celda(ws, r, 2, MIN_U[c], "0.0", dato=True)
        for i, p in enumerate(PLANTAS):
            c0 = 3 + 4 * i
            a, b = fi(p, "tipo"), fi(p, "tipo", 16)
            q = L(17 + c)
            celda(ws, r, c0, f"=INDEX(Importados!$B${fi(p, 'rub', c)}:$M${fi(p, 'rub', c)},{nmes[p]})"
                             f"*SUMPRODUCT(Importados!$N${a}:$N${b},Importados!${q}${a}:${q}${b}/Importados!$X${a}:$X${b})"
                             f"/{dias[p]}", "#,##0.0")
            celda(ws, r, c0 + 1, f"=Importados!{q}${fi(p, 'lote')}", "0")
            celda(ws, r, c0 + 2, f"=2*{L(c0)}{r}/{L(c0 + 1)}{r}", "0.0")
            celda(ws, r, c0 + 3, f"={L(c0 + 2)}{r}*({sp('parada')}+{sp('manejo')})/60+2*{L(c0)}{r}*$B{r}/60", "0.0")
        celda(ws, r, 13, "Criterio · § 10.4")
    tot_r = r7 + 1
    celda(ws, tot_r, 1, "Total", neg=True)
    for c in (3, 5, 6, 7, 9, 10):
        celda(ws, tot_r, c, f"=SUM({L(c)}{C_R1}:{L(c)}{r7})", "#,##0.0", res=True)

    # 4. camión y viajes ({c} = columna de la planta, {h} = horas por turno de la planta)
    cab(ws, C_CAM - 1, ["4. Camión y viajes (§ 10.5)", "Comercial", "Industrial"])
    R = C_CAM
    cam = [("Equipos que se retiran por día (u)", lambda i: f"={L(6 + 3 * i)}{tot_t}", "#,##0.0", "Total por tipo"),
           ("Sustitutos que salen por día (u)", lambda i: f"={L(7 + 3 * i)}{tot_t}", "#,##0.0", "Total por tipo"),
           ("Carga por sentido por día (u)", lambda i: f"={{c}}{R}+{{c}}{R + 1}", "#,##0.0",
            "Sale con recargados + sustitutos y vuelve con lo retirado + sustitutos devueltos"),
           ("Peso medio por u. (kg)", lambda i: f"={L(6 + 3 * i)}{tot_t + 1}", "0.0", "Tabla por tipo"),
           ("Jaulas por u.", lambda i: f"={L(8 + 3 * i)}{tot_t}/{L(6 + 3 * i)}{tot_t}", "0.000",
            "Jaulas por día ÷ u/día (cada tipo en su jaula)"),
           ("Capacidad por viaje por espacio (u)", lambda i: f"={sp('pos')}/({{c}}{R + 4}*{sp('medio')})", "#,##0",
            "Posiciones de la caja ÷ (jaulas por u. × factor de jaulas a medio llenar)"),
           ("Capacidad por viaje por peso (u)", lambda i: f"={sp('neta')}/{{c}}{R + 3}", "#,##0",
            "Carga neta ÷ peso medio"),
           ("Capacidad por viaje (u)", lambda i: f"=MIN({{c}}{R + 5},{{c}}{R + 6})", "#,##0", "La menor de las dos"),
           ("Viajes por día por carga", lambda i: f"={{c}}{R + 2}/{{c}}{R + 7}", "0.00",
            "Carga por sentido ÷ capacidad por viaje"),
           ("Paradas por día", lambda i: f"={L(5 + 4 * i)}{tot_r}", "0.0", "Total por rubro"),
           ("Horas de parada y manipuleo por día", lambda i: f"={L(6 + 4 * i)}{tot_r}", "0.0", "Total por rubro"),
           ("Viajes por día por tiempo", lambda i: f"={{c}}{R + 10}/({{h}}-{sp('iv')})", "0.00",
            "Horas de parada ÷ (horas del turno − ida y vuelta): cada viaje entra en un turno"),
           ("Viajes por día", lambda i: f"=ROUNDUP(MAX({{c}}{R + 8},{{c}}{R + 11}),0)", "0",
            "El mayor de los dos, redondeado"),
           ("Horas de camión por día", lambda i: f"={{c}}{R + 12}*{sp('iv')}+{{c}}{R + 10}", "0.0",
            "Viajes × ida y vuelta + horas de parada y manipuleo"),
           ("Camiones necesarios", lambda i: f"=ROUNDUP({{c}}{R + 13}/({{h}}*{sp('turnos')}*{sp('umbral')}),0)", "0",
            "Horas de camión ÷ (horas por turno × turnos de logística × ocupación máxima)"),
           ("Ocupación de la flota", lambda i: f"={{c}}{R + 13}/({{c}}{R + 14}*{{h}}*{sp('turnos')})", "0%",
            "Horas de camión ÷ horas disponibles"),
           ("Peso por viaje (kg)", lambda i: f"={{c}}{R + 2}*{{c}}{R + 3}/{{c}}{R + 12}", "#,##0",
            "Control: tiene que quedar por debajo de la carga neta"),
           ("Inversión en camiones (ARS)",
            lambda i: f'=IF({sp("precio")}="","a cotizar",{{c}}{R + 14}*{sp("precio")})', "#,##0",
            "Camiones × precio (Supuestos)")]
    for j, (t, fx, fmt, como) in enumerate(cam):
        r = R + j
        clave = t in ("Viajes por día", "Camiones necesarios")
        celda(ws, r, 1, t, neg=clave)
        for i, p in enumerate(PLANTAS):
            f = fx(i).replace("{c}", L(2 + i)).replace("{h}", horas[p])
            celda(ws, r, 2 + i, f, fmt, res=clave).alignment = g.DER

    # 5. pulmón
    r1 = R + len(cam) + 2
    cab(ws, r1 - 1, ["5. Pulmón", "Comercial", "", "", "", "Industrial", "", "", "", "", "", "", "Dónde y cómo van"])
    ws.merge_cells(start_row=r1 - 1, start_column=2, end_row=r1 - 1, end_column=5)
    ws.merge_cells(start_row=r1 - 1, start_column=6, end_row=r1 - 1, end_column=9)
    cab(ws, r1, ["Zona", "Matafuegos (u)", "Jaulas de manuales", "Jaulas de rodantes", "Posiciones 1,2 × 1,0 m",
                 "Matafuegos (u)", "Jaulas de manuales", "Jaulas de rodantes", "Posiciones 1,2 × 1,0 m", "", "", "", ""])
    t_e, t_c = f"$E${C_T1}:$E${t17}", f"$C${C_T1}:$C${t17}"
    zonas = [("Recepción", "Jaulas A (1 kg) y B por cliente, en orden de llegada; rodantes en su jaula"),
             ("Despacho", "Jaulas armadas por recorrido y en el orden de entrega"),
             ("Sustitutos en planta", "Estantería con faja amarilla, por tipo; rodantes en el piso"),
             ("Jaulas vacías", "Apiladas, junto a despacho")]
    for j, (z, donde) in enumerate(zonas):
        r = r1 + 1 + j
        celda(ws, r, 1, z)
        for i in range(2):
            c0 = 2 + 4 * i
            u, s_, jj = L(6 + 3 * i), L(7 + 3 * i), L(8 + 3 * i)
            cm, cr, cp = L(c0 + 1), L(c0 + 2), L(c0 + 3)
            if z in ("Recepción", "Despacho"):
                d = sp("d_rec" if z == "Recepción" else "d_desp")
                vals = [f"={u}{tot_t}*{d}", f"=SUMPRODUCT(({t_e}=0)*{jj}{C_T1}:{jj}{t17})*{d}*{sp('medio')}",
                        f"=SUMPRODUCT(({t_e}=1)*{jj}{C_T1}:{jj}{t17})*{d}", f"=ROUNDUP({cm}{r},0)+ROUNDUP({cr}{r},0)"]
            elif z == "Sustitutos en planta":
                vals = [f"={s_}{tot_t}", f"=SUMPRODUCT(({t_e}=0)*{s_}{C_T1}:{s_}{t17}/{t_c})",
                        f"=SUMPRODUCT(({t_e}=1)*{s_}{C_T1}:{s_}{t17}/{t_c})",
                        f"=ROUNDUP({cm}{r}/{sp('niveles')},0)+ROUNDUP({cr}{r},0)"]
            else:
                vals = ["—", "—", "—", f"=ROUNDUP(SUM({cp}{r1 + 1}:{cp}{r1 + 3})*{sp('vacias')}/{sp('apil')},0)"]
            for n, v in enumerate(vals):
                celda(ws, r, c0 + n, v, "0" if n == 3 else ("#,##0" if n == 0 else "0.0")).alignment = g.DER
        celda(ws, r, 13, donde)
    pul = r1 + 5
    extra = [("Total de posiciones", lambda cp, cm, i: f"=SUM({cp}{r1 + 1}:{cp}{r1 + 4})", "0"),
             ("Superficie del pulmón con pasillo (m²)", lambda cp, cm, i: f"={cp}{pul}*{sp('m2pos')}", "0.0"),
             ("Stock de sustitutos 2035 (u)", lambda cp, cm, i: f"=Importados!$P${fi(PLANTAS[i], 'tot')}", "#,##0"),
             ("Jaulas A y B a tener (planta + camiones + vacías)",
              lambda cp, cm, i: (f"=ROUNDUP(SUM({cm}{r1 + 1}:{cm}{r1 + 3}),0)+{L(2 + i)}{R + 14}*{sp('pos')}"
                                 f"+{cp}{r1 + 4}*{sp('apil')}"), "0")]
    for j, (t, fx, fmt) in enumerate(extra):
        rr = pul + j
        celda(ws, rr, 1, t, neg=True)
        for i in range(2):
            c0 = 2 + 4 * i
            celda(ws, rr, c0 + 3, fx(L(c0 + 3), L(c0 + 1), i), fmt, res=True)
    ws.freeze_panes = "B3"
    return {"tot_t": tot_t, "tot_r": tot_r, "pul": pul}


def h_resumen(wb, rf):
    ws = hoja(wb, "Resumen", [52, 16, 16, 60])
    tit(ws, 1, "Recargas — logística de las dos plantas (mes pico de 2035)")
    nota(ws, 2, "Desarrollo y fuentes en el Word de justificación de recargas · § 10")
    cab(ws, 4, ["Concepto", "Comercial", "Industrial", "De dónde sale"])
    tt, tr, pu = rf["tot_t"], rf["tot_r"], rf["pul"]
    filas = [("Recargas del año 2035 (u)", lambda i, p: f"=Importados!$N${fi(p, 'tot')}", "#,##0", "Importados"),
             ("Mes pico", lambda i, p: f"='Cálculo'!{L(2 + i)}{C_MES + 1}", None, "Cálculo, punto 1"),
             ("Llegan por día en el mes pico (u)", lambda i, p: f"='Cálculo'!{L(6 + 3 * i)}{tt}", "#,##0",
              "Cálculo, punto 2"),
             ("Sustitutos que salen por día (u)", lambda i, p: f"='Cálculo'!{L(7 + 3 * i)}{tt}", "#,##0",
              "Cálculo, punto 2"),
             ("Stock de sustitutos (u)", lambda i, p: f"=Importados!$P${fi(p, 'tot')}", "#,##0",
              "Dimensionamiento, paso 3"),
             ("Paradas por día", lambda i, p: f"='Cálculo'!{L(5 + 4 * i)}{tr}", "#,##0", "Cálculo, punto 3"),
             ("Viajes por día", lambda i, p: f"='Cálculo'!{L(2 + i)}{C_CAM + 12}", "0", "Cálculo, punto 4"),
             ("Camiones", lambda i, p: f"='Cálculo'!{L(2 + i)}{C_CAM + 14}", "0", "Cálculo, punto 4"),
             ("Ocupación de la flota", lambda i, p: f"='Cálculo'!{L(2 + i)}{C_CAM + 15}", "0%", "Cálculo, punto 4"),
             ("Posiciones del pulmón", lambda i, p: f"='Cálculo'!{L(5 + 4 * i)}{pu}", "0", "Cálculo, punto 5"),
             ("Superficie del pulmón (m²)", lambda i, p: f"='Cálculo'!{L(5 + 4 * i)}{pu + 1}", "0.0",
              "Cálculo, punto 5"),
             ("Jaulas A y B a tener", lambda i, p: f"='Cálculo'!{L(5 + 4 * i)}{pu + 3}", "0", "Cálculo, punto 5")]
    for j, (t, fx, fmt, de) in enumerate(filas):
        r = 5 + j
        celda(ws, r, 1, t, neg=True)
        for i, p in enumerate(PLANTAS):
            celda(ws, r, 2 + i, fx(i, p), fmt, res=True).alignment = g.DER
        celda(ws, r, 4, de)


def generar():
    datos = {p: datos_dim(p) for p in PLANTAS}
    wb = Workbook()
    wb.remove(wb.active)
    g.h_links(wb, LINKS)
    h_supuestos(wb)
    h_importados(wb, datos)
    rf = h_calculo(wb)
    h_resumen(wb, rf)
    wb.move_sheet("Resumen", offset=-(len(wb.sheetnames) - 2))
    wb.active = wb.sheetnames.index("Resumen")
    ruta = os.path.join(g.SALIDA, "Recargas.xlsx")
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    print(generar())
