"""Planilla de zonas y equipos del layout v10 (código, ubicación, superficie y fuente de cada uno)."""

import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from . import modelo as M
from .laminas import NOMBRE_CAT, resumen

CAT = dict(NOMBRE_CAT)


def generar(ruta):
    wb = Workbook()
    neg, nor = Font(name="Arial", size=9, bold=True), Font(name="Arial", size=9)
    gris, gris2 = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="F2F2F2")
    fino = Side(style="thin", color="808080")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    wrap = Alignment(wrap_text=True, vertical="top")

    def hoja(ws, titulo, notas, cols, filas, anchos):
        ws.append([titulo]); ws["A1"].font = Font(name="Arial", size=11, bold=True)
        for n in notas:
            ws.append([n]); ws.cell(ws.max_row, 1).font = Font(name="Arial", size=9, italic=True)
        ws.append([])
        ws.append(cols)
        h = ws.max_row
        for c in range(1, len(cols) + 1):
            x = ws.cell(h, c); x.font = neg; x.fill = gris; x.border = borde; x.alignment = wrap
        for f in filas:
            ws.append(list(f))
            r = ws.max_row
            tot = isinstance(f[0], str) and f[0].startswith("TOTAL")
            for c in range(1, len(cols) + 1):
                x = ws.cell(r, c); x.font = neg if tot else nor; x.border = borde; x.alignment = wrap
                if tot:
                    x.fill = gris2
                if isinstance(x.value, float):
                    x.number_format = "0.00"
        for i, w in enumerate(anchos):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = ws.cell(h + 1, 1)

    tot, adm = resumen()
    filas = [(CAT.get(k, k), round(tot.get(k, 0.0), 1), round(100 * tot.get(k, 0.0) / (M.LX * M.LY), 1))
             for k in ["MP", "MAN", "PINT", "TERM", "CAR", "PT", "AUX", "LIB", "CIRC"]]
    filas.append(("TOTAL NAVE (%.1f × %.1f m)" % (M.LX, M.LY), round(M.LX * M.LY, 1), 100.0))
    filas.append(("Anexo de administración y servicios al personal", round(adm, 1), "-"))
    hoja(wb.active, "Layout v10: superficies por destino",
         ["Cada m² de la nave está en una zona con código (verificado: sin superposiciones ni huecos sin nombre)."],
         ["Destino", "m²", "% de la nave"], filas, [60, 12, 14])
    wb.active.title = "Resumen"

    def zf(z):
        r = z.rect
        return (z.cod, z.nombre, CAT.get(z.cat if z.tipo != "pasillo" else "CIRC", z.cat), z.tipo, r.x0, r.y0, r.x1, r.y1,
                round(r.w, 2), round(r.h, 2), round(r.area, 2), z.req if z.req else "-", z.fuente)
    cols = ["Código", "Zona", "Destino", "Tipo", "x0", "y0", "x1", "y1", "Ancho (m)", "Largo (m)", "m² del plano",
            "m² de la fuente", "Fuente"]
    an = [10, 52, 22, 9, 7, 7, 7, 7, 9, 9, 10, 10, 90]
    hoja(wb.create_sheet("Zonas de la nave"), "Zonas de la nave",
         ["Origen en la esquina SO interior de la nave; x al este, y al norte (m)."], cols, [zf(z) for z in M.ZONAS], an)
    hoja(wb.create_sheet("Anexo administración"), "Locales del anexo de administración y servicios", [], cols,
         [zf(z) for z in M.ADMIN], an)
    filas = [(e.cod, e.nombre, e.zona, e.tipo, e.rect.x0, e.rect.y0, e.rect.x1, e.rect.y1, round(e.rect.w, 2),
              round(e.rect.h, 2), e.frente, e.op, "Sí" if e.ae else "", e.fuente) for e in M.EQUIPOS]
    hoja(wb.create_sheet("Equipos"), "Equipos, pulmones y almacenamiento",
         ["Frente: lado del operario (0,95 m libres verificados). AE: necesita el autoelevador (toca un carril)."],
         ["Código", "Equipo", "Zona", "Símbolo", "x0", "y0", "x1", "y1", "Medida x (m)", "Medida y (m)", "Frente",
          "Operarios", "AE", "Fuente de la medida"], filas, [11, 60, 9, 14, 7, 7, 7, 7, 10, 10, 7, 9, 5, 70])
    filas = []
    for k, (mat, cil) in M.PICO_MES.items():
        filas.append((k, mat, cil, round(mat + cil, 2), round((mat + cil) / M.SEMANAS_MES, 1),
                      round((mat + cil) / M.DIAS_MES, 2)))
    tm = sum(v[0] for v in M.PICO_MES.values())
    tc = sum(v[1] for v in M.PICO_MES.values())
    filas.append(("TOTAL", round(tm, 2), round(tc, 2), round(tm + tc, 2), round((tm + tc) / M.SEMANAS_MES, 1),
                  round((tm + tc) / M.DIAS_MES, 2)))
    hoja(wb.create_sheet("Demanda de carros"), "Demanda pico 2035 de carros (X-CAR, hoja Plan mensual)",
         ["Mes pico: marzo a mayo, julio y agosto. 20 días hábiles por mes = 4 semanas (X-CAR, parametros_carros)."],
         ["Modelo", "Matafuegos por mes", "Cilindros por mes", "Total por mes", "Por semana", "Por día"], filas,
         [12, 18, 18, 14, 12, 10])
    filas = [(k[0], k[1], k[2], k[3], k[4], k[5]) for k in M.STOCK_CARROS]
    hoja(wb.create_sheet("Lugares de carros"), "Lugares de carros rodantes en planta",
         ["Cada lugar: carro parado sobre sus ruedas y su pata, 0,68 × 0,75 m, en filas enfrentadas de a dos."],
         ["Zona", "Puesto", "Qué guarda", "Criterio", "Cuenta (pico 2035)", "Lugares dibujados"], filas,
         [8, 30, 45, 70, 22, 10])
    filas = [(p.cod, {"N": "Norte", "S": "Sur", "E": "Este", "O": "Oeste"}[p.muro], p.a, p.b, round(p.b - p.a, 2), p.alto,
              p.tipo, p.uso) for p in M.PUERTAS]
    hoja(wb.create_sheet("Puertas"), "Puertas de la nave", ["a y b: coordenadas sobre el muro (x en N y S, y en E y O)."],
         ["Código", "Muro", "a", "b", "Ancho (m)", "Alto (m)", "Tipo", "Uso"], filas, [8, 8, 7, 7, 9, 8, 12, 80])
    wb.save(ruta)


if __name__ == "__main__":
    generar(sys.argv[1])
