"""Arma simple/index.html (trazabilidad interna FLAMA, sin QR) desde plantilla.html + app.js con los datos de los libros
de referencias/extra_fran:
  MP-FLAMA2035.xlsx        Stock_MP (las 125 MP: seguridad, consumo mientras llega, stock máximo, bulto, almacén) y
                           Almacenes (lugares y cómo se dispone cada almacén)
  PTPINTADO-FLAMA2035.xlsx Stock_PT (stock máximo y pallets de cada producto), Almacen_PT (rack, piso, andén) y
                           Despacho (días fijos de cada vehículo)
  Balanceo-FLAMA2035.xlsx  Carga (puestos de manuales, carros y recargas, con su número)
y de la app de gestión interna FLAMA (APP_SISTEMA_DE_GESTION_INTERNA_FLAMA): el plano rev10 (layout_rev10.json, copia de
public/data/layout.json) y los logos (img/), para que las dos apps se vean y se ubiquen igual.
Los registros (conteos, pedidos, formularios, pallets) no se inventan: arrancan vacíos y los carga la app.
Uso: python3 build.py"""
import base64
import json
import math
import os

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
REF = os.path.join(os.path.dirname(os.path.dirname(AQUI)), "referencias", "extra_fran")


def libro(nombre):
    return load_workbook(os.path.join(REF, nombre), data_only=True)


def r1(x, d=1):
    return round(float(x), d) if isinstance(x, (int, float)) else None


def stock_mp(wb):
    """Una fila por MP (Stock_MP filas 6 a 130). Punto de pedido = seguridad + consumo mientras llega (L + K): el
    stock que tiene que quedar cuando sale el pedido para que la entrega llegue antes de tocar la seguridad."""
    out = []
    for i, r in enumerate(wb["Stock_MP"].iter_rows(min_row=6, max_row=130, values_only=True), 6):
        if not r[1]:
            continue
        ss, k, mx, cont = float(r[11] or 0), float(r[10] or 0), float(r[18] or 0), float(r[13] or 1)
        out.append(dict(ref=r[0], cod=r[1], mat=r[2], um=r[3], prov=r[4], pico=r1(r[7], 2), pmin=r[8], pmax=r[9],
                        ss=r1(ss), k=r1(k), pp=r1(ss + k), max=r1(mx), bulto=r[12], cont=cont, cada=r[14],
                        alm=r[17], lug=(math.ceil(float(r[20])) if isinstance(r[20], (int, float)) else r[20]),
                        guarda=r[21], linea=r[22], fila=i))
    return out


def almacenes(wb):
    out = []
    for r in wb["Almacenes"].iter_rows(min_row=19, max_row=31, values_only=True):
        if r[0]:
            out.append(dict(alm=r[0], que=r[1], est=r[2], lug=r[4], m2=r1(r[6]), como=r[8]))
    return out


def stock_pt(wb):
    ws = wb["Stock_PT"]
    out = []
    for r in ws.iter_rows(min_row=63, max_row=78, values_only=True):
        out.append(dict(prod=r[0], ventas=r1(r[1], 0), dia=r1(r[5]), ss=r[7], max=r1(r[9], 0), porp=r[10],
                        pallets=r[11]))
    a = {r[0]: r[1] for r in wb["Almacen_PT"].iter_rows(min_row=7, max_row=31, values_only=True) if r[0]}
    carros = {r[0]: r[1] for r in wb["Almacen_PT"].iter_rows(min_row=22, max_row=25, values_only=True) if r[0]}
    desp = [dict(veh=r[0], dias=r[3], salidas=r[2]) for r in wb["Despacho"].iter_rows(min_row=25, max_row=28,
                                                                                         values_only=True) if r[0]]
    return out, dict(modulos=a["Módulos de rack"], lugares=a["Lugares que da el rack"], anden=a["Lugares de pallet en el andén"],
                     pesados=a["Pallets que van solo en el piso (más de la mitad de la carga del nivel)"]), carros, desp


def puestos(wb):
    """Puestos del balanceo (hoja Carga): manuales filas 34 a 59, carros 74 a 91, recargas 104 a 126."""
    ws = wb["Carga"]
    pz = lambda a, b: [dict(n=str(r[0]), nom=r[1]) for r in ws.iter_rows(min_row=a, max_row=b, values_only=True)
                       if r[0] is not None and r[1]]
    return dict(man=pz(34, 59), car=pz(74, 91), rec=pz(104, 126))


# Foto de cada cámara (fotos/<archivo>.jpg o .png): si existe se incrusta achicada a 960 px; si no, la app dibuja el
# esquema del almacén y deja subir la foto desde la pantalla de cámaras.
FOTOS = {"Chapa (pilas en el piso)": "chapa", "Flejes (rack + volteador)": "flejes", "Cantilever de caños y barras": "cantilever",
         "Depósito de polvo químico": "polvo", "Embalaje": "embalaje", "Identificación": "identificacion",
         "Insumos de montaje (manuales)": "insumos_manuales", "Insumos de recargas": "insumos_recargas",
         "Insumos de soldadura y superficie": "soldadura", "Jaula de gases": "gases",
         "Rack de casquetes y cuplas (carros)": "casquetes", "Rack de insumos de montaje (carros)": "insumos_carros",
         "Almacén de producto terminado": "pt"}


def fotos():
    import io
    from PIL import Image
    out = {}
    for alm, nombre in FOTOS.items():
        for ext in (".jpg", ".jpeg", ".png"):
            ruta = os.path.join(AQUI, "fotos", nombre + ext)
            if os.path.exists(ruta):
                im = Image.open(ruta).convert("RGB")
                im.thumbnail((960, 960))
                buf = io.BytesIO()
                im.save(buf, "JPEG", quality=72)
                out[alm] = "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
                break
    return out


def main():
    mp = libro("MP-FLAMA2035.xlsx")
    pt = libro("PTPINTADO-FLAMA2035.xlsx")
    bal = libro("Balanceo-FLAMA2035.xlsx")
    spt, apt, carros, desp = stock_pt(pt)
    layout = json.load(open(os.path.join(AQUI, "layout_rev10.json"), encoding="utf-8"))
    D = dict(mp=stock_mp(mp), layout=layout, fotos=fotos(), alm=almacenes(mp), pt=spt, apt=apt, carros=carros, desp=desp, puestos=puestos(bal),
             fuente=dict(mp="MP-FLAMA2035.xlsx › Stock_MP (filas 6 a 130) y Almacenes (filas 19 a 31)",
                         pt="PTPINTADO-FLAMA2035.xlsx › Stock_PT (filas 63 a 78), Almacen_PT y Despacho (filas 25 a 28)",
                         bal="Balanceo-FLAMA2035.xlsx › Carga (manuales 34-59, carros 74-91, recargas 104-126)"))
    html = open(os.path.join(AQUI, "plantilla.html"), encoding="utf-8").read()
    js = open(os.path.join(AQUI, "app.js"), encoding="utf-8").read()
    for marca, archivo in (("/*LOGO_LLAMA*/", "logo-llama.png"), ("/*LOGO_PALABRA*/", "logo-palabra.png")):
        b64 = base64.b64encode(open(os.path.join(AQUI, "img", archivo), "rb").read()).decode()
        html = html.replace(marca, "data:image/png;base64," + b64)
    html = html.replace("/*DATOS*/", "const D = " + json.dumps(D, ensure_ascii=False) + ";").replace("/*APP*/", js)
    # index.html completo (para Vercel o abrir local) y artifact.html sin envoltorio (lo arma el visor de artifacts)
    with open(os.path.join(AQUI, "artifact.html"), "w", encoding="utf-8") as fh:
        fh.write(html)
    with open(os.path.join(AQUI, "index.html"), "w", encoding="utf-8") as fh:
        fh.write('<!doctype html>\n<html lang="es">\n<head>\n<meta charset="utf-8">\n'
                 '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
                 + html.replace("<style>", "<style>\nbody { margin: 0 }", 1)
                 .replace('<div class="contexto">', '</head>\n<body>\n<div class="contexto">', 1) + "\n</body>\n</html>\n")
    print("ok", len(D["mp"]), "MP ·", len(D["pt"]), "productos de PT ·",
          {k: len(v) for k, v in D["puestos"].items()}, "puestos")


if __name__ == "__main__":
    main()
