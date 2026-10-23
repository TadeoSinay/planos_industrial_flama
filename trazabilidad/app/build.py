"""Arma app/index.html (sistema de gestión FLAMA) desde plantilla.html + js/*.js con los datos reales del proyecto:
Plan mensual de Aprovechamiento_Manuales y _Carros (órdenes y series de cada mes del año hasta el mes en curso), plan de
recargas de las dos plantas, política de stock del MRP y modelos del BOM. Lo que no sale de esos archivos (estado de
cada lote, clientes, stock físico, reclamos) lo simula la app con semilla fija y lo marca como dato de prueba.
Uso: python3 build.py [año] [mes 1-12]  (por defecto, octubre de 2026)."""
import json
import os
import sys

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(os.path.dirname(os.path.dirname(AQUI)), "salida")
PLANOS = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(AQUI))), "planos_flama-", "salida")
URL = "https://trazabilidad-flama.vercel.app/m/"
JS = ["base.js", "datos.js", "shell.js", "produccion.js", "logistica.js", "servicio.js", "tablero.js", "calidad.js", "gestion.js", "app.js"]
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
# tamaño de lote = lo que entra en el carro de transporte interno del modelo (criterio, lo fija PCP)
LOTE = {"ABC 1 kg": 60, "ABC 2,5 kg": 48, "ABC 5 kg": 48, "ABC 10 kg": 24, "Carro ABC 25 kg": 4, "Carro ABC 50 kg": 4,
        "Carro ABC 70 kg": 2, "Carro ABC 100 kg": 2}


def plan(archivo, anio):
    """Filas del año en la hoja «Plan mensual»: [(producto, tipo, [12 valores])] y el rango leído."""
    ws = load_workbook(os.path.join(SALIDA, archivo), read_only=True, data_only=True)["Plan mensual"]
    filas, desde, hasta = [], None, None
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        if r[0] == anio and r[1] not in (None, "Producto"):
            filas.append((r[1], r[2], [float(x or 0) for x in r[3:15]]))
            desde = desde or i
            hasta = i
    return filas, f"A{desde}:P{hasta}"


def plan_recargas(archivo, anio, mes):
    """Recipientes por tipo del mes en «Plan mensual» de Aprovechamiento_Recargas_*: bloque del año, filas hasta el total."""
    ws = load_workbook(os.path.join(SALIDA, archivo), read_only=True, data_only=True)["Plan mensual"]
    out, dentro = [], False
    for r in ws.iter_rows(values_only=True):
        if r[0] == anio and r[1] == "Agente":
            dentro = True
            continue
        if dentro:
            if r[0] is None or str(r[0]).startswith("Total"):
                break
            out.append({"tipo": r[0], "agente": r[1], "cant": round(float(r[2 + mes] or 0), 1)})
    return out


def insumos():
    """POLITICA_STOCK del MRP (2035): código, material, UM, proveedor, plazo y punto de pedido; y el consumo 2026 del
    CONSOLIDADO para el ritmo actual."""
    wb = load_workbook(os.path.join(PLANOS, "mrp", "FLAMA_MRP.xlsx"), read_only=True, data_only=True)
    c26 = {}
    for r in wb["CONSOLIDADO"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            c26[r[0]] = float(r[7] or 0)
    out = []
    for r in wb["POLITICA_STOCK"].iter_rows(min_row=5, values_only=True):
        if not r[1] or not r[2]:
            continue
        out.append({"cod": r[1], "mat": r[2], "um": r[3], "prov": r[4], "dia35": round(float(r[6] or 0), 3),
                    "pmin": r[8], "pmax": r[9], "seg": round(float(r[12] or 0), 1), "pdp": round(float(r[13] or 0), 1),
                    "anio26": round(c26.get(r[1], 0), 1)})
    return out


def modelos():
    """Hoja MODELOS del BOM: carga nominal y banda de carga de los ABC propios."""
    ws = load_workbook(os.path.join(PLANOS, "bom", "FLAMA_BOM.xlsx"), read_only=True, data_only=True)["MODELOS"]
    return [{"cod": r[0], "nombre": r[1], "iram": r[2], "carga": r[7], "um": r[8], "min": r[12], "max": r[13]}
            for r in ws.iter_rows(min_row=3, values_only=True) if r[0] and str(r[0]).startswith("FL_MAT_ABC")]


def lotes_del_anio(man, car, anio, hasta_mes):
    """Órdenes de cada mes desde enero, de matafuegos terminados y de cilindros vendidos vacíos: las series del año son
    correlativas para todos."""
    serie, lotes = 1, []
    for k in range(hasta_mes):
        n = 0
        for prod, tipo, v in man + car:
            cant = round(v[k])
            if tipo not in ("Terminado", "Cilindro") or cant == 0:
                continue
            base = prod.replace("Cilindros ", "")                       # cilindro vendido vacío: mismo recipiente
            tam = LOTE.get(base, LOTE.get("Carro " + base, 48))
            hecho = 0
            while hecho < cant:
                c = min(tam, cant - hecho)
                n += 1
                carro = prod.startswith("Carro") or any(f"ABC {x} kg" in prod for x in (25, 50, 70, 100))
                lotes.append({"op": f"OP {anio % 100:02d}-{k + 1:02d}-{n:02d}", "mes": k + 1, "modelo": prod, "cant": c,
                              "desde": serie, "hasta": serie + c - 1, "carro": carro, "cilindro": tipo == "Cilindro"})
                serie += c
                hecho += c
    return lotes


def main(anio=2026, mes=10):
    man, rango_m = plan("Aprovechamiento_Manuales.xlsx", anio)
    car, rango_c = plan("Aprovechamiento_Carros.xlsx", anio)
    k = mes - 1
    filas = [{"producto": p, "tipo": t, "cant": round(v[k]), "linea": "Carros" if p.startswith("Carro") else "Manuales"}
             for p, t, v in man + car if round(v[k])]
    datos = {"anio": anio, "mes": mes, "mesNombre": MESES[k], "url": URL, "rangos": {"man": rango_m, "car": rango_c},
             "plan": filas, "lotes": lotes_del_anio(man, car, anio, mes),
             "recargas": {"comercial": plan_recargas("Aprovechamiento_Recargas_Planta_Comercial.xlsx", anio, mes),
                          "industrial": plan_recargas("Aprovechamiento_Recargas_Planta_Industrial.xlsx", anio, mes)},
             "insumos": insumos(), "modelos": modelos()}
    with open(os.path.join(AQUI, "plantilla.html"), encoding="utf-8") as fh:
        html = fh.read()
    html = html.replace("/*__DATOS__*/null", json.dumps(datos, ensure_ascii=False, separators=(",", ":")))
    html = html.replace("/*__JS__*/", "\n".join(open(os.path.join(AQUI, "js", f), encoding="utf-8").read() for f in JS))
    with open(os.path.join(AQUI, "index.html"), "w", encoding="utf-8") as fh:
        fh.write(html)
    print("ok", MESES[k], anio, len(datos["lotes"]), "órdenes en el año,", len(datos["insumos"]), "insumos")


if __name__ == "__main__":
    main(*(int(x) for x in sys.argv[1:3]))
