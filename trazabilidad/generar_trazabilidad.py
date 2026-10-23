"""Sistema de trazabilidad FLAMA (salida/Trazabilidad_FLAMA.xlsx): explicación corta de cómo se graba, se lee y qué
muestra el QR (solapa técnica y solapa comercial, las dos con la historia del equipo)."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SALIDA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "salida",
                      "Trazabilidad_FLAMA.xlsx")
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
FT = Font(name="Arial", size=13, bold=True)
FS = Font(name="Arial", size=9, italic=True, color="595959")
FW = Font(name="Arial", size=10, bold=True, color="FFFFFF")
NEGRO = PatternFill("solid", fgColor="000000")
GRIS = PatternFill("solid", fgColor="D9D9D9")
TEC = PatternFill("solid", fgColor="DDEBF7")
COM = PatternFill("solid", fgColor="E2EFDA")
EJ = PatternFill("solid", fgColor="FFF2CC")
fino = Side(style="thin", color="A6A6A6")
B = Border(left=fino, right=fino, top=fino, bottom=fino)
W = Alignment(wrap_text=True, vertical="top")
C = Alignment(wrap_text=True, vertical="center", horizontal="center")


def hoja(wb, nombre, anchos, titulo, nota):
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False
    for i, a in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = a
    ws["A1"] = titulo
    ws["A1"].font = FT
    ws["A2"] = nota
    ws["A2"].font = FS
    return ws


def tabla(ws, r, cab, filas, relleno=None, alto=None):
    for j, t in enumerate(cab, 1):
        x = ws.cell(r, j, t)
        x.font, x.fill, x.alignment, x.border = FW, NEGRO, C, B
    for f in filas:
        r += 1
        for j, v in enumerate(f, 1):
            x = ws.cell(r, j, v)
            x.font, x.alignment, x.border = (FB if j == 1 else F), W, B
            if relleno:
                x.fill = relleno
        if alto:
            ws.row_dimensions[r].height = alto
    return r + 2


def seccion(ws, r, t, n, relleno=GRIS):
    for j in range(1, n + 1):
        ws.cell(r, j).fill = relleno
    ws.cell(r, 1, t).font = FB
    return r + 1


wb = Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------------- 1. Resumen
ws = hoja(wb, "Resumen", [26, 95], "TRAZABILIDAD FLAMA — cómo funciona en una hoja",
          "Un número por recipiente desde el corte hasta la baja. Lo grabado cumple la norma; el código lleva a la "
          "historia.")
r = tabla(ws, 4, ["Tema", "Cómo se hace"], [
    ("Qué se graba", "En el cuerpo, franja baja (pollera): texto legible «FLAMA S.A. · N° 26-000123 · 26» + cuño DPS "
                     "15 × 7 (IRAM 3523 / 3550 5.1; Res. 349/07 anexo IV; carros: + presión de ensayo) y un "
                     "DataMatrix ECC 200 con el mismo n° de serie (GTIN del modelo + serie, formato GS1)."),
    ("Con qué", "Micropercusión Pannier ADP120160D (la que ya tiene manuales): graba texto, logos y códigos 2D en una "
                "ventana de 120 × 160 mm, marcado profundo pensado para seguir legible después del granallado y la "
                "pintura (ficha del fabricante). Se muda del puesto 6 («numerado de cúpula») a la salida de la "
                "guillotina 2 y graba la planchuela en plano. Carros: su marcadora del puesto 2.1 (confirmar con el "
                "proveedor que haga DataMatrix)."),
    ("Cuándo", "Al cortar: ahí nace el n° de serie y queda atado a la hoja y al lote de chapa. Después del grabado van "
               "el cilindrado, la soldadura y la prueba hidráulica (la PH valida la pared ya marcada)."),
    ("Cómo se lee en planta", "Lector DPM (para códigos grabados) fijo en PH, pintura, carga, etiquetado y despacho; "
                              "uno de mano en recargas. Cada lectura deja fecha, puesto, operario y resultado."),
    ("Qué ve el cliente", "La etiqueta lleva un QR impreso con el mismo n° de serie (GS1 Digital Link): se lee con el "
                          "celular y abre la ficha del equipo con dos solapas, TÉCNICA y COMERCIAL, las dos con la "
                          "historia."),
    ("Si se pierde la etiqueta", "El recargador repinta (IRAM 3517-2 9.10) y la etiqueta se va, pero el grabado queda: "
                                 "lo lee con el lector DPM y reimprime el QR. El número nunca se pierde."),
    ("Quién ve qué", "Público (QR): solapa comercial y técnica resumida. Recargador habilitado, inspector DPS/AGC "
                     "(con usuario): solapa técnica completa y carga de servicios. FLAMA interno: además lotes, "
                     "operarios y parámetros de máquina."),
    ("Qué norma lo pide", "IRAM 3517-2:2020: sistema informatizado de seguimiento y trazabilidad del recargador, "
                          "registro de 6 años por extintor (9.4.16) y n° de serie en la etiqueta de servicio (9.4.14). "
                          "El DataMatrix y el QR se suman al grabado legible, no lo reemplazan."),
])
ws.cell(r, 1, "Lo que falta confirmar").font = FB
ws.cell(r + 1, 1, "• Profundidad del grabado: prueba con piezas reales (grabar → cilindrar → granallar → pintar → "
                  "leer). • DataMatrix de la marcadora de carros. • Con DPS: el cuerpo bajo como «pollera» en 1 kg "
                  "y carros.").font = F
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
ws.cell(r + 1, 1).alignment = W
ws.row_dimensions[r + 1].height = 30

# ---------------------------------------------------------------- 2. Flujo
ws = hoja(wb, "Flujo", [5, 22, 40, 22, 46, 10],
          "FLUJO — qué pasa en cada etapa y qué queda registrado",
          "T = lo muestra la solapa técnica · C = la comercial · I = solo interno FLAMA.")
r = tabla(ws, 4, ["N°", "Etapa", "Qué se hace", "Medio", "Qué queda registrado", "Solapa"], [
    (1, "Recepción de chapa", "Se carga la hoja con su certificado", "Remito y certificado",
     "Proveedor, colada, espesor, certificado de calidad", "I"),
    (2, "Corte y grabado", "Guillotina 2 + Pannier: se graba la planchuela", "Pannier ADP120160D",
     "N° de serie, modelo, fecha, hoja y lote de chapa", "T"),
    (3, "Conformado y soldadura", "Cilindrado, soldaduras, encastre, bordón", "Lector fijo",
     "Fecha, soldador o máquina, lote de alambre y gas", "I"),
    (4, "Prueba hidráulica", "Ensayo al 100 %", "Lector fijo + banco PH",
     "Presión, tiempo, resultado; si falla, serie anulado", "T"),
    (5, "Granallado y pintura", "Terminación", "Lector fijo", "Fecha, lote de pintura", "T"),
    (6, "Carga y presurizado", "Polvo, válvula, N₂, ensayo de pérdidas", "Lector fijo + balanza",
     "Lote de polvo, masa de carga, presión, n° de válvula y manómetro, resultado de pérdidas", "T"),
    (7, "Etiquetado", "Etiqueta con QR del mismo serie; sello IRAM, oblea PBA", "Impresora + lector",
     "N° de estampilla IRAM, oblea, precinto", "T"),
    (8, "Embalaje y despacho", "Pallet y remito", "Lector de mano", "Pallet, remito, cliente, destino (CABA/PBA)",
     "C"),
    (9, "Instalación", "El cliente o el instalador escanea el QR", "Celular",
     "Propiedad y ubicación (opcional, la carga el cliente)", "C"),
    (10, "Recarga y mantenimiento", "Se escanea el grabado o el QR y se carga el servicio",
     "Lector de mano o celular", "Fecha, tipo de servicio, PH si hubo, marbete, tarjeta DPS/AGC, prestador", "T y C"),
    (11, "Baja", "Inutilización (IRAM 3517-2 9.12)", "Lector de mano", "Fecha, motivo, acta (anexo G)", "T y C"),
])
r = seccion(ws, r, "Casos especiales", 6)
r = tabla(ws, r, ["", "Caso", "Cómo entra al sistema", "", "", ""], [
    ("", "Cilindro vendido vacío", "Igual que el fabricado hasta la pintura (etapas 1 a 5) y despacho. Si después "
                                   "lo carga otro, el grabado sigue identificándolo", "", "", ""),
    ("", "Revendido", "No se toca su marcado. En la recepción se carga el n° grabado por su fabricante, el lote y el "
                      "remito, y se le pega la etiqueta con QR FLAMA", "", "", ""),
    ("", "Recarga de otra marca", "Se carga su n° grabado y se le pega una etiqueta de servicio con QR", "", "", ""),
    ("", "Sustituto", "Mismo QR; la solapa marca «sustituto habilitado» y a qué propiedad cubre", "", "", ""),
])
for rr in range(r - 6, r):
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)

# ---------------------------------------------------------------- 3. Solapas del QR
ws = hoja(wb, "QR - solapas", [30, 38, 4, 30, 38],
          "QUÉ MUESTRA EL QR — solapa TÉCNICA y solapa COMERCIAL (las dos con la historia)",
          "Ejemplo inventado de un ABC 5 kg (celdas amarillas = datos de ejemplo).")
r = 4
for col, (titulo, relleno) in ((1, ("SOLAPA TÉCNICA", TEC)), (4, ("SOLAPA COMERCIAL", COM))):
    x = ws.cell(r, col, titulo)
    x.font, x.fill = FW, NEGRO
    ws.cell(r, col + 1).fill = NEGRO
tec = [("Identificación", None),
       ("N° de serie", "26-000123"), ("Modelo", "ABC 5 kg — IRAM 3523"), ("Fabricante", "FLAMA S.A."),
       ("Fabricado", "03/2026"), ("Grabado en el cuerpo", "FLAMA S.A. · N° 26-000123 · 26 + DPS"),
       ("Fabricación", None),
       ("Lote de chapa", "LAF 1,6 · colada 4521 (Pradecon)"), ("Prueba hidráulica de fábrica", "03/2026 · aprobada"),
       ("Carga", "Polvo ABC 60 · lote 2603-07 · 5,15 kg"), ("Presión", "1,4 MPa (N₂)"),
       ("Válvula / manómetro", "HZ M-30 n° 88412 / 1,4 MPa"), ("Sello IRAM / oblea PBA", "A 26 1111053 / 00457"),
       ("Vencimientos", None),
       ("Próxima recarga", "03/2027"), ("Próxima PH", "03/2031 (cada 5 años, IRAM 3517-2 anexo E)"),
       ("Fin de vida útil", "2046 (20 años, IRAM 3517-2 anexo E)")]
com = [("Identificación", None),
       ("Equipo", "Matafuego ABC 5 kg"), ("Apto para", "Fuegos A, B y C"), ("N° de serie", "26-000123"),
       ("Estado", "VIGENTE"), ("Dónde está", "Local 3, depósito (lo carga el cliente)"),
       ("Garantía y certificados", None),
       ("Sello IRAM", "Sí — ver certificado"), ("Habilitación", "DPS (PBA) / AGC (CABA)"),
       ("Garantía", "Según condiciones de venta"),
       ("Próximos pasos", None),
       ("Próxima recarga", "03/2027 — botón «pedir recarga»"), ("Próxima PH", "03/2031"),
       ("Reemplazo", "2046"),
       ("Contacto", None),
       ("Servicio técnico", "FLAMA S.A. — teléfono / WhatsApp / web"), ("Comprado en", "Distribuidor X, 04/2026")]
for col, lista, relleno in ((1, tec, TEC), (4, com, COM)):
    rr = r + 1
    for k, v in lista:
        if v is None:
            for j in (col, col + 1):
                ws.cell(rr, j).fill = relleno
            ws.cell(rr, col, k).font = FB
        else:
            a = ws.cell(rr, col, k)
            b = ws.cell(rr, col + 1, v)
            a.font, b.font = F, F
            a.border, b.border = B, B
            b.fill = EJ
            a.alignment = b.alignment = W
        rr += 1
r = rr + 1
for col, titulo, relleno in ((1, "HISTORIA — vista técnica", TEC), (4, "HISTORIA — vista comercial", COM)):
    for j in (col, col + 1):
        ws.cell(r, j).fill = relleno
    ws.cell(r, col, titulo).font = FB
hist_t = [("03/2026", "Fabricado · PH aprobada · carga 5,15 kg · FLAMA"),
          ("04/2026", "Despachado · remito 0001-00045 · PBA"),
          ("03/2027", "Recarga anual · polvo recuperado + reposición · marbete 2027 · tarjeta DPS 12345"),
          ("03/2028", "Recarga anual · repintado (IRAM 3517-2 9.10) · nueva etiqueta QR"),
          ("03/2031", "Recarga + PH aprobada · oblea de vencimiento de PH 2036")]
hist_c = [("03/2026", "Fabricado y controlado"),
          ("04/2026", "Entregado"),
          ("03/2027", "Recarga anual realizada"),
          ("03/2028", "Recarga anual y pintura nueva"),
          ("03/2031", "Recarga y prueba hidráulica aprobadas")]
for col, lista in ((1, hist_t), (4, hist_c)):
    rr = r + 1
    for f, t in lista:
        a = ws.cell(rr, col, f)
        b = ws.cell(rr, col + 1, t)
        for x in (a, b):
            x.font, x.border, x.alignment, x.fill = F, B, W, EJ
        rr += 1
r = rr + 1
ws.cell(r, 1, "Las dos solapas leen la misma historia; la comercial la cuenta en lenguaje del cliente y sin datos de "
              "proceso. Los lotes, operarios y parámetros de máquina quedan solo en la vista interna de FLAMA.").font = FS
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.cell(r, 1).alignment = W
ws.row_dimensions[r].height = 28

# ---------------------------------------------------------------- 4. Qué información da
ws = hoja(wb, "Qué información da", [30, 70],
          "QUÉ INFORMACIÓN TE DA EL SISTEMA (para FLAMA)",
          "Lo que sale de juntar todas las lecturas por n° de serie.")
tabla(ws, 4, ["Pregunta", "Respuesta del sistema"], [
    ("¿De qué chapa salió este equipo?", "Hoja, colada y certificado del proveedor (llamado a revisión por lote)."),
    ("¿Qué pasó en cada puesto?", "Fecha, operario o máquina y resultado de cada etapa; dónde se rechazó."),
    ("¿Cuántos fallaron la PH y por qué?", "Rechazos por modelo, lote de chapa y soldador."),
    ("¿Dónde está cada equipo vendido?", "Cliente, remito, destino y, si el cliente lo carga, la ubicación."),
    ("¿Cuándo vence cada uno?", "Recarga, PH y vida útil: base para avisar y vender el servicio."),
    ("¿Quién hizo cada recarga?", "Prestador, fecha, tarjeta DPS/AGC; registro de 6 años (IRAM 3517-2 9.4.16)."),
    ("¿Es un equipo FLAMA verdadero?", "El serie existe, coincide grabado y QR y no está dado de baja."),
    ("¿Numeración sin huecos?", "Cada serie está fabricado, anulado (rechazo) o dado de baja: auditable."),
], alto=30)

for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
wb.save(SALIDA)
print("guardado", SALIDA)
