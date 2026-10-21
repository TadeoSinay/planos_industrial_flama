"""Analisis_alternativas_corte_chapa.xlsx: alternativas de compra y corte de chapa para cuerpos, cúpulas y fondos.

Hojas: Links (links y todos los IMPORTRANGE), Resumen (comparación y decisión), Datos (todo lo editable con su
fuente), Demanda, Piezas (medidas del plano y disco por embutido), Corte (cómo sale cada pieza de hoja, bobina, fleje
y bobina madre; dos guillotinas en serie, 1.er corte por el lado corto) y Cálculo (año por año de cada alternativa).
Todo con fórmulas.

Alternativas: A bobina por formato + fleje a medida; B hojas para todo; C bobina madre por espesor con corte
longitudinal propio; D hojas para cuerpos + fleje a medida para discos; Referencia: cuerpos comprados a medida."""

import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), "recargas"))
import generar_dim_recargas as g  # noqa: E402
import datos_corte as DC  # noqa: E402

SALIDA = os.path.join(os.path.dirname(AQUI), "salida", "Analisis_alternativas_corte_chapa.xlsx")
ANIOS = DC.ANIOS
YC = [L(3 + i) for i in range(10)]            # C..L = 2026..2035
FORM = DC.FORM
N = {}                                        # nombre -> celda absoluta
F0, F1, F2, F3 = "#,##0", "#,##0.0", "#,##0.00", "0.000"
FP = "0.0%"
OPC = ["A", "B", "C", "D", "Ref"]
OPC_NOM = {"A": "A: bobina por formato + fleje", "B": "B: hojas para todo", "C": "C: bobina madre propia",
           "D": "D: hojas + fleje", "Ref": "Referencia: cuerpos a medida"}


# ------------------------------------------------------------------ estilo
def c(ws, ref, v, fmt=None, dato=False, neg=False, al=None):
    col = "".join(ch for ch in ref if ch.isalpha())
    row = int("".join(ch for ch in ref if ch.isdigit()))
    x = g.celda(ws, row, _ci(col), v, fmt, neg=neg)
    if dato:
        x.fill = g.GRIS_C
    if al is not None:
        x.alignment = al
    return x


def _ci(col):
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def fx(s):
    """Reemplaza {nombre} por la celda de Datos."""
    return s.format(**N)


def titulo(ws, t, sub):
    g.tit(ws, 1, t)
    g.nota(ws, 2, sub)


def seccion(ws, r, t):
    x = ws.cell(r, 1, t)
    x.font = Font(name="Arial", size=11, bold=True)


# ------------------------------------------------------------------ Datos
GEN = [
    ("anio", "Año de referencia", 2035, "año", "0", "Año de diseño del proyecto (Proyecciones)."),
    ("tasa", "Tasa de descuento anual en USD", 0.12, "%", FP, "Evaluación económica del Excel V2 (12 % real)."),
    ("mant", "Mantenimiento anual de equipos", 0.03, "% de la inversión", FP, "Criterio del grupo."),
    ("scrap", "Precio de venta del scrap", 250, "USD/t", F0,
     "Reaquila (mar-2026): acero reciclado $300-1.200/kg con dólar ~$1.480."),
    ("dens", "Densidad del acero SAE 1010", 7.87, "g/cm³", "0.00", "Planilla «Dimensiones de chapa»."),
    ("op_ars", "Costo mensual de un operario con cargas", 1337922, "ARS/mes", F0,
     "Planilla «Costo oportunidad operario-máquina», CCT 260/75, ago-2026."),
    ("h_mes", "Horas por mes", 150, "h/mes", F0, "Excel V2."),
    ("h_anio", "Horas productivas por operario y año", 1800, "h/año", F0, "Excel V2 (150 h × 12)."),
    ("tc", "Tipo de cambio", 1420, "ARS/USD", F0, "Valor fijo del proyecto (criterio del grupo)."),
    ("c_hora", "Costo hora de operario", "={op_ars}/{h_mes}/{tc}", "USD/h", F2, "Cálculo."),
]
CORTE = [
    ("margen", "Margen mínimo de borde a lo ancho de una bobina", 10, "mm", F0,
     "Criterio del grupo: bordes de laminación y tolerancia de ancho."),
    ("despunte", "Despunte de cabeza y cola de bobina", 0.005, "%", FP, "Criterio del grupo."),
    ("pf", "Puente de avance entre discos", 1.5, "× espesor", "0.0",
     "MISUMI, Press Dies Tutorial: 1,5 veces el espesor, mínimo 0,7 mm."),
    ("pmin", "Puente mínimo entre discos", 0.7, "mm", "0.0", "MISUMI."),
    ("pb", "Puente al borde del fleje", 1.5, "× puente de avance", "0.0",
     "MISUMI: 1,5 veces el de avance (la guía del fleje deja juego)."),
    ("mezcla", "¿Mezclar 70 y 100 kg en la misma hoja?", "SI", "SI/NO", None,
     "Proceso de carros: la hoja mixta lleva 1 cuerpo de 100 kg y 3 de 70 kg."),
    ("mix100", "Cuerpos de 100 kg por hoja mixta", 1, "u", F0, "Proceso de carros."),
    ("mix70", "Cuerpos de 70 kg por hoja mixta", 3, "u", F0, "Proceso de carros."),
    ("comb", "¿Combinar 25 y 50 kg en la hoja de 3,2 mm (corte en tres etapas)?", "SI", "SI/NO", None,
     "Hoja Nesting: el aprovechamiento de la chapa de 3,2 mm sube del 85 % al 90 %."),
    ("tresb", "¿Discos en fleje de dos filas al tresbolillo?", "NO", "SI/NO", None,
     "Corte, sección C: necesita matriz de dos discos o alimentador en zigzag (sin cotización)."),
]
MAQ = [
    ("gpm1", "Guillotina 1 (franjas) Ermaksan CNC HGS 2600-6: golpes por minuto", 20, "golpes/min", F0,
     "Aprovechamiento Manuales del grupo. La ficha de la máquina (resale.info, 2018) da 24."),
    ("L1", "Guillotina 1: largo de corte", 2600, "mm", F0, "Ficha Ermaksan CNC HGS 2600-6: 2.600 × 6 mm."),
    ("tope1", "Guillotina 1: recorrido del tope trasero", 1000, "mm", F0,
     "Ficha Ermaksan CNC HGS 2600-6: tope trasero motorizado de 1.000 mm."),
    ("gpm2", "Guillotina 2 (piezas) QC12Y 6 × 2500: golpes por minuto", 20, "golpes/min", F0,
     "Confirmado por el grupo (igual que la HG 6 × 2500 de la ficha Molinari Panther)."),
    ("L2", "Guillotina 2: largo de corte para franjas en paralelo", 2500, "mm", F0,
     "Publicación Alibaba 1600900259147: 6 × 2500, cuchilla de 2.550 mm."),
    ("tope2", "Guillotina 2: recorrido del tope trasero", 800, "mm", F0,
     "Publicación Alibaba 1600900259147: 20 a 800 mm."),
    ("ciclo", "Hojas por ciclo como máximo (franjas de hojas enteras en paralelo en G2)", 4, "hojas", F0,
     "Criterio del grupo (comentarios del corte de 5 y 10 kg): sin franjas sueltas en el pulmón."),
    ("g1_eur", "Guillotina 1 Ermaksan CNC HGS 2600-6 usada (2018)", 26000, "EUR", F0,
     "Anuncio resale.info N.° 165879844 (nov-2024), sin flete ni nacionalización."),
    ("eurusd", "Tipo de cambio euro / dólar", "=1/0.8723", "USD/EUR", "0.000",
     "Pound Sterling Live: USD/EUR 0,8723 al cierre del 21/09/2026."),
    ("linea", "Línea de corte transversal (A y C): golpes por hora", 1200, "golpes/h", F0,
     "Criterio del grupo: 20 golpes por minuto."),
    ("slit_v", "Cortadora longitudinal (C): velocidad efectiva", 40, "m/min", F0,
     "Criterio del grupo (máquinas de 60 a 120 m/min nominales)."),
    ("slit_arm", "Cortadora: armado de cuchillas por bobina madre", 45, "min", F0, "Criterio del grupo."),
    ("slit_op", "Operarios en la cortadora", 2, "personas", F0, "Criterio del grupo."),
    ("tira", "Carga manual de una tira en el balancín (B)", 15, "s/tira", F0, "Criterio del grupo."),
    ("cambio", "Cambio de fleje en el desbobinador de la prensa", 15, "min/fleje", F0, "Criterio del grupo."),
]
STOCK = [
    ("seg", "Stock de seguridad", 2, "semanas", F0, "Criterio del grupo."),
    ("lote", "Lote de compra de hojas y de chapa a medida", 1, "meses", F0, "Criterio del grupo: compra mensual."),
    ("cob", "Cobertura objetivo de un rollo o bobina", 1, "meses", F0, "Criterio del grupo."),
    ("cobmax", "Cobertura máxima de stock por rollo", 3, "meses", F0, "Criterio del grupo."),
    ("pmin_rollo", "Peso mínimo de rollo", 0.5, "t", "0.0", "Criterio del grupo."),
    ("dext", "Diámetro exterior máximo de bobina", 1500, "mm", F0, "Criterio del grupo."),
    ("dint", "Diámetro interior de bobina", 508, "mm", F0, "Estándar de mercado."),
    ("t_mm", "Peso máximo de bobina por mm de ancho", "=PI()/4*({dext}^2-{dint}^2)*{dens}/1000000000", "t/mm",
     "0.00000", "Cálculo con el diámetro exterior máximo."),
    ("des_lin", "Desbobinador de la línea de corte transversal", 10, "t", F0, "A especificar al comprar."),
    ("des_pre", "Desbobinador de fleje de la prensa", 3, "t", F0, "A especificar al comprar."),
    ("des_slit", "Desbobinador de la cortadora longitudinal", 14, "t", F0, "Pradecon: la bobina madre pesa 13-14 t."),
    ("madre_dif", "Precio de la bobina madre sobre la flejada", 0, "%", FP,
     "Pradecon (mail 28/09/2026): no cobra el flejado."),
    ("grua", "¿El puente grúa ya está en el layout por otro motivo?", "NO", "SI/NO", None,
     "Si es SI, no se carga a A ni a C."),
]
COT = [  # id, documento, producto, espesor, formato, moneda, precio/kg, TC, nota
    ("Q1", "Pradecon 173170", "Bobina LAF flejada", 0.9, "Ancho 1.220", "USD", 1.26285, 1545, "Flejado sin cargo"),
    ("Q2", "Pradecon 173170", "Bobina LAF flejada", 1.25, "Ancho 1.220", "USD", 1.25056, 1545, ""),
    ("Q3", "Pradecon 173170", "Bobina LAF flejada", 1.6, "Ancho 1.220 o 1.000", "USD", 1.25056, 1545, ""),
    ("Q4", "Pradecon 173170", "Bobina LAF flejada", 2.0, "Ancho 1.220 o 1.000", "USD", 1.26285, 1545, ""),
    ("Q5", "Pradecon 173170", "Bobina LAC flejada", 3.2, "Ancho 1.245 o 1.500", "USD", 1.11373, 1545, ""),
    ("Q6", "Pradecon 173170", "Hoja LAC", 4.75, "1500 × 3000", "USD", 1.05792, 1545, "Mismo producto que Q11"),
    ("Q7", "Pradecon 173195", "Hoja LAF C16", 1.6, "1000 × 2000", "ARS", 2051.02, 1545, ""),
    ("Q8", "Pradecon 173195", "Hoja LAF C18", 1.25, "1220 × 2440", "ARS", 1993.51, 1545, ""),
    ("Q9", "Pradecon 173195", "Hoja LAF C14", 2.0, "1500 × 3000", "ARS", 1978.22, 1545, ""),
    ("Q10", "Pradecon 173195", "Hoja LAC 1/8", 3.2, "1500 × 3000", "ARS", 1702.49, 1545, ""),
    ("Q11", "Pradecon 173195", "Hoja LAC 3/16", 4.75, "1500 × 3000", "ARS", 1685.56, 1545, ""),
    ("Q12", "Hierros Parrotta 413329", "Hoja doble decapada N.° 20", 0.9, "1220 × 2440", "USD", 1.3561, 1,
     "Neto con 6,84 % de descuento, contado"),
    ("Q13", "Hierros Parrotta 413329", "Hoja doble decapada N.° 20", 0.9, "1000 × 2000", "USD", 1.3993, 1, ""),
    ("Q14", "Hierros Parrotta 413329", "Hoja doble decapada N.° 18 (1,2)", 1.25, "1220 × 2440", "USD", 1.3464, 1, ""),
    ("Q15", "Hierros Parrotta 413329", "Hoja doble decapada N.° 16", 1.6, "1220 × 2440", "USD", 1.3443, 1, ""),
    ("Q16", "Hierros Parrotta 413329", "Hoja doble decapada N.° 14", 2.0, "1220 × 2440", "USD", 1.3605, 1, ""),
    ("Q17", "Hierros Parrotta 413329", "Hoja LAC 1/8", 3.2, "1500 × 3000", "USD", 1.1607, 1, ""),
    ("Q18", "Hierros Parrotta 413329", "Hoja LAC 3/16", 4.75, "1500 × 3000", "USD", 1.1678, 1, ""),
    ("Q19", "Pacheco Chapas PR 42021", "Rollo de fleje LAF", 1.25, "Ancho 200", "USD", 1.186, 1530, ""),
    ("Q20", "Pacheco Chapas PR 42021", "Rollo de fleje LAF", 1.6, "Ancho 200 y 250", "USD", 1.186, 1530, ""),
    ("Q21", "Pacheco Chapas PR 42021", "Rollo de fleje LAF", 2.0, "Ancho 249 y 300", "USD", 1.198, 1530, ""),
    ("Q22", "Pacheco Chapas PR 42021", "Plancha a medida LAF", 1.25, "1220 × 540", "USD", 1.426, 1530, ""),
    ("Q23", "Pacheco Chapas PR 42021", "Plancha a medida LAF", 1.6, "1000 × 475", "USD", 1.453, 1530, ""),
    ("Q24", "Pacheco Chapas PR 42021", "Plancha a medida LAF", 2.0, "1500 × 563", "USD", 1.429, 1530, ""),
    ("Q25", "Pacheco Chapas PR 42021", "Plancha a medida LAC", 3.2, "1500 × 2584 / 1988", "USD", 1.256, 1530, ""),
    ("Q26", "Pacheco Chapas PR 42024", "Plancha a medida LAC", 4.75, "1245 × 1360 / 1800", "USD", 1.236, 1530, ""),
]
ESP = [0.9, 1.25, 1.6, 2.0, 3.2, 4.75]
# precio por espesor y forma: bobina (A), hoja (B y D), fleje a medida, a medida (Referencia); madre = bobina
PRECIO = {0.9: ("Q1", "Q12", "Q1", "Q12"), 1.25: ("Q2", "Q8", "Q19", "Q22"), 1.6: ("Q3", "Q7", "Q20", "Q23"),
          2.0: ("Q4", "Q9", "Q21", "Q24"), 3.2: ("Q5", "Q10", "Q5", "Q25"), 4.75: ("Q5", "Q11", "Q5", "Q26")}
P_NOTA = {0.9: "A medida: sin cotización, se toma la hoja Q12.", 4.75: "Bobina: sin cotización, se toma la de 3,2 (Q5).",
          3.2: "Fleje: la bobina flejada Q5.", 1.25: "", 1.6: "", 2.0: ""}
EQ = [  # equipo, precio USD, cantidades A, B, C, D, Ref ("grua" = según Datos), fuente
    ("Línea de corte transversal 1-6 mm × 1.600 (desbobinador, aplanadora, alimentador, cizalla)", 80000,
     (1, 0, 1, 0, 0), "ESTIMADO del grupo: Hener da USD 45.000 para 0,2-3 mm × 1.250; para 4,75 mm, más pesada."),
    ("Guillotina 1 (franjas) Ermaksan CNC HGS 2600-6, tope 1.000 mm", "={g1_eur}*{eurusd}", (1, 1, 0, 1, 0),
     "Anuncio resale.info (usada, 2018), en EUR pasado a USD (sección 3)."),
    ("Guillotina 2 (piezas) QC12Y 6 × 2500, tope 800 mm", 10700, (0, 1, 0, 1, 0),
     "Publicación Alibaba 1600900259147: FOB China, 1 a 2 unidades."),
    ("Cortadora longitudinal 0,3-3,2 × 1.600", 60000, (0, 0, 1, 0, 0),
     "ESTIMADO del grupo: Cangzhou Zunsheng USD 28.600-59.620 FOB."),
    ("Puente grúa monoviga 10 t con carrilera", 30000, ("grua", 0, "grua", 0, 0),
     "Dongqi Crane: USD 10.000-50.000 (10 t básico)."),
    ("Pinza en C portabobinas y cunas", 6000, (1, 0, 1, 0, 0), "ESTIMADO del grupo."),
    ("Desbobinador, enderezador y alimentador servo de la prensa de discos", 12000, (1, 0, 1, 1, 1),
     "ESTIMADO del grupo (en cotización)."),
    ("Estantería porta-flejes", 3000, (1, 0, 1, 1, 1), "ESTIMADO del grupo."),
    ("Autoelevador eléctrico 2,5 t", 20000, (0, 1, 0, 1, 0), "Mercado Libre GBA, nuevos 2026: USD 13.700-26.700."),
    ("Mesa elevadora de tijera 3 t", 7000, (0, 1, 0, 1, 0), "ESTIMADO del grupo."),
    ("Mesa de bolas entre la mesa elevadora y la guillotina", 1500, (0, 1, 0, 1, 0), "ESTIMADO del grupo."),
    ("Racks por espesor y carros porta-desarrollos (pulmón entre guillotinas)", 2000, (0, 1, 0, 1, 0),
     "ESTIMADO del grupo."),
    ("Mesa de alimentación de tiras al balancín", 1500, (0, 1, 0, 0, 0), "ESTIMADO del grupo."),
]


def h_datos(wb):
    ws = g.hoja(wb, "Datos", [56, 14, 16, 12, 12, 12, 12, 20, 70])
    titulo(ws, "Datos", "Gris claro: dato editable. Todo lo demás del libro sale de acá, de Demanda y de Piezas.")
    r = 4
    for nombre, filas in (("1. Generales", GEN), ("2. Corte y puentes", CORTE), ("3. Máquinas", MAQ),
                          ("4. Compras y stock", STOCK)):
        seccion(ws, r, nombre)
        g.cab(ws, r + 1, ["Concepto", "Valor", "Unidad"])
        ws.merge_cells(start_row=r + 1, start_column=4, end_row=r + 1, end_column=9)
        g.cab(ws, r + 1, ["Fuente o criterio"], c0=4)
        r += 2
        for k, a, v, u, fmt, fte in filas:
            N[k] = f"Datos!$B${r}"
            c(ws, f"A{r}", a)
            esf = isinstance(v, str) and v.startswith("=")
            c(ws, f"B{r}", fx(v) if esf else v, fmt, dato=not esf)
            c(ws, f"C{r}", u)
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)
            c(ws, f"D{r}", fte)
            r += 1
        r += 1
    # cotizaciones
    seccion(ws, r, "5. Cotizaciones de acero recibidas (sin IVA)")
    g.cab(ws, r + 1, ["Cotización", "Id", "Espesor (mm)", "Moneda", "Precio por kg", "TC", "USD/kg", "Formato",
                      "Producto y nota"])
    r += 2
    N["cot"] = {}
    for q, doc, prod, e, form, mon, p, tc, nota in COT:
        c(ws, f"A{r}", doc)
        c(ws, f"B{r}", q)
        c(ws, f"C{r}", e, "0.00")
        c(ws, f"D{r}", mon)
        c(ws, f"E{r}", p, "0.0000" if mon == "USD" else F2, dato=True)
        c(ws, f"F{r}", tc, F0, dato=True)
        c(ws, f"G{r}", f'=IF(D{r}="USD",E{r},E{r}/F{r})', "0.0000")
        c(ws, f"H{r}", form)
        c(ws, f"I{r}", prod + (f". {nota}." if nota else "."))
        N["cot"][q] = f"Datos!$G${r}"
        r += 1
    r += 1
    seccion(ws, r, "6. Precio del acero por espesor y forma de compra (USD/kg)")
    g.cab(ws, r + 1, ["Espesor (mm)", "Bobina (A)", "Hoja (B y D)", "Fleje a medida", "Bobina madre (C)",
                      "A medida (Ref.)", "", "", "Cotización usada"])
    r += 2
    N["p0"] = r
    for e in ESP:
        b, h, f, m = PRECIO[e]
        c(ws, f"A{r}", e, "0.00")
        c(ws, f"B{r}", "=" + N["cot"][b], "0.0000")
        c(ws, f"C{r}", "=" + N["cot"][h], "0.0000")
        c(ws, f"D{r}", "=" + N["cot"][f], "0.0000")
        c(ws, f"E{r}", fx(f"=B{r}*(1+{{madre_dif}})"), "0.0000")
        c(ws, f"F{r}", "=" + N["cot"][m], "0.0000")
        c(ws, f"I{r}", f"Bobina {b} · hoja {h} · fleje {f} · a medida {m}. {P_NOTA[e]}".strip())
        r += 1
    N["p1"] = r - 1
    p0, p1 = N["p0"], N["p1"]
    for k, col in (("pB", "B"), ("pH", "C"), ("pF", "D"), ("pM", "E"), ("pR", "F")):
        N[k] = f"Datos!${col}${p0}:${col}${p1}"
    N["pE"] = f"Datos!$A${p0}:$A${p1}"
    r += 1
    seccion(ws, r, "7. Equipos de corte y manejo (USD sin IVA)")
    g.cab(ws, r + 1, ["Equipo", "Precio (USD)", "A", "B", "C", "D", "Ref.", "", "Fuente"])
    r += 2
    e0 = r
    for nom, p, cant, fte in EQ:
        c(ws, f"A{r}", nom)
        esf = isinstance(p, str)
        c(ws, f"B{r}", fx(p) if esf else p, F0, dato=not esf)
        for j, q in enumerate(cant):
            v = fx('=IF({grua}="SI",0,1)') if q == "grua" else q
            c(ws, f"{L(3 + j)}{r}", v, F0, dato=q != "grua")
        c(ws, f"I{r}", fte)
        r += 1
    e1 = r - 1
    c(ws, f"A{r}", "Inversión por alternativa", neg=True)
    for j, o in enumerate(OPC):
        col = L(3 + j)
        c(ws, f"{col}{r}", f"=SUMPRODUCT($B${e0}:$B${e1},{col}{e0}:{col}{e1})", F0, neg=True)
        N[f"inv{o}"] = f"Datos!${col}${r}"
    g.nota(ws, r + 1, "Prensa de discos y cilindradoras están en todas las alternativas y no entran. "
                      "La guillotina 2 está a precio FOB.")
    ws.freeze_panes = "B4"


def precio(e_ref, k):
    return f"INDEX({N[k]},MATCH({e_ref},{N['pE']},0))"


# ------------------------------------------------------------------ Links (todos los IMPORTRANGE) y Demanda
LINEA = ["Manual"] * 4 + ["Rodante"] * 4
LK = {"proy": "Links!$B$4", "com": "Links!$B$5", "ind": "Links!$B$6"}
IMP = {}      # bloque -> primera fila de datos en Links


def h_links(wb):
    """Primera hoja: links (B4:B6) y los bloques IMPORTRANGE. Demanda los lee con IFERROR(…;0): sin link la fuente
    queda en 0 y la columna Estado lo avisa."""
    s = DC.D["S_ALTA"]
    ws = g.hoja(wb, "Links", [44, 11] + [11] * 10 + [30])
    g.tit(ws, 1, "Links e IMPORTRANGE")
    g.nota(ws, 2, "Pegar en B el link de cada Google Sheet. La primera vez, tocar la primera celda de cada bloque y "
                  "«Permitir acceso». Todos los IMPORTRANGE del libro están en esta hoja.")
    g.cab(ws, 3, ["Archivo", "Link", "", "", "", "", "Qué se importa", "", "", "", "", "Estado"])
    for c0, c1 in ((2, 6), (7, 11)):
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c1)
    bloques = [
        ("proy_t", "proy", "Proyecciones · Indexado Dim Técnico B3:K10: matafuegos terminados",
         [("Indexado Dim Técnico!B3:K10", DC.D["TERM"])], "Terminados"),
        ("proy_c", "proy", "Proyecciones · Indexado Dim Técnico B15:K22: cilindros vendidos vacíos",
         [("Indexado Dim Técnico!B15:K22", DC.D["CIL"])], "Cilindros"),
        ("ind", "ind", f"Recargas Planta Industrial · sobrestock_sustitutos B{s}:K{s + 3} y B{s + 13}:K{s + 16}: "
                       "sustitutos ABC terminados",
         [(f"sobrestock_sustitutos!B{s}:K{s + 3}", DC.D["IND"][:4]),
          (f"sobrestock_sustitutos!B{s + 13}:K{s + 16}", DC.D["IND"][4:])], "Sustitutos"),
        ("com", "com", f"Recargas Planta Comercial · sobrestock_sustitutos B{s}:K{s + 3} y B{s + 13}:K{s + 16}: "
                       "sustitutos ABC cilindros vacíos",
         [(f"sobrestock_sustitutos!B{s}:K{s + 3}", DC.D["COM"][:4]),
          (f"sobrestock_sustitutos!B{s + 13}:K{s + 16}", DC.D["COM"][4:])], "Sustitutos")]
    r = 9
    for clave, lk, t, imps, tipo in bloques:
        seccion(ws, r, t)
        g.cab(ws, r + 1, ["Formato", "Qué es"] + [str(a) for a in ANIOS])
        IMP[clave] = r + 2
        rr = r + 2
        for rng, vals in imps:
            g.importrange(ws, rr, 3, LK[lk], rng, vals, F0)
            rr += len(vals)
        for i, f in enumerate(FORM):
            c(ws, f"A{r + 2 + i}", f)
            c(ws, f"B{r + 2 + i}", tipo)
        r += 11
    filas = [("Proyecciones", DC.D["URL"], "Indexado Dim Técnico B3:K10 y B15:K22 → Demanda", "proy_t"),
             ("Aprovechamiento Recargas Planta Comercial", DC.D["PEGAR"],
              f"sobrestock_sustitutos B{s}:K{s + 3} y B{s + 13}:K{s + 16} → Demanda", "com"),
             ("Aprovechamiento Recargas Planta Industrial", DC.D["PEGAR"],
              f"sobrestock_sustitutos B{s}:K{s + 3} y B{s + 13}:K{s + 16} → Demanda", "ind")]
    for i, (a, u, q, blk) in enumerate(filas):
        rr = 4 + i
        c(ws, f"A{rr}", a, neg=True)
        c(ws, f"B{rr}", u, dato=True)
        c(ws, f"G{rr}", q)
        for c0, c1 in ((2, 6), (7, 11)):
            ws.merge_cells(start_row=rr, start_column=c0, end_row=rr, end_column=c1)
        c(ws, f"L{rr}", f'=IF(ISERROR(C{IMP[blk]}),"Falta el link o «Permitir acceso»","OK")')
    g.nota(ws, 7, "Sin link, esa fuente suma 0 en Demanda (la columna Estado lo avisa).")
    ws.freeze_panes = "C4"


def h_demanda(wb):
    ws = g.hoja(wb, "Demanda", [34, 10] + [11] * 10)
    titulo(ws, "Demanda: cuerpos a fabricar por formato", "Terminados + cilindros vacíos + sustitutos de las dos "
                                                         "plantas de recargas, todo desde la hoja Links. Cada unidad "
                                                         "lleva un cuerpo; las de 1 a 10 kg, además, cúpula y fondo.")
    seccion(ws, 4, "Unidades a fabricar = terminados + cilindros + sustitutos")
    g.cab(ws, 5, ["Formato", "Línea"] + [str(a) for a in ANIOS])
    for i, f in enumerate(FORM):
        r = 6 + i
        c(ws, f"A{r}", f)
        c(ws, f"B{r}", LINEA[i])
        for col in YC:
            c(ws, f"{col}{r}", "=" + "+".join(f"IFERROR(Links!{col}{IMP[k] + i},0)" for k in IMP), F0, neg=True)
    c(ws, "A14", "Total", neg=True)
    for col in YC:
        c(ws, f"{col}14", f"=SUM({col}6:{col}13)", F0, neg=True)
    g.nota(ws, 16, "Discos de rodantes (25 a 100 kg): se compran embutidos, no entran en este análisis.")
    ws.freeze_panes = "C6"


# ------------------------------------------------------------------ Piezas
PZ0 = 6        # filas de cuerpos (8)
DS0 = 18       # filas de discos (8)
DIAM = {"1 kg": 76.2, "2,5 kg": 124.0, "5 kg": 153.0, "10 kg": 181.5, "25 kg": 276.5, "50 kg": 320.0,
        "70 kg": 390.0, "100 kg": 390.0}


def h_piezas(wb):
    ws = g.hoja(wb, "Piezas", [10, 10, 30, 10, 10, 10, 10, 11, 11, 12, 11, 11, 11, 9, 10, 10, 13, 8, 11, 44, 3,
                               12, 10, 14])
    titulo(ws, "Piezas del plano: cuerpos y discos de cúpula y fondo",
           "Medidas de los planos de recipiente (modelo 3D). El disco sale de la pieza embutida por igualdad de "
           "superficie, con su relación de embutido.")
    seccion(ws, 4, "A) Cuerpos (desarrollo = π·(Ø − e), fibra media)")
    g.cab(ws, 5, ["Formato", "Línea", "Plano", "Ø ext. (mm)", "Espesor (mm)", "Alto (mm)", "Desarrollo (mm)",
                  "Masa (kg)", "Acero"])
    for i, f in enumerate(FORM):
        r = PZ0 + i
        al, de, e, ha, hl, ac = DC.CUERPO[f]
        c(ws, f"A{r}", f)
        c(ws, f"B{r}", LINEA[i])
        c(ws, f"C{r}", DC.PLANO[f])
        c(ws, f"D{r}", DIAM[f], "0.0", dato=True)
        c(ws, f"E{r}", e, "0.00", dato=True)
        c(ws, f"F{r}", al, "0.0", dato=True)
        c(ws, f"G{r}", f"=ROUND(PI()*(D{r}-E{r})*2,0)/2", "0.0")
        c(ws, f"H{r}", fx(f"=F{r}*G{r}*E{r}*{{dens}}/1000000"), F3)
        c(ws, f"I{r}", ac)
    g.nota(ws, PZ0 + 8, "Alto: cuerpo del modelo 3D del plano (incluye el labio del encastre). 1 kg en chapa de 0,9 mm "
                        "por decisión del grupo.")
    seccion(ws, DS0 - 2, "B) Discos: igual superficie que la pieza embutida (sin sobremedida de recorte)")
    hdr = ["Formato", "Pieza", "Perfil del plano", "Espesor e (mm)", "Ø ext. (mm)", "Altura h (mm)",
           "Pollera o pestaña p (mm)", "a = Ø/2 − e/2 (mm)", "b o profundidad (mm)", "Área media (mm²)",
           "D0 = √(4A/π) (mm)", "D0 adoptado (mm)", "d = 2a (mm)", "m = d/D0", "e/D0 × 100", "m1 mínimo",
           "¿Un solo embutido?", "h/d", "Masa del disco (kg)", "Plano"]
    g.cab(ws, DS0 - 1, hdr)
    for i, d in enumerate(DC.DISC):
        r = DS0 + i
        cas = d["perfil"] == "casquete"
        c(ws, f"A{r}", d["formato"])
        c(ws, f"B{r}", d["pieza"])
        c(ws, f"C{r}", "Semielipsoide + pollera" if cas else "Paraboloide + pestaña")
        for col, v, fmt in (("D", d["e"], "0.00"), ("E", d["d_ext"], "0.0"), ("F", d["h"], "0.0"),
                            ("G", d["p"], "0.0")):
            c(ws, f"{col}{r}", v, fmt, dato=True)
        h, b = f"H{r}", f"I{r}"
        ecc = f"SQRT(1-({b}/{h})^2)"
        c(ws, f"H{r}", f"=E{r}/2-D{r}/2", F2)
        c(ws, f"I{r}", f"=F{r}-G{r}-D{r}/2" if cas else f"=F{r}", F2)
        c(ws, f"J{r}", (f"=PI()*{h}^2+PI()*{b}^2/(2*{ecc})*LN((1+{ecc})/(1-{ecc}))+2*PI()*{h}*G{r}" if cas else
                        f"=PI()*{h}^4/(6*{b}^2)*((1+4*{b}^2/{h}^2)^1.5-1)+2*PI()*{h}*G{r}"), F0)
        c(ws, f"K{r}", f"=SQRT(4*J{r}/PI())", "0.0")
        c(ws, f"L{r}", f"=CEILING(K{r},0.5)", "0.0", neg=True)
        c(ws, f"M{r}", f"=2*H{r}", "0.0")
        c(ws, f"N{r}", f"=M{r}/L{r}", F3)
        c(ws, f"O{r}", f"=D{r}/L{r}*100", F2)
        c(ws, f"P{r}", f"=INDEX($X${DS0}:$X${DS0 + 5},MATCH(O{r},$V${DS0}:$V${DS0 + 5},1))", F2)
        c(ws, f"Q{r}", f'=IF(N{r}>=P{r},"Sí","No: dos embutidos")')
        c(ws, f"R{r}", f"=F{r}/E{r}" if cas else f"=(F{r}+G{r})/E{r}", F2)
        c(ws, f"S{r}", fx(f"=PI()*L{r}^2/4*D{r}*{{dens}}/1000000"), F3)
        c(ws, f"T{r}", DC.PLANO[d["formato"]] + ("" if cas else "; Ø del fondo = Ø del cuerpo − 2 e"))
    g.cab(ws, DS0 - 1, ["e/D0 × 100 desde", "hasta", "m1 (1.er embutido)"], c0=22)
    for i, (a, b, m) in enumerate(DC.M1):
        for j, v in enumerate((a, b, m)):
            c(ws, f"{L(22 + j)}{DS0 + i}", v, F2, dato=True)
    g.nota(ws, DS0 + 6, "Valor alto de cada rango (lado seguro).", c=22)
    fuentes = [
        "Disco y pieza tienen la misma superficie (el espesor no cambia): Harsle, «How to Calculate Deep Drawing "
        "Parts» (harslepress.com).",
        "Superficie del semielipsoide achatado y del paraboloide de revolución: fórmulas de geometría (MathWorld, "
        "«Oblate Spheroid» y «Paraboloid»).",
        "m1: relación de embutido del primer paso, copa sin brida, con prensachapas (EasyChair, preprint Vxvp, tabla 2). "
        "Límite general D0/d ≤ 2 (Univ. Mustansiriyah, apunte de embutido).",
        "Sin sobremedida de recorte: la prensa corta y embute en el mismo golpe, sin recorte después (W-MAN, puesto 4), "
        "y la tabla de recorte de Harsle arranca en h/d = 0,5 (estas piezas tienen h/d de 0,12 a 0,36).",
        "Geometría (Ø, alturas, pollera de 6 mm, pestaña de 10 mm, espesores): planos de recipiente (modelo 3D)."]
    seccion(ws, DS0 + 9, "Fuentes")
    for i, t in enumerate(fuentes):
        g.nota(ws, DS0 + 10 + i, t)
    ws.freeze_panes = "C6"


def pz(col, i):
    return f"Piezas!${col}${PZ0 + i}"


def ds(col, i):
    return f"Piezas!${col}${DS0 + i}"


# ------------------------------------------------------------------ Cálculo: plan de filas (se arma antes que Corte)
U0, D0, NETO = 8, 16, 25
NP0, NPROM = 26, 52   # hoja Nesting: primera fila de los patrones de 3,2 mm y de sus promedios por hoja
KP = (1, 3, 4)       # patrones de 3,2 mm que usa el plan (P2 = ⅓ P1 + ⅔ P3: mismas hojas, no suma)
GRUPOS = ["0,9", "1,25", "1,6", "2,0", "3,2", "4,75"]
GK = {"0,9": "09", "1,25": "125", "1,6": "16", "2,0": "20", "3,2": "32", "4,75": "475"}
R = {"U0": U0, "U5": U0 + 5, "U6": U0 + 6, "U7": U0 + 7, "U3": U0 + 3, "D0": D0, "D7": D0 + 7, "NETO": NETO}
PLAN = []          # (clave, etiqueta, unidad, formato, plantilla o None para título)


def _plan():
    def t(x):
        PLAN.append((None, x, None, None, None))

    def f(k, lab, uni, fmt, tpl):
        PLAN.append((k, lab, uni, fmt, tpl))
    t("A — Bobina por formato para cuerpos + fleje a medida para discos")
    f("A_tc", "Acero de bobina para cuerpos", "t", F1, "=SUMPRODUCT({c}${U0}:{c}${U7},Corte!$F${BB0}:$F${BB7})/1000")
    f("A_uc", "Costo de la bobina para cuerpos", "USD", F0, "=SUMPRODUCT({c}${U0}:{c}${U7},Corte!$G${BB0}:$G${BB7})")
    f("A_tf", "Acero de fleje para discos", "t", F1, "=SUMPRODUCT({c}${D0}:{c}${D7},Corte!$I${DC0}:$I${DC7})/1000")
    f("A_uf", "Costo del fleje para discos", "USD", F0, "=SUMPRODUCT({c}${D0}:{c}${D7},Corte!$J${DC0}:$J${DC7})")
    f("A_t", "Acero comprado", "t", F1, "={c}{A_tc}+{c}{A_tf}")
    f("A_u", "Costo del acero", "USD", F0, "={c}{A_uc}+{c}{A_uf}")
    f("A_sc", "Scrap", "t", F1, "={c}{A_t}-{c}{NETO}")
    f("A_vs", "Venta de scrap (resta)", "USD", F0, "=-{c}{A_sc}*{scrap}")
    stock = ("=SUMPRODUCT(({P}+{c}${a}:{c}${b}*{K}/1000*{cobmax}/12-ABS({P}-{c}${a}:{c}${b}*{K}/1000*{cobmax}/12))"
             "/4+{c}${a}:{c}${b}*{K}/1000*{seg}/52)")
    f("A_stc", "Stock de bobinas de cuerpo", "t", F1, stock.replace("{P}", "Corte!$H${BB0}:$H${BB7}").replace(
        "{K}", "Corte!$F${BB0}:$F${BB7}").replace("{a}", "{U0}").replace("{b}", "{U7}"))
    f("A_stf", "Stock de rollos de fleje", "t", F1, stock.replace("{P}", "Corte!$K${DC0}:$K${DC7}").replace(
        "{K}", "Corte!$I${DC0}:$I${DC7}").replace("{a}", "{D0}").replace("{b}", "{D7}"))
    f("A_st", "Stock promedio", "t", F1, "={c}{A_stc}+{c}{A_stf}")
    f("A_fin", "Costo financiero del stock", "USD", F0, "=IF({c}{A_t}=0,0,{c}{A_st}*{c}{A_u}/{c}{A_t}*{tasa})")
    f("A_fl", "Rollos de fleje por año", "rollos", F1,
      "=SUMPRODUCT({c}${D0}:{c}${D7}*Corte!$I${DC0}:$I${DC7}/1000/Corte!$K${DC0}:$K${DC7})")
    f("A_h", "Horas-hombre de corte: línea + guillotina 1 + cambios de fleje", "h", F0,
      "=SUMPRODUCT({c}${U0}:{c}${U7}/Corte!$C${BB0}:$C${BB7})/{linea}+SUM({c}${U0}:{c}${U7})/({gpm1}*60)"
      "+{c}{A_fl}*{cambio}/60")
    f("A_mo", "Mano de obra de corte", "USD", F0, "={c}{A_h}*{c_hora}")
    f("A_mt", "Mantenimiento de equipos", "USD", F0, "={invA}*{mant}")
    f("A_cost", "COSTO ANUAL A", "USD", F0, "={c}{A_u}+{c}{A_vs}+{c}{A_fin}+{c}{A_mo}+{c}{A_mt}")
    t("B — Hojas para todo: cuerpos con dos guillotinas y discos de tiras de hoja")
    f("B_mix", "Hojas mixtas 70 + 100 kg", "hojas", F1,
      '=IF({mezcla}="SI",MIN({c}{U7}/{mix100},{c}{U6}/{mix70}),0)')
    f("B_h70", "Hojas de 70 kg solas", "hojas", F1, "=({c}{U6}-{mix70}*{c}{B_mix})/Corte!$I${HS6}")
    f("B_h100", "Hojas de 100 kg solas", "hojas", F1, "=({c}{U7}-{mix100}*{c}{B_mix})/Corte!$I${HS7}")

    n25, n50 = "{c}{U4}", "{c}{U5}"
    f("B_p1", "Hojas de 3,2 mm, patrón P1 (6 de 25 kg + 3 de 50 kg)", "hojas", F1,
      f'=IF({{comb}}="SI",MIN({n25}/Nesting!$E${NP0},{n50}/Nesting!$F${NP0}),0)')
    f("B_p3", "Hojas de 3,2 mm, patrón P3 (6 de 50 kg)", "hojas", F1,
      f'=IF({{comb}}="SI",({n50}-{{c}}{{B_p1}}*Nesting!$F${NP0})/Nesting!$F${NP0 + 2},{n50}/Corte!$I${{HS5}})')
    f("B_p4", "Hojas de 3,2 mm, patrón P4 (9 de 25 kg)", "hojas", F1,
      f'=IF({{comb}}="SI",({n25}-{{c}}{{B_p1}}*Nesting!$E${NP0})/Nesting!$E${NP0 + 3},{n25}/Corte!$I${{HS4}})')

    def hojas(col):
        if col in ("P", "R"):     # golpes: los de cada patrón (hoja Nesting)
            c32 = "+".join(f"{{c}}{{B_p{k}}}*Nesting!${'G' if col == 'P' else 'H'}${NP0 + k - 1}" for k in KP)
        else:                     # masa y USD: misma hoja de 3,2 mm para los tres patrones
            c32 = f"({{c}}{{B_p1}}+{{c}}{{B_p3}}+{{c}}{{B_p4}})*Corte!${col}${{HS4}}"
        return (f"SUMPRODUCT({{c}}${{U0}}:{{c}}${{U3}}/Corte!$I${{HS0}}:$I${{HS3}}*Corte!${col}${{HS0}}:${col}${{HS3}})"
                f"+{c32}+{{c}}{{B_h70}}*Corte!${col}${{HS6}}+{{c}}{{B_h100}}*Corte!${col}${{HS7}}"
                f"+{{c}}{{B_mix}}*Corte!${col}${{HMIX}}")
    f("B_hc", "Hojas para cuerpos", "hojas", F0,
      "=SUMPRODUCT({c}${U0}:{c}${U3}/Corte!$I${HS0}:$I${HS3})+{c}{B_p1}+{c}{B_p3}+{c}{B_p4}"
      "+{c}{B_h70}+{c}{B_h100}+{c}{B_mix}")
    f("B_tc", "Acero de hojas para cuerpos", "t", F1, "=(" + hojas("W") + ")/1000")
    f("B_uc", "Costo de las hojas para cuerpos", "USD", F0, "=" + hojas("X"))
    f("B_g1", "Horas de la guillotina 1 (1.er corte: franjas)", "h", F1,
      "=(" + hojas("P") + ")/{gpm1}/60")
    f("B_g2", "Horas de la guillotina 2 (2.º corte, franjas en paralelo)", "h", F1, "=(" + hojas("R") + ")/{gpm2}/60")
    f("B_td", "Acero de hojas para discos", "t", F1, "=SUMPRODUCT({c}${U0}:{c}${U3},Corte!$E${UF0}:$E${UF3})/1000")
    f("B_ud", "Costo de las hojas para discos", "USD", F0, "=SUMPRODUCT({c}${U0}:{c}${U3},Corte!$F${UF0}:$F${UF3})")
    f("B_ti", "Tiras de hoja para discos", "tiras", F0, "=SUMPRODUCT({c}${U0}:{c}${U3},Corte!$G${UF0}:$G${UF3})")
    f("B_t", "Acero comprado", "t", F1, "={c}{B_tc}+{c}{B_td}")
    f("B_u", "Costo del acero", "USD", F0, "={c}{B_uc}+{c}{B_ud}")
    f("B_sc", "Scrap", "t", F1, "={c}{B_t}-{c}{NETO}")
    f("B_vs", "Venta de scrap (resta)", "USD", F0, "=-{c}{B_sc}*{scrap}")
    f("B_st", "Stock promedio (hojas)", "t", F1, "={c}{B_t}*({lote}/2/12+{seg}/52)")
    f("B_fin", "Costo financiero del stock", "USD", F0, "=IF({c}{B_t}=0,0,{c}{B_st}*{c}{B_u}/{c}{B_t}*{tasa})")
    f("B_h", "Horas-hombre de corte: guillotinas + tiras + carga manual", "h", F0,
      "={c}{B_g1}+{c}{B_g2}+{c}{B_ti}/{gpm1}/60+{c}{B_ti}*{tira}/3600")
    f("B_mo", "Mano de obra de corte", "USD", F0, "={c}{B_h}*{c_hora}")
    f("B_mt", "Mantenimiento de equipos", "USD", F0, "={invB}*{mant}")
    f("B_cost", "COSTO ANUAL B", "USD", F0, "={c}{B_u}+{c}{B_vs}+{c}{B_fin}+{c}{B_mo}+{c}{B_mt}")
    t("C — Bobina madre por espesor con corte longitudinal propio")
    for gr in GRUPOS:
        f(f"C_m{GK[gr]}", f"Metros de madre de {gr} mm", "m", F0, None)
    for gr in GRUPOS:
        f(f"C_t{GK[gr]}", f"Acero de madre de {gr} mm", "t", F1, None)
    f("C_t", "Acero comprado", "t", F1, "=SUM({c}{C_t09}:{c}{C_t475})")
    f("C_u", "Costo del acero", "USD", F0, "=SUMPRODUCT({c}{C_t09}:{c}{C_t475},Corte!$E${GS0}:$E${GS5})*1000")
    f("C_sc", "Scrap", "t", F1, "={c}{C_t}-{c}{NETO}")
    f("C_vs", "Venta de scrap (resta)", "USD", F0, "=-{c}{C_sc}*{scrap}")
    f("C_st", "Stock promedio (madres y flejes en proceso)", "t", F1,
      "=SUMPRODUCT((Corte!$G${GS0}:$G${GS5}+{c}{C_t09}:{c}{C_t475}*{cobmax}/12-ABS(Corte!$G${GS0}:$G${GS5}"
      "-{c}{C_t09}:{c}{C_t475}*{cobmax}/12))/2+{c}{C_t09}:{c}{C_t475}*{seg}/52)")
    f("C_fin", "Costo financiero del stock", "USD", F0, "=IF({c}{C_t}=0,0,{c}{C_st}*{c}{C_u}/{c}{C_t}*{tasa})")
    f("C_hm", "Horas de la cortadora longitudinal (máquina)", "h", F1,
      "=SUM({c}{C_m09}:{c}{C_m475})/({slit_v}*60)+SUMPRODUCT({c}{C_t09}:{c}{C_t475}/Corte!$G${GS0}:$G${GS5})"
      "*{slit_arm}/60")
    f("C_fl", "Flejes de disco generados", "flejes", F1,
      "=SUMPRODUCT({c}{C_t09}:{c}{C_t475}/Corte!$G${GS0}:$G${GS5}*Corte!$F${GS0}:$F${GS5})")
    f("C_h", "Horas-hombre de corte: cortadora + línea + cambios de fleje", "h", F0,
      "={c}{C_hm}*{slit_op}+SUM({c}${U0}:{c}${U7})/{linea}+{c}{C_fl}*{cambio}/60")
    f("C_mo", "Mano de obra de corte", "USD", F0, "={c}{C_h}*{c_hora}")
    f("C_mt", "Mantenimiento de equipos", "USD", F0, "={invC}*{mant}")
    f("C_cost", "COSTO ANUAL C", "USD", F0, "={c}{C_u}+{c}{C_vs}+{c}{C_fin}+{c}{C_mo}+{c}{C_mt}")
    t("D — Hojas para cuerpos (dos guillotinas en serie) + fleje a medida para discos")
    f("D_t", "Acero comprado: hojas de cuerpo + fleje", "t", F1, "={c}{B_tc}+{c}{A_tf}")
    f("D_u", "Costo del acero", "USD", F0, "={c}{B_uc}+{c}{A_uf}")
    f("D_sc", "Scrap", "t", F1, "={c}{D_t}-{c}{NETO}")
    f("D_vs", "Venta de scrap (resta)", "USD", F0, "=-{c}{D_sc}*{scrap}")
    f("D_st", "Stock promedio: hojas + rollos de fleje", "t", F1, "={c}{B_tc}*({lote}/2/12+{seg}/52)+{c}{A_stf}")
    f("D_fin", "Costo financiero del stock", "USD", F0, "=IF({c}{D_t}=0,0,{c}{D_st}*{c}{D_u}/{c}{D_t}*{tasa})")
    f("D_h", "Horas-hombre de corte: guillotinas 1 y 2 + cambios de fleje", "h", F0,
      "={c}{B_g1}+{c}{B_g2}+{c}{A_fl}*{cambio}/60")
    f("D_mo", "Mano de obra de corte", "USD", F0, "={c}{D_h}*{c_hora}")
    f("D_mt", "Mantenimiento de equipos", "USD", F0, "={invD}*{mant}")
    f("D_cost", "COSTO ANUAL D", "USD", F0, "={c}{D_u}+{c}{D_vs}+{c}{D_fin}+{c}{D_mo}+{c}{D_mt}")
    t("Referencia — cuerpos comprados cortados a medida + fleje a medida para discos")
    f("R_tc", "Cuerpos a medida", "t", F1, "=SUMPRODUCT({c}${U0}:{c}${U7},Piezas!$H${PZ0}:$H${PZ7})/1000")
    f("R_uc", "Costo de los cuerpos a medida", "USD", F0, "=SUMPRODUCT({c}${U0}:{c}${U7},Corte!$I${BB0}:$I${BB7})")
    f("R_t", "Acero comprado", "t", F1, "={c}{R_tc}+{c}{A_tf}")
    f("R_u", "Costo del acero", "USD", F0, "={c}{R_uc}+{c}{A_uf}")
    f("R_sc", "Scrap (esqueleto de los discos)", "t", F1, "={c}{R_t}-{c}{NETO}")
    f("R_vs", "Venta de scrap (resta)", "USD", F0, "=-{c}{R_sc}*{scrap}")
    f("R_st", "Stock promedio: cuerpos + rollos de fleje", "t", F1, "={c}{R_tc}*({lote}/2/12+{seg}/52)+{c}{A_stf}")
    f("R_fin", "Costo financiero del stock", "USD", F0, "=IF({c}{R_t}=0,0,{c}{R_st}*{c}{R_u}/{c}{R_t}*{tasa})")
    f("R_h", "Horas-hombre: cambios de fleje", "h", F0, "={c}{A_fl}*{cambio}/60")
    f("R_mo", "Mano de obra", "USD", F0, "={c}{R_h}*{c_hora}")
    f("R_mt", "Mantenimiento de equipos", "USD", F0, "={invRef}*{mant}")
    f("R_cost", "COSTO ANUAL REFERENCIA", "USD", F0, "={c}{R_u}+{c}{R_vs}+{c}{R_fin}+{c}{R_mo}+{c}{R_mt}")
    t("Opción D: compras por año")
    for i, fm in enumerate(FORM):
        tpl = {4: "={c}{B_p4}", 5: "={c}{B_p3}", 6: "={c}{B_h70}", 7: "={c}{B_h100}"}.get(
            i, f"={{c}}{{U{i}}}/Corte!$I${{HS{i}}}")
        f(f"Dh{i}", f"Hojas para cuerpos de {fm}" + (" solos" if i in (4, 5, 6, 7) else ""), "hojas", F0, tpl)
    f("Dhc", "Hojas combinadas 25 + 50 kg (patrón P1)", "hojas", F0, "={c}{B_p1}")
    f("Dhm", "Hojas mixtas 70 + 100 kg", "hojas", F0, "={c}{B_mix}")
    for i, d in enumerate(DC.DISC):
        f(f"Dr{i}", f"Rollos de fleje: {d['pieza'].lower()} {d['formato']}", "rollos", F1,
          f"={{c}}{{D{i}}}*Corte!$I${{DC{i}}}/1000/Corte!$K${{DC{i}}}")
    r = NETO + 2
    for i, (k, lab, uni, fmt, tpl) in enumerate(PLAN):
        if k is None:
            r += 1 if i else 0
            R[f"T{i}"] = r
        else:
            R[k] = r
        r += 1
    for i in range(8):
        R[f"U{i}"] = U0 + i
        R[f"D{i}"] = D0 + i


_plan()
HS0, HMIX, EX0, BB0, DC0, UF0, GS0, GD0 = 6, 14, 19, 32, 44, 55, 63, 71
R.update({"HS0": HS0, "HS5": HS0 + 5, "HS6": HS0 + 6, "HS7": HS0 + 7, "HMIX": HMIX, "BB0": BB0, "BB7": BB0 + 7,
          "DC0": DC0, "DC7": DC0 + 7, "UF0": UF0, "UF3": UF0 + 3, "GS0": GS0, "GS5": GS0 + 5,
          "PZ0": PZ0, "PZ7": PZ0 + 7})
for _i in range(8):
    R[f"HS{_i}"] = HS0 + _i
    R[f"DC{_i}"] = DC0 + _i


# ------------------------------------------------------------------ Corte
HOJA_TIRAS = [(1220, 2440), (1220, 2440), (1220, 2440), (1220, 2440), (1220, 2440), (1000, 2000), (1000, 2000),
              (1500, 3000)]
BOBINA_A = [1220, 1220, 1000, 1000, 1500, 1500, 1245, 1245]
GRP = {}


def anio_ref(rango):
    return f"INDEX({rango},{N['anio']}-2025)"


def h_corte(wb):
    ws = g.hoja(wb, "Corte", [24, 9, 9, 9, 9, 8, 9, 8, 8, 8, 9, 9, 9, 8, 8, 8, 8, 8, 8, 8, 7, 9, 9, 9, 34, 2,
                              9, 9, 9, 9, 9, 9])
    titulo(ws, "Corte: cómo sale cada pieza de la hoja, la bobina, el fleje y la bobina madre",
           "Medidas de la hoja Piezas. G1 corta franjas por el lado corto de la hoja; pasan por un pulmón y G2 apoya "
           "las franjas de hojas enteras una al lado de la otra y corta una pieza de cada una por golpe.")
    seccion(ws, 4, "A) Cuerpos en hoja estándar con dos guillotinas en serie (B y D)")
    g.cab(ws, 5, ["Formato", "Hoja: lado corto (mm)", "Hoja: lado largo (mm)", "Orientación (1 alto, 2 desarrollo "
                  "a lo largo)", "Franja = 1.er corte y tope de G1 (mm)", "Franjas por hoja", "Pieza = 2.º corte (mm)",
                  "Piezas por franja", "Piezas por hoja", "Aprovech.", "Masa de la franja (kg)",
                  "Tope de G2 (pieza u orilla, mm)", "Franjas que entran en G2", "Hojas por ciclo",
                  "Tandas de G2 por ciclo", "Golpes G1 por hoja", "Golpes G2 por tanda", "Golpes G2 por hoja",
                  "Min G1 por hoja", "Min G2 por hoja", "Manda", "Cuerpos por hora", "Masa de la hoja (kg)",
                  "USD por hoja", "Control"])
    g.cab(ws, 5, ["Tandas por hoja con ciclo de 1 hoja", "2 hojas", "3 hojas", "4 hojas",
                  "Piezas por hoja, orientación 1", "orientación 2"], c0=27)
    for i, f in enumerate(FORM):
        r = HS0 + i
        al, de, e = pz("F", i), pz("G", i), pz("E", i)
        B, C = f"B{r}", f"C{r}"

        def feas(a, b):
            return fx(f"AND({a}<={{tope1}},{a}<={{L2}},{B}<={{L1}},OR({b}<={{tope2}},AND(INT({B}/{b})=1,"
                      f"{B}-{b}<={{tope2}})))")
        c(ws, f"A{r}", f)
        c(ws, f"B{r}", DC.CUERPO[f][3], F0, dato=True)
        c(ws, f"C{r}", DC.CUERPO[f][4], F0, dato=True)
        c(ws, f"AE{r}", f"=IF({feas(al, de)},INT({C}/{al})*INT({B}/{de}),0)", F0)
        c(ws, f"AF{r}", f"=IF({feas(de, al)},INT({C}/{de})*INT({B}/{al}),0)", F0)
        c(ws, f"D{r}", f"=IF(AE{r}>AF{r},1,IF(AF{r}>AE{r},2,IF(INT({C}/{al})<=INT({C}/{de}),1,2)))", F0)
        c(ws, f"E{r}", f"=IF(D{r}=1,{al},{de})", "0.0")
        c(ws, f"F{r}", f"=INT({C}/E{r})", F0)
        c(ws, f"G{r}", f"=IF(D{r}=1,{de},{al})", "0.0")
        c(ws, f"H{r}", f"=INT({B}/G{r})", F0)
        c(ws, f"I{r}", f"=F{r}*H{r}", F0, neg=True)
        c(ws, f"J{r}", f"=I{r}*{al}*{de}/({B}*{C})", FP)
        c(ws, f"M{r}", fx(f"=INT({{L2}}/E{r})"), F0)
        for k in range(1, 5):
            c(ws, f"{L(26 + k)}{r}", fx(f"=IFERROR(IF({k}>{{ciclo}},1E+9,CEILING({k}*F{r}/M{r},1)/{k}),1E+9)"), F2)
        _hoja_maquinas(ws, r, e, f"C{r}-F{r}*E{r}")
        base = (f'IF(I{r}=0,"No entra ninguna pieza",IF({B}>{{L1}},"Hoja más ancha que la cuchilla de G1",'
                f'IF(E{r}>{{tope1}},"Franja mayor que el tope de G1",IF(L{r}>{{tope2}},"2.º corte mayor que el tope '
                f'de G2",IF(M{r}<1,"La franja no entra en G2","OK")))))')
        if f == "70 kg":
            base = f'IF({base}="OK",IF({{mezcla}}="SI","OK; con mezcla va en la hoja mixta","OK"),{base})'
        c(ws, f"Y{r}", fx("=" + base))
    r = HMIX
    a70, d70, a100, d100, e = pz("F", 6), pz("G", 6), pz("F", 7), pz("G", 7), pz("E", 7)
    ancho = fx(f"({{mix100}}*{a100}+{{mix70}}*{a70})")
    c(ws, f"A{r}", "Hoja mixta 70 + 100 kg", neg=True)
    c(ws, f"B{r}", f"=B{HS0 + 7}", F0)
    c(ws, f"C{r}", f"=C{HS0 + 7}", F0)
    c(ws, f"D{r}", 1, F0)
    c(ws, f"E{r}", f"=MAX({a100},{a70})", "0.0")
    c(ws, f"F{r}", fx("={mix100}+{mix70}"), F0)
    c(ws, f"G{r}", f"=MAX({d100},{d70})", "0.0")
    c(ws, f"H{r}", f"=INT(B{r}/G{r})", F0)
    c(ws, f"I{r}", f"=F{r}*H{r}", F0, neg=True)
    c(ws, f"J{r}", fx(f"=({{mix100}}*{a100}*{d100}+{{mix70}}*{a70}*{d70})/(B{r}*C{r})"), FP)
    c(ws, f"M{r}", fx(f"=INT({{L2}}/MIN({a100},{a70}))"), F0)
    for k in range(1, 5):
        c(ws, f"{L(26 + k)}{r}", fx(f"=IFERROR(IF({k}>{{ciclo}},1E+9,MAX(CEILING({k}*F{r}/M{r},1),"
                                    f"CEILING({k}*{ancho}/{{L2}},1))/{k}),1E+9)"), F2)
    _hoja_maquinas(ws, r, e, f"C{r}-{ancho}")
    c(ws, f"Y{r}", fx(f'=IF(AND({d70}<=B{r},{d100}<=B{r},{ancho}<=C{r},E{r}<={{tope1}},L{r}<={{tope2}}),'
                      f'"OK: "&{{mix100}}&" de 100 kg y "&{{mix70}}&" de 70 kg","No entra")'))
    g.nota(ws, HMIX + 1, "1.er corte por el lado corto (franjas de hasta 1.500 mm de largo). G1 hace solo el 1.er corte y "
                         "G2 solo el 2.º: la pieza, o la orilla si la pieza supera el tope. Franjas en paralelo: hojas "
                         "enteras, con el ciclo que da menos tandas por hoja (Datos, «Hojas por ciclo»).")
    # para los Aprovechamientos de Manuales y Carros
    seccion(ws, EX0 - 2, "Para los Aprovechamientos de Manuales (1 a 10 kg) y Carros (25 a 100 kg): IMPORTRANGE")
    g.cab(ws, EX0 - 1, ["Formato", "Piezas por hoja", "Golpes G1 por hoja", "Golpes G2 por hoja",
                        "Tandas de G2 por hoja", "Golpes por minuto G1", "Golpes por minuto G2"])
    for i in range(9):
        r, rs = EX0 + i, HS0 + i
        c(ws, f"A{r}", f"=A{rs}")
        base = [f"I{rs}", f"P{rs}", f"R{rs}", f"O{rs}/N{rs}"]
        for j, col in enumerate("BCDE"):
            v = base[j] if i not in (4, 5) else fx(f'IF({{comb}}="SI",Nesting!$B${NPROM + j},{base[j]})')
            c(ws, f"{col}{r}", "=" + v, F0 if j == 0 and i not in (4, 5) else F2)
        c(ws, f"F{r}", fx("={gpm1}"), F0)
        c(ws, f"G{r}", fx("={gpm2}"), F0)
    # B) bobina por formato
    seccion(ws, BB0 - 2, "B) Cuerpos en bobina de ancho estándar con línea de corte transversal (A) y a medida (Ref.)")
    g.cab(ws, BB0 - 1, ["Formato", "Ancho de bobina (mm)", "Piezas por golpe", "Avance por golpe (mm)", "Aprovech.",
                        "kg comprados por cuerpo", "USD por cuerpo", "Peso de bobina (t)",
                        "USD por cuerpo a medida (Ref.)"])
    for i, f in enumerate(FORM):
        r = BB0 + i
        al, de, e = pz("F", i), pz("G", i), pz("E", i)
        a1 = fx(f"INT((B{r}-{{margen}})/{al})*{al}/B{r}")
        a2 = fx(f"INT((B{r}-{{margen}})/{de})*{de}/B{r}")
        c(ws, f"A{r}", f)
        c(ws, f"B{r}", BOBINA_A[i], F0, dato=True)
        c(ws, f"C{r}", fx(f"=IF({a1}>={a2},INT((B{r}-{{margen}})/{al}),INT((B{r}-{{margen}})/{de}))"), F0)
        c(ws, f"D{r}", f"=IF({a1}>={a2},{de},{al})", "0.0")
        c(ws, f"E{r}", fx(f"=MAX({a1},{a2})*(1-{{despunte}})"), FP)
        c(ws, f"F{r}", f"={pz('H', i)}/E{r}", F3)
        c(ws, f"G{r}", f"=F{r}*{precio(e, 'pB')}", F2)
        cons = anio_ref(f"Demanda!$C${6 + i}:$L${6 + i}") + f"*F{r}/1000"
        c(ws, f"H{r}", fx(f"=MIN(ROUNDDOWN(MIN(B{r}*{{t_mm}},{{des_lin}})*2,0)/2,MAX({{pmin_rollo}},"
                          f"ROUNDUP({cons}/12*{{cob}}*2,0)/2))"), "0.0")
        c(ws, f"I{r}", f"={pz('H', i)}*{precio(e, 'pR')}", F2)
    # C) discos
    seccion(ws, DC0 - 2, "C) Cúpulas y fondos: fleje a medida (A, D, Ref.) o tiras de hoja (B)")
    g.cab(ws, DC0 - 1, ["Pieza", "Ø disco (mm)", "Espesor (mm)", "Puente de avance (mm)", "Puente al borde (mm)",
                        "Ancho de fleje (mm)", "Paso (mm)", "Aprovech. del fleje", "kg de fleje por disco",
                        "USD de fleje por disco", "Peso de rollo (t)", "Hoja: lado corto = largo de la tira", "lado largo",
                        "Tiras por hoja", "Discos por tira", "Discos por hoja", "Aprovech. de la hoja",
                        "kg de hoja por disco", "USD de hoja por disco", "Fleje 1 fila: ancho (mm)",
                        "Aprovech. 1 fila", "Fleje 2 filas al tresbolillo: ancho (mm)", "Aprovech. 2 filas",
                        "kg por disco, 1 fila", "kg por disco, 2 filas"])
    fmi = {f: i for i, f in enumerate(FORM)}
    for i, d in enumerate(DC.DISC):
        r = DC0 + i
        e = f"C{r}"
        c(ws, f"A{r}", f"{d['pieza']} {d['formato']}")
        c(ws, f"B{r}", f"={ds('L', i)}", "0.0")
        c(ws, f"C{r}", f"={ds('D', i)}", "0.00")
        c(ws, f"D{r}", fx(f"=MAX({{pf}}*C{r},{{pmin}})"), F2)
        c(ws, f"E{r}", fx(f"={{pb}}*D{r}"), F2)
        c(ws, f"F{r}", fx(f'=IF({{tresb}}="SI",V{r},T{r})'), "0.0")
        c(ws, f"G{r}", f"=B{r}+D{r}", "0.0")
        c(ws, f"H{r}", fx(f'=IF({{tresb}}="SI",W{r},U{r})'), FP)
        c(ws, f"T{r}", f"=B{r}+2*E{r}", "0.0")
        c(ws, f"U{r}", fx(f"=PI()*B{r}^2/4/(T{r}*G{r})*(1-{{despunte}})"), FP)
        c(ws, f"V{r}", f"=B{r}+G{r}*SQRT(3)/2+2*E{r}", "0.0")
        c(ws, f"W{r}", fx(f"=2*PI()*B{r}^2/4/(V{r}*G{r})*(1-{{despunte}})"), FP)
        c(ws, f"X{r}", f"={ds('S', i)}/U{r}", F3)
        c(ws, f"Y{r}", f"={ds('S', i)}/W{r}", F3)
        c(ws, f"I{r}", f"={ds('S', i)}/H{r}", F3)
        c(ws, f"J{r}", f"=I{r}*{precio(e, 'pF')}", F3)
        fr = fmi[d["formato"]]
        cons = anio_ref(f"Demanda!$C${6 + fr}:$L${6 + fr}") + f"*I{r}/1000"
        c(ws, f"K{r}", fx(f"=MIN(ROUNDDOWN(MIN(F{r}*{{t_mm}},{{des_pre}})*2,0)/2,MAX({{pmin_rollo}},"
                          f"ROUNDUP({cons}/12*{{cob}}*2,0)/2))"), "0.0")
        hw, hl = HOJA_TIRAS[i]
        c(ws, f"L{r}", hw, F0, dato=True)
        c(ws, f"M{r}", hl, F0, dato=True)
        # tiras cortadas por el lado corto de la hoja (mismo criterio que los cuerpos)
        c(ws, f"N{r}", f"=INT(M{r}/T{r})", F0)
        c(ws, f"O{r}", f"=INT((L{r}-D{r})/G{r})", F0)
        c(ws, f"P{r}", f"=N{r}*O{r}", F0)
        c(ws, f"Q{r}", f"=P{r}*PI()*B{r}^2/4/(L{r}*M{r})", FP)
        c(ws, f"R{r}", fx(f"=L{r}*M{r}*C{r}*{{dens}}/1000000/P{r}"), F3)
        c(ws, f"S{r}", f"=R{r}*{precio(e, 'pH')}", F3)
    seccion(ws, UF0 - 2, "Por unidad de matafuego (cúpula + fondo)")
    g.cab(ws, UF0 - 1, ["Formato", "Masa neta (kg)", "kg de fleje", "USD de fleje", "kg de hoja", "USD de hoja",
                        "Tiras de hoja"])
    for i, f in enumerate(FORM[:4]):
        r = UF0 + i
        a, b = DC0 + 2 * i, DC0 + 2 * i + 1
        c(ws, f"A{r}", f)
        c(ws, f"B{r}", f"={ds('S', 2 * i)}+{ds('S', 2 * i + 1)}", F3)
        c(ws, f"C{r}", f"=I{a}+I{b}", F3)
        c(ws, f"D{r}", f"=J{a}+J{b}", F3)
        c(ws, f"E{r}", f"=R{a}+R{b}", F3)
        c(ws, f"F{r}", f"=S{a}+S{b}", F3)
        c(ws, f"G{r}", f"=1/O{a}+1/O{b}", F3)
    _bobina_madre(ws)
    ws.freeze_panes = "B6"


def _hoja_maquinas(ws, r, e, resto):
    c(ws, f"K{r}", fx(f"=E{r}*B{r}*{e}*{{dens}}/1000000"), F1)
    c(ws, f"L{r}", fx(f"=IF(G{r}<={{tope2}},G{r},B{r}-H{r}*G{r})"), "0.0")
    c(ws, f"N{r}", f"=MATCH(MIN(AA{r}:AD{r}),AA{r}:AD{r},0)", F0)
    c(ws, f"O{r}", f"=ROUND(INDEX(AA{r}:AD{r},N{r})*N{r},0)", F0)
    c(ws, f"P{r}", f"=IF({resto}>0.001,F{r},F{r}-1)", F0)
    c(ws, f"Q{r}", fx(f"=IF(G{r}<={{tope2}},IF(B{r}-H{r}*G{r}>0.001,H{r},H{r}-1),1)"), F0)
    c(ws, f"R{r}", f"=Q{r}*O{r}/N{r}", F2)
    c(ws, f"S{r}", fx(f"=P{r}/{{gpm1}}"), F2)
    c(ws, f"T{r}", fx(f"=R{r}/{{gpm2}}"), F2)
    c(ws, f"U{r}", f'=IF(S{r}>=T{r},"G1","G2")', neg=True)
    c(ws, f"V{r}", f"=I{r}*60/MAX(S{r},T{r})", F0)
    c(ws, f"W{r}", fx(f"=B{r}*C{r}*{e}*{{dens}}/1000000"), F1)
    c(ws, f"X{r}", f"=W{r}*{precio(e, 'pH')}", F2)


def _bobina_madre(ws):
    lp = DC.patrones_c()
    seccion(ws, GS0 - 2, "D) Bobina madre por espesor con corte longitudinal propio (C): patrones de tiras")
    g.cab(ws, GS0 - 1, ["Grupo", "Ancho de madre (mm)", "Cuerpo a lo ancho (1 alto, 2 desarrollo)",
                        "kg por metro", "USD/kg", "Tiras de disco por ancho", "Peso de madre (t)", "Control"])
    r = GD0
    for j, gr in enumerate(GRUPOS):
        rs = GS0 + j
        madre = {"4,75": 1245}.get(gr, DC.MADRE.get(gr))
        orient = lp[gr]["orient"] if gr in lp else 2
        c(ws, f"A{rs}", f"{gr} mm")
        c(ws, f"B{rs}", madre, F0, dato=True)
        c(ws, f"C{rs}", orient if gr != "4,75" else "2 (fijo)", F0, dato=gr != "4,75")
        esp = float(gr.replace(",", "."))
        c(ws, f"D{rs}", fx(f"=B{rs}*{esp}*{{dens}}/1000"), F2)
        c(ws, f"E{rs}", f"={precio(esp, 'pM')}", "0.0000")
        t_row = R[f"C_t{GK[gr]}"]
        cons = anio_ref(f"Cálculo!$C${t_row}:$L${t_row}")
        c(ws, f"G{rs}", fx(f"=MIN(ROUNDDOWN(MIN(B{rs}*{{t_mm}},{{des_slit}})*2,0)/2,MAX({{pmin_rollo}},"
                           f"ROUNDUP({cons}/12*{{cob}}*2,0)/2))"), "0.0")
        # detalle del grupo
        seccion(ws, r, f"Grupo {gr} mm")
        g.cab(ws, r + 1, ["Tira", "Formato", "Ancho (mm)", "Avance por pieza (mm)", "Patrón 1", "Patrón 2", "Patrón 3",
                          "Tiras por metro de madre"])
        if gr == "4,75":
            tiras = [("Cuerpos 70 y 100 kg (franja compartida)", "70 y 100 kg", "c2")]
            pats, part = [(1,)], [1.0]
        else:
            tiras = [(t[0], t[3], "c" if t[0].startswith("Cuerpo") else "d") for t in lp[gr]["tiras"]]
            pats, part = lp[gr]["pats"], [float(x) for x in lp[gr]["part"]]
        r0 = r + 2
        rp = r0 + len(tiras)
        info = []
        for i, (nom, fmt_, tipo) in enumerate(tiras):
            rr = r0 + i
            c(ws, f"A{rr}", nom)
            c(ws, f"B{rr}", fmt_)
            if tipo == "c":
                k = FORM.index(fmt_)
                c(ws, f"C{rr}", f"=IF($C${rs}=1,{pz('F', k)},{pz('G', k)})", "0.0")
                c(ws, f"D{rr}", f"=IF($C${rs}=1,{pz('G', k)},{pz('F', k)})", "0.0")
                info.append((k, f"Corte!$D${rr}", f"Corte!$H${rr}"))
            elif tipo == "c2":
                c(ws, f"C{rr}", f"={pz('G', 7)}", "0.0")
                c(ws, f"D{rr}", "alto de cada uno")
                info.append(("4,75", None, f"Corte!$H${rr}"))
            else:
                pieza, fm = nom.split(" ", 1)
                di = next(x for x, d in enumerate(DC.DISC) if d["pieza"] == pieza and d["formato"] == fm)
                c(ws, f"C{rr}", f"=F{DC0 + di}", "0.0")
                c(ws, f"D{rr}", f"=G{DC0 + di}", "0.0")
                info.append((FORM.index(fm), f"Corte!$D${rr}", f"Corte!$H${rr}"))
            for jj, col in enumerate("EFG"):
                c(ws, f"{col}{rr}", pats[jj][i] if jj < len(pats) else 0, F0, dato=True)
            c(ws, f"H{rr}", f"=E{rr}*$E${rp}+F{rr}*$F${rp}+G{rr}*$G${rp}", F3)
        c(ws, f"A{rp}", "Participación (% de los metros de madre)")
        for jj, col in enumerate("EFG"):
            c(ws, f"{col}{rp}", part[jj] if jj < len(part) else 0, FP, dato=True)
        c(ws, f"H{rp}", f'=IF(ABS(SUM(E{rp}:G{rp})-1)<0.001,"Suma 100 %","Revisar: no suma 100 %")')
        c(ws, f"A{rp + 1}", "Ancho usado (mm)")
        c(ws, f"A{rp + 2}", "¿Entra en la madre? (orilla ≥ margen)")
        for col in "EFG":
            c(ws, f"{col}{rp + 1}", f"=SUMPRODUCT($C${r0}:$C${rp - 1},{col}{r0}:{col}{rp - 1})", "0.0")
            c(ws, f"{col}{rp + 2}", fx(f'=IF({col}{rp + 1}=0,"sin uso",IF($B${rs}-{col}{rp + 1}>={{margen}},'
                                       f'"OK","No entra"))'))
        discos = [f"H{r0 + i}" for i, t in enumerate(tiras) if t[2] == "d"]
        c(ws, f"F{rs}", "=" + "+".join(discos) if discos else 0, F2)
        c(ws, f"H{rs}", f'=IF(AND(H{rp}="Suma 100 %",COUNTIF(E{rp + 2}:G{rp + 2},"No entra")=0),"OK","Revisar")')
        GRP[gr] = dict(info=info, fila=rs)
        r = rp + 4
    g.nota(ws, r - 1, "Patrones y participaciones por programación lineal sobre la demanda 2026-2035 (mínimo acero "
                      "comprado); si cambia la demanda o una medida, hay que recalcularlos.")


# ------------------------------------------------------------------ Nesting
CAND = [("1 kg", 1220, 2440, "Q12"), ("1 kg", 1000, 2000, "Q13"), ("2,5 kg", 1220, 2440, "Q8"),
        ("5 kg", 1000, 2000, "Q7"), ("5 kg", 1220, 2440, "Q15"), ("10 kg", 1500, 3000, "Q9"),
        ("10 kg", 1220, 2440, "Q16"), ("25 kg", 1500, 3000, "Q10"), ("25 kg", 1500, 3000, "Q17"),
        ("50 kg", 1500, 3000, "Q10"), ("50 kg", 1500, 3000, "Q17"), ("70 kg", 1500, 3000, "Q11"),
        ("70 kg", 1500, 3000, "Q18"), ("100 kg", 1500, 3000, "Q11"), ("100 kg", 1500, 3000, "Q18")]
# patrones de 3,2 mm (1500 × 3000): (nombre, franjas, 25 kg, 50 kg, etapas, golpes G1, golpes G2, tandas de G2)
PAT32 = [("P1", "3 franjas de 995,5: 1 de 50 kg (640) + 2 de 25 kg (tramo de 858,5 cortado en 493 + 493)", 6, 3, 3,
          3, 6, 3),
         ("P2", "2 franjas de 995,5 con 2 de 50 kg + 1 franja como P1 (= ⅓ P1 + ⅔ P3)", 2, 5, 3, 3, 4, 2),
         ("P3", "3 franjas de 995,5 con 2 de 50 kg", 0, 6, 2, 3, 3, 1.5),
         ("P4", "3 franjas de 858,5 con 3 de 25 kg", 9, 0, 2, 3, 4.5, 1.5)]
PAT475 = [("3 de 70 kg + 1 de 100 kg (hoja mixta)", 3, 1), ("4 de 70 kg", 4, 0), ("3 de 100 kg", 0, 3),
          ("1 de 70 kg + 2 de 100 kg", 1, 2)]
NA0, NC0, ND0 = 6, 59, 67      # filas: hojas candidatas, patrones de 4,75 mm, discos


def h_nesting(wb):
    ws = g.hoja(wb, "Nesting", [30, 11, 11, 9, 11, 10, 12, 12, 10, 10, 11, 10, 11, 12, 30])
    titulo(ws, "Nesting: ¿es el mejor corte?",
           "Todas las hojas cotizadas y todos los patrones de guillotina con el 1.er corte por el lado corto (hasta "
           "tres etapas). Los patrones salen de enumerarlos todos; la mezcla, de programación lineal.")
    seccion(ws, NA0 - 2, "A) La hoja que conviene para cada formato (corte en dos etapas)")
    g.cab(ws, NA0 - 1, ["Formato", "Hoja: lado corto (mm)", "lado largo (mm)", "Espesor (mm)", "Cotización",
                        "USD/kg", "Franjas × piezas", "Piezas por hoja", "Aprovech.", "Masa de la hoja (kg)",
                        "kg por cuerpo", "USD por cuerpo", "La más barata", "En Corte"])
    for j, (fm, W, C, q) in enumerate(CAND):
        r = NA0 + j
        i = FORM.index(fm)
        al, de, e = pz("F", i), pz("G", i), pz("E", i)

        def feas(a, b):
            return fx(f"AND({a}<={{tope1}},{a}<={{L2}},B{r}<={{L1}},OR({b}<={{tope2}},AND(INT(B{r}/{b})=1,"
                      f"B{r}-{b}<={{tope2}})))")
        p1 = f"IF({feas(al, de)},INT(C{r}/{al})*INT(B{r}/{de}),0)"
        p2 = f"IF({feas(de, al)},INT(C{r}/{de})*INT(B{r}/{al}),0)"
        c(ws, f"A{r}", fm)
        c(ws, f"B{r}", W, F0, dato=True)
        c(ws, f"C{r}", C, F0, dato=True)
        c(ws, f"D{r}", f"={e}", "0.00")
        c(ws, f"E{r}", q)
        c(ws, f"F{r}", "=" + N["cot"][q], "0.0000")
        c(ws, f"G{r}", f'=IF({p1}>={p2},INT(C{r}/{al})&" × "&INT(B{r}/{de}),INT(C{r}/{de})&" × "&INT(B{r}/{al}))')
        c(ws, f"H{r}", f"=MAX({p1},{p2})", F0)
        c(ws, f"I{r}", f"=IF(H{r}=0,0,H{r}*{al}*{de}/(B{r}*C{r}))", FP)
        c(ws, f"J{r}", fx(f"=B{r}*C{r}*D{r}*{{dens}}/1000000"), F1)
        c(ws, f"K{r}", f"=IF(H{r}=0,0,J{r}/H{r})", F3)
        c(ws, f"L{r}", f"=IF(H{r}=0,0,K{r}*F{r})", F2)
        filas = [NA0 + k for k, x in enumerate(CAND) if x[0] == fm]
        rng = f"$L${filas[0]}:$L${filas[-1]}"
        c(ws, f"M{r}", f'=IF(L{r}=MIN({rng}),"Sí","")', neg=True)
        c(ws, f"N{r}", f'=IF(AND(B{r}=Corte!$B${HS0 + i},C{r}=Corte!$C${HS0 + i},F{r}={precio(e, "pH")}),"Sí","")')
    r = NA0 + len(CAND)
    g.nota(ws, r, "Cotizaciones de la hoja Datos (sección 5). 25 y 50 kg y 70 y 100 kg se pueden combinar en la misma "
                  "hoja: secciones B y C.")
    # B) 3,2 mm combinada
    rb = NP0 - 3
    seccion(ws, rb, "B) 3,2 mm: 25 y 50 kg en la misma hoja de 1.500 × 3.000 (corte en tres etapas)")
    g.nota(ws, rb + 1, "En una franja de 995,5 mm (desarrollo del 50 kg) entran, a lo largo, una pieza de 50 kg (640) "
                       "y un tramo de 858,5 que se corta a lo largo en dos piezas de 25 kg (493 + 493).")
    g.cab(ws, NP0 - 1, ["Patrón", "Franjas", "", "", "25 kg por hoja", "50 kg por hoja", "Golpes G1 por hoja",
                        "Golpes G2 por hoja", "Tandas de G2 por hoja", "Etapas", "Aprovech."])
    ws.merge_cells(start_row=NP0 - 1, start_column=2, end_row=NP0 - 1, end_column=4)
    a25, d25, a50, d50 = pz("F", 4), pz("G", 4), pz("F", 5), pz("G", 5)
    for k, (nom, txt, n25, n50, et, g1, g2, ta) in enumerate(PAT32):
        r = NP0 + k
        c(ws, f"A{r}", nom, neg=True)
        c(ws, f"B{r}", txt)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        for col, v, fmt in (("E", n25, F0), ("F", n50, F0), ("G", g1, F2), ("H", g2, F2), ("I", ta, F2),
                            ("J", et, F0)):
            c(ws, f"{col}{r}", v, fmt, dato=True)
        c(ws, f"K{r}", f"=(E{r}*{a25}*{d25}+F{r}*{a50}*{d50})/(Corte!$B${HS0 + 4}*Corte!$C${HS0 + 4})", FP)
    r = NP0 + len(PAT32)
    g.nota(ws, r, "Golpes de G2 en P1: cada tanda de 2 franjas lleva 2 golpes (pieza de 50 kg y tramo) y cada tanda "
                  "de 2 tramos girados, 2 más (493 y el recorte). El plan usa P1, P3 y P4: P2 da las mismas hojas.")
    rc = r + 2
    seccion(ws, rc, "Controles de P1 (con las medidas del plano)")
    W, C = f"Corte!$B${HS0 + 4}", f"Corte!$C${HS0 + 4}"
    ctl = [("Dos piezas de 25 kg a lo ancho de la franja de 50 kg", f"=2*{a25}", f"={d50}", "<="),
           ("Pieza de 50 kg + tramo de 25 kg a lo largo de la franja", f"={a50}+{d25}", f"={W}", "<="),
           ("Tres franjas a lo largo de la hoja", f"=3*{d50}", f"={C}", "<="),
           ("Franja dentro del tope de G1", f"={d50}", fx("={tope1}"), "<="),
           ("Corte a lo largo (493) dentro del tope de G2", f"={a25}", fx("={tope2}"), "<=")]
    g.cab(ws, rc + 1, ["Control", "Medida (mm)", "Límite (mm)", "", "Resultado"])
    for k, (t, v, lim, op) in enumerate(ctl):
        r = rc + 2 + k
        c(ws, f"A{r}", t)
        c(ws, f"B{r}", v, "0.0")
        c(ws, f"C{r}", lim, "0.0")
        c(ws, f"E{r}", f'=IF(B{r}<=C{r},"OK (sobran "&TEXT(C{r}-B{r},"0.0")&" mm)","No entra")')
    r = rc + 2 + len(ctl)
    g.nota(ws, r, "Sobran 1,5 mm a lo largo de la franja: la chapa laminada en caliente viene con tolerancia de ancho "
                  "+0 (nunca más angosta), así que el recorte final siempre existe.")
    # plan del año de referencia
    rp = r + 2
    seccion(ws, rp, "Plan del año de referencia: cada formato en su hoja o combinados (mezcla de patrones)")
    g.cab(ws, rp + 1, ["Concepto", "Separados (2 etapas)", "Combinados", "Diferencia"])
    U4, U5 = (anio_ref(f"Cálculo!$C${R[k]}:$L${R[k]}") for k in ("U4", "U5"))
    pk = {k: anio_ref(f"Cálculo!$C${R['B_p' + str(k)]}:$L${R['B_p' + str(k)]}") for k in KP}
    masa = f"Corte!$W${HS0 + 4}"
    filas = [("Cuerpos de 25 kg", f"={U4}", f"={U4}", F0),
             ("Cuerpos de 50 kg", f"={U5}", f"={U5}", F0),
             ("Hojas de 3,2 mm", f"={U4}/Corte!$I${HS0 + 4}+{U5}/Corte!$I${HS0 + 5}",
              "=" + "+".join(pk[k] for k in KP), F1),
             ("Acero (t)", f"=B{{h}}*{masa}/1000", f"=C{{h}}*{masa}/1000", F1),
             ("Aprovechamiento", f"=({U4}*{a25}*{d25}+{U5}*{a50}*{d50})/(B{{h}}*{W}*{C})",
              f"=({U4}*{a25}*{d25}+{U5}*{a50}*{d50})/(C{{h}}*{W}*{C})", FP),
             ("Costo del acero (USD)", f"=B{{t}}*1000*{precio(pz('E', 4), 'pH')}",
              f"=C{{t}}*1000*{precio(pz('E', 4), 'pH')}", F0)]
    rh = rp + 2 + 2
    rt = rp + 2 + 3
    for k, (t, v1, v2, fmt) in enumerate(filas):
        r = rp + 2 + k
        c(ws, f"A{r}", t)
        c(ws, f"B{r}", v1.replace("{h}", str(rh)).replace("{t}", str(rt)), fmt)
        c(ws, f"C{r}", v2.replace("{h}", str(rh)).replace("{t}", str(rt)), fmt)
        c(ws, f"D{r}", f"=C{r}-B{r}", fmt)
    r = rp + 2 + len(filas)
    g.nota(ws, r + 1, "Plan: P1 hasta cubrir el 25 kg (o el 50 kg, si es menos de la mitad) y el resto en P3 o P4. Es el "
                      "mínimo de hojas (Word «Esquema de corte en guillotina», 5.2).")
    c(ws, f"A{r}", "Combinar conviene", neg=True)
    c(ws, f"B{r}", f'=IF(C{rt}<B{rt},"Sí: ahorra "&TEXT(B{rh}-C{rh},"0")&" hojas y USD "&TEXT(B{rt + 2}-C{rt + 2},'
                   f'"#,##0")&" en el año","No")', neg=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    # promedios del plan para los Aprovechamientos (25 y 50 kg)
    re_ = r + 2
    seccion(ws, re_, "Promedio por hoja de 3,2 mm con el plan elegido (va a Corte, tabla para los Aprovechamientos)")
    tot = "(" + "+".join(pk[k] for k in KP) + ")"
    prom = [("Piezas por hoja", f"=IF({tot}=0,0,({U4}+{U5})/{tot})"),
            ("Golpes G1 por hoja", "=IF({t}=0,0,(" + "+".join(f"{pk[k]}*G{NP0 + k - 1}" for k in KP) + ")/{t})"),
            ("Golpes G2 por hoja", "=IF({t}=0,0,(" + "+".join(f"{pk[k]}*H{NP0 + k - 1}" for k in KP) + ")/{t})"),
            ("Tandas de G2 por hoja", "=IF({t}=0,0,(" + "+".join(f"{pk[k]}*I{NP0 + k - 1}" for k in KP) + ")/{t})")]
    for k, (t, v) in enumerate(prom):
        r = re_ + 1 + k
        c(ws, f"A{r}", t)
        c(ws, f"B{r}", v.replace("{t}", tot), F2)
    assert re_ + 1 == NPROM, re_
    # C) 4,75 mm
    seccion(ws, NC0 - 2, "C) 4,75 mm: 70 y 100 kg en la hoja de 1.500 × 3.000")
    g.cab(ws, NC0 - 1, ["Patrón", "70 kg por hoja", "100 kg por hoja", "", "Aprovech."])
    a70, d70, a100, d100 = pz("F", 6), pz("G", 6), pz("F", 7), pz("G", 7)
    for k, (t, n70, n100) in enumerate(PAT475):
        r = NC0 + k
        c(ws, f"A{r}", t)
        c(ws, f"B{r}", n70, F0, dato=True)
        c(ws, f"C{r}", n100, F0, dato=True)
        c(ws, f"E{r}", f"=(B{r}*{a70}*{d70}+C{r}*{a100}*{d100})/(Corte!$B${HS0 + 7}*Corte!$C${HS0 + 7})", FP)
    g.nota(ws, NC0 + len(PAT475), "La mezcla óptima es la del plan actual: hojas mixtas hasta agotar los 70 kg y el resto "
                                  "de 100 kg solos (77,5 % en 2035).")
    # D) discos
    seccion(ws, ND0 - 2, "D) Discos: fleje de una fila o de dos filas al tresbolillo (detalle en Corte, sección C)")
    g.cab(ws, ND0 - 1, ["Pieza", "Ancho 1 fila (mm)", "Aprovech. 1 fila", "Ancho 2 filas (mm)", "Aprovech. 2 filas",
                        "kg de fleje por año, 1 fila", "kg, 2 filas", "Ahorro (kg)"])
    for i, d in enumerate(DC.DISC):
        r, rc_ = ND0 + i, DC0 + i
        ud = anio_ref(f"Cálculo!$C${R[f'D{i}']}:$L${R[f'D{i}']}")
        c(ws, f"A{r}", f"={'Corte!A' + str(rc_)}")
        c(ws, f"B{r}", f"=Corte!$T${rc_}", "0.0")
        c(ws, f"C{r}", f"=Corte!$U${rc_}", FP)
        c(ws, f"D{r}", f"=Corte!$V${rc_}", "0.0")
        c(ws, f"E{r}", f"=Corte!$W${rc_}", FP)
        c(ws, f"F{r}", f"={ud}*Corte!$X${rc_}", F0)
        c(ws, f"G{r}", f"={ud}*Corte!$Y${rc_}", F0)
        c(ws, f"H{r}", f"=F{r}-G{r}", F0)
    r = ND0 + len(DC.DISC)
    c(ws, f"A{r}", "Total", neg=True)
    for col in "FGH":
        c(ws, f"{col}{r}", f"=SUM({col}{ND0}:{col}{r - 1})", F0, neg=True)
    g.nota(ws, r + 1, "Dos filas al tresbolillo: centros de una fila corridos medio paso y separados paso × √3/2, así "
                      "entre discos de filas vecinas queda el mismo puente de avance.")
    ws.freeze_panes = "B4"


# ------------------------------------------------------------------ Cálculo
def _metros_c(gr, col):
    info = GRP[gr]["info"]
    if gr == "4,75":
        return (f"=({col}{R['U6']}*{pz('F', 6)}+{col}{R['U7']}*{pz('F', 7)})/1000/{info[0][2]}")
    partes = []
    for k, av, h in info:
        u = f"{col}{U0 + k}"
        partes.append(f"IF({h}=0,IF({u}>0,1000000000000,0),{u}*{av}/1000/{h})")
    return "=MAX(" + ",".join(partes) + ")"


def h_calculo(wb):
    ws = g.hoja(wb, "Cálculo", [62, 9] + [12] * 10)
    titulo(ws, "Cálculo año por año de cada alternativa",
           "No se edita: sale de Datos, Demanda, Piezas y Corte. Los costos van en USD sin IVA.")
    g.cab(ws, 4, ["Concepto", "Unidad"] + [str(a) for a in ANIOS])
    c(ws, "A5", "Período k (para descontar)")
    c(ws, "A6", "Factor de descuento 1/(1+i)^k")
    for i, col in enumerate(YC):
        c(ws, f"{col}5", i + 1, F0)
        c(ws, f"{col}6", fx(f"=1/(1+{{tasa}})^{col}5"), F3)
    seccion(ws, U0 - 1, "Unidades a fabricar (cuerpos)")
    for i, f in enumerate(FORM):
        c(ws, f"A{U0 + i}", f)
        c(ws, f"B{U0 + i}", "u")
        for col in YC:
            c(ws, f"{col}{U0 + i}", f"=Demanda!{col}{6 + i}", F0)
    for i, d in enumerate(DC.DISC):
        r = D0 + i
        c(ws, f"A{r}", f"{d['pieza']} {d['formato']}")
        c(ws, f"B{r}", "u")
        for col in YC:
            c(ws, f"{col}{r}", f"={col}{U0 + FORM.index(d['formato'])}", F0)
    c(ws, f"A{NETO}", "Acero neto en piezas (cuerpos + cúpulas + fondos)", neg=True)
    c(ws, f"B{NETO}", "t")
    for col in YC:
        c(ws, f"{col}{NETO}", f"=(SUMPRODUCT({col}${U0}:{col}${U0 + 7},Piezas!$H${PZ0}:$H${PZ0 + 7})+"
                              f"SUMPRODUCT({col}${D0}:{col}${D0 + 7},Piezas!$S${DS0}:$S${DS0 + 7}))/1000", F1, neg=True)
    for i, (k, lab, uni, fmt, tpl) in enumerate(PLAN):
        if k is None:
            seccion(ws, R[f"T{i}"], lab)
            continue
        r = R[k]
        tot = lab.startswith("COSTO ANUAL")
        c(ws, f"A{r}", lab, neg=tot)
        c(ws, f"B{r}", uni)
        for col in YC:
            if k.startswith("C_m") and k[3:] in GK.values():
                gr = next(x for x in GRUPOS if GK[x] == k[3:])
                v = _metros_c(gr, col)
            elif k.startswith("C_t") and k[3:] in GK.values():
                gr = next(x for x in GRUPOS if GK[x] == k[3:])
                v = fx(f"={col}{R['C_m' + k[3:]]}*Corte!$D${GRP[gr]['fila']}/1000/(1-{{despunte}})")
            else:
                v = tpl.format(c=col, **R, **{a: b for a, b in N.items() if isinstance(b, str)})
            c(ws, f"{col}{r}", v, fmt, neg=tot)
    ws.freeze_panes = "C5"


# ------------------------------------------------------------------ Resumen
COMPRA = {"A": ("Bobina de ancho estándar por formato, línea de corte transversal", "Fleje a medida por disco"),
          "B": ("Hojas estándar, dos guillotinas en serie", "Tiras de hoja cargadas a mano"),
          "C": ("Bobina madre por espesor, cortada en flejes propios", "Flejes propios de la bobina madre"),
          "D": ("Hojas estándar, dos guillotinas en serie", "Fleje a medida por disco"),
          "Ref": ("Cuerpos comprados cortados a medida", "Fleje a medida por disco")}
CLAVE = {"A": "A", "B": "B", "C": "C", "D": "D", "Ref": "R"}


def h_resumen(wb):
    ws = g.hoja(wb, "Resumen", [46, 24, 24, 24, 24, 24, 3, 40])
    titulo(ws, "Corte de chapa: comparación de alternativas 2026-2035",
           "Cuerpos de 1 a 100 kg, cúpulas y fondos de 1 a 10 kg. Valor actual = inversión + costos de 10 años "
           "descontados. Menor es mejor.")
    g.cab(ws, 4, ["Indicador"] + [OPC_NOM[o] for o in OPC])
    filas = [("Cuerpos: se compra", "txt0", None), ("Cúpulas y fondos: se compra", "txt1", None),
             ("Inversión en equipos (USD)", "inv", F0), ("Acero comprado 2026-2035 (t)", "acero", F0),
             ("Aprovechamiento de chapa 2026-2035", "aprov", FP), ("Horas-hombre de corte en el año de referencia",
                                                                   "horas", F0),
             ("Costo anual en el año de referencia (USD)", "anual", F0),
             ("Valor actual a 10 años (USD)", "va", F0), ("Ranking (1 = conviene; la referencia no corta en planta)",
                                                          "rank", F0),
             ("Diferencia contra la que conviene (USD)", "dif", F0)]
    rva = 4 + 1 + [k for _, k, _ in filas].index("va")
    for i, (lab, k, fmt) in enumerate(filas):
        r = 5 + i
        c(ws, f"A{r}", lab, neg=k == "va")
        for j, o in enumerate(OPC):
            col = L(2 + j)
            p = CLAVE[o]
            tot = f"{p}_t" if o != "C" else "C_t"
            h = {"A": "A_h", "B": "B_h", "C": "C_h", "D": "D_h", "Ref": "R_h"}[o]
            cost = {"A": "A_cost", "B": "B_cost", "C": "C_cost", "D": "D_cost", "Ref": "R_cost"}[o]
            v = {
                "txt0": COMPRA[o][0], "txt1": COMPRA[o][1],
                "inv": f"={N['inv' + o]}",
                "acero": f"=SUM(Cálculo!$C${R[tot]}:$L${R[tot]})",
                "aprov": f"=SUM(Cálculo!$C${NETO}:$L${NETO})/{col}{5 + 3}",
                "horas": "=" + anio_ref(f"Cálculo!$C${R[h]}:$L${R[h]}"),
                "anual": "=" + anio_ref(f"Cálculo!$C${R[cost]}:$L${R[cost]}"),
                "va": f"={col}7+SUMPRODUCT(Cálculo!$C${R[cost]}:$L${R[cost]},Cálculo!$C$6:$L$6)",
                "rank": f"=RANK({col}{rva},$B${rva}:$E${rva},1)" if o != "Ref" else "—",
                "dif": f"={col}{rva}-MIN($B${rva}:$E${rva})",
            }[k]
            x = c(ws, f"{col}{r}", v, fmt, neg=k == "va")
            if k.startswith("txt"):
                x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        ws.row_dimensions[r].height = 30 if k.startswith("txt") else 16
    r = 5 + len(filas) + 1
    nombres = ",".join(f'"{OPC_NOM[o]}"' for o in OPC[:4])
    c(ws, f"A{r}", "Conviene (de las que cortan en planta)", neg=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    c(ws, f"B{r}", f"=INDEX({{{nombres}}},MATCH(MIN($B${rva}:$E${rva}),$B${rva}:$E${rva},0))", neg=True)
    rd = R["B_uc"]
    c(ws, f"A{r + 1}", "Margen de D: suba del precio de la hoja que la iguala con A")
    c(ws, f"B{r + 1}", f"=($B${rva}-$E${rva})/SUMPRODUCT(Cálculo!$C${rd}:$L${rd},Cálculo!$C$6:$L$6)", FP)
    rf = R["A_uf"]
    c(ws, f"A{r + 2}", "Margen de D: suba del precio del fleje que la iguala con B")
    c(ws, f"B{r + 2}", f"=($C${rva}-$E${rva})/SUMPRODUCT(Cálculo!$C${rf}:$L${rf},Cálculo!$C$6:$L$6)", FP)
    g.nota(ws, r + 4, "Datos editables en la hoja Datos (gris claro). Demanda por IMPORTRANGE (hoja Links). "
                      "Medidas de cuerpos y discos del plano (hoja Piezas).")
    ch = BarChart()
    ch.title = "Valor actual a 10 años (USD): menor es mejor"
    ch.add_data(Reference(ws, min_col=2, max_col=6, min_row=rva, max_row=rva), from_rows=True,
                titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=2, max_col=6, min_row=4, max_row=4))
    ch.legend = None
    ch.y_axis.scaling.min = 0
    ch.y_axis.title = "USD"
    ch.height, ch.width = 7.5, 22
    ws.add_chart(ch, f"A{r + 6}")
    ws.freeze_panes = "B5"


def main():
    wb = Workbook()
    wb.remove(wb.active)
    h_resumen_placeholder = wb.create_sheet("Resumen")
    h_datos(wb)
    h_links(wb)
    h_demanda(wb)
    h_piezas(wb)
    h_corte(wb)
    h_nesting(wb)
    h_calculo(wb)
    wb.remove(h_resumen_placeholder)
    h_resumen(wb)
    orden = ["Links", "Resumen", "Datos", "Demanda", "Piezas", "Corte", "Nesting", "Cálculo"]
    wb._sheets = [wb[n] for n in orden]
    wb.active = 1
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    # valores calculados guardados en el archivo (vista previa de Drive) y IMPORTRANGE como lo exporta Google
    import actualizar_manuales_carros as amc
    amc.recalc(SALIDA)
    amc.arreglar_dummy(SALIDA)
    print("guardado", SALIDA)


if __name__ == "__main__":
    main()
